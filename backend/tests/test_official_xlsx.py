"""地政局正式範本 xlsx 寫入器：金山範本案與決賽樹林案都要能填進範本，格位照 docs/12。"""
import io
import json
import zipfile
from pathlib import Path

import pytest
from openpyxl import load_workbook

from app.engine.rules import load_ruleset
from app.engine.tables import run_case
from app.output.official_xlsx import fill_table3, fill_table4, fill_table5, official_zip, workbook_bytes

FIX = Path(__file__).resolve().parents[2] / "fixtures"


def _load(name: str) -> dict:
    d = json.loads((FIX / name).read_text(encoding="utf-8"))
    return {k: d[k] for k in ("case", "sections", "subject_parcel", "comparables")}


def _run(data: dict):
    reg = load_ruleset(data["case"]["rulesets"]["regional"])
    ind = load_ruleset(data["case"]["rulesets"]["individual"])
    return reg, ind, run_case(reg, ind, data)


def _ws(wb):
    return load_workbook(io.BytesIO(workbook_bytes(wb))).active


def test_jinshan_sample_fills_official_templates():
    data = _load("sample_case_P002-00.json")
    reg, ind, res = _run(data)
    meta = {"appraiser": "王小明", "fill_date": "114 年 9 月 18 日", "notes": {"subject": "S", "comparables": {"1": "C1"}, "case": "ALL"}}
    t4 = _ws(fill_table4(data, res["table4"], ind, meta))
    assert t4["P1"].value == "1140901-99-001" and t4["L1"].value == "1140901"
    assert t4["G5"].value == 184763 and t4["J6"].value == 2.0 and t4["G8"].value == "P002-00"
    assert t4["D17"].value == "金山國小" and t4["E17"].value == 150 and t4["H17"].value == 100          # 接近條件：名稱＋距離分格
    assert t4["E16"].value == 18 and t4["H16"].value == 6                                              # 面前道路寬度來自 front_road.width_m（比準地中山路 18、比較標的金包里街 6）
    assert t4["G29"].value == 13.0 and t4["G32"].value == 212958                                        # 個別因素合計、比準地比較價格（範本驗收值）
    assert t4["D33"].value == "S" and t4["G33"].value == "C1" and t4["D34"].value == "ALL"
    assert "114 年 9 月 18 日" in t4["A35"].value and t4["K36"].value == "不動產估價師：王小明"
    t5 = _ws(fill_table5(data, res["table5"], reg, meta))
    assert t5["A2"].value == "案號：1140901-99-001" and t5["C3"].value == "P002-00" and t5["E3"].value == "P002-00"
    assert t5["C5"].value == "優" and t5["E5"].value == "優" and t5["G5"].value == 0.0
    assert t5["G42"].value == 0.0                                                                       # 範本案區域因素總修正數 0
    t3 = fill_table3(data, reg, meta)
    assert t3.sheetnames == ["表3 P002-00"]
    ws = _ws(t3)
    assert ws["B3"].value == "1140901" and ws["G3"].value == "P002-00" and ws["H6"].value == "70%"
    assert ws["F15"].value == "國光客運金山站" and ws["E15"].value == "●" and ws["J15"].value == 300 and "●本區段外" in ws["H15"].value
    assert ws["R14"].value == "名稱：金山變電所" and ws["U14"].value == 700 and "●本區段外" in ws["T14"].value
    assert ws["I39"].value.startswith("●本區段內") and ws["F39"].value.startswith("傳統市場：")
    assert ws["R46"].value == "不動產估價師：王小明"


def test_shulin_competition_case_fills_all_four_sections_and_skips_no_adjust_items():
    data = _load("shulin_case_1110901.json")
    reg, ind, res = _run(data)
    t5s = res["table5"]
    assert set(t5s) == {1, 2, 3}
    for t in t5s.values():                                                      # 表5-1 備註：使用分區、建蔽率、容積率併同個別因素修正 → 免修正
        skipped = {r.rule_id: r for r in t.rows if r.rule_id in ("R1-2", "R1-3", "R1-4")}
        assert all(r.pct is None for r in skipped.values()) and all(r.subject_level for r in skipped.values())
        assert t.rows[[r.rule_id for r in t.rows].index("R1-6")].pct == 0.0     # 三級「有無限制建築」布林 → 「無」
        assert t.rows[[r.rule_id for r in t.rows].index("R4-1")].subject_level == "優"   # 建築基地改良四項 → 優
    meta = {"appraiser": "", "fill_date": "", "notes": {}}
    t5 = _ws(fill_table5(data, t5s, reg, meta))
    assert t5["G6"].value == "-" and t5["G7"].value == "-" and t5["G8"].value == "-"
    assert t5["B40"].value == "其他影響因素" and t5["G42"].value == t5s[1].total_pct
    t3 = fill_table3(data, reg, meta)
    assert t3.sheetnames == ["表3 P001-00", "表3 P002-00", "表3 P003-00", "表3 P004-00"]
    ws = load_workbook(io.BytesIO(workbook_bytes(t3)))["表3 P002-00"]
    assert ws["G11"].value == "樹人街" and ws["J11"].value == 7 and ws["G12"].value == 6 and ws["H7"].value == "200%"
    assert ws["I23"].value == "部分規劃及闢建" and ws["R42"].value == "70%" and "●住宅用" in ws["Q44"].value
    assert "■整平或填挖基地" in ws["E31"].value and "■埋設管道" in ws["E32"].value and "□水土保持" in ws["E31"].value
    t4 = _ws(fill_table4(data, res["table4"], ind, meta))
    assert t4["G4"].value == "新北市樹林區樹德段284地號" and t4["G5"].value == 130167 and t4["J6"].value == 5.96
    assert t4["O5"].value == 170909 and t4["R6"].value == 5.49 and t4["G8"].value == "P002-00"
    blob = official_zip(data, res, reg, ind, meta)
    names = zipfile.ZipFile(io.BytesIO(blob)).namelist()
    assert len(names) == 3 and all(n.startswith("1110901-99-XXX_") and n.endswith(".xlsx") for n in names)


@pytest.mark.parametrize("obs,expected", [([], "劣"), (["整平或填挖基地"], "稍劣"), (["a", "b"], "普通"), (["a", "b", "c"], "稍優"), ({"a": True, "b": True, "c": True, "d": True}, "優")])
def test_enum_count_normalization(obs, expected):
    from app.engine.rules import grade
    reg = load_ruleset("shulin_residential_regional")
    rule = next(r for r in reg.rules if r.id == "R4-1")
    assert grade(rule, obs) == expected
