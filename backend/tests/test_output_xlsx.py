"""表5 / 表4 Excel 輸出：用 fixture 產出，openpyxl 讀回，數值與 expected 一致；免修正顯示「-」；布林細項顯示「無」。"""
import io
import json
from pathlib import Path

import pytest
from openpyxl import load_workbook

from app.output.xlsx import export_xlsx

ROOT = Path(__file__).resolve().parents[2]
FIX = json.loads((ROOT / "fixtures" / "sample_case_P002-00.json").read_text(encoding="utf-8"))


@pytest.fixture(scope="module")
def wb():
    buf = io.BytesIO()
    data = {k: FIX[k] for k in ("case", "sections", "subject_parcel", "comparables")}
    export_xlsx(data, buf, meta={"appraiser": "測試估價師", "fill_date": "114 年 09 月 18 日",
                                 "notes": {"case": "全案備註", "comparables": {"1": "比較標的1備註"}}})
    buf.seek(0)
    return load_workbook(buf)


def _rows(ws):
    return [[c.value for c in row] for row in ws.iter_rows()]


def _find_row(rows, col, text):
    for r in rows:
        v = r[col]
        if isinstance(v, str) and v.replace("\n", "") == text:
            return r
    raise AssertionError(f"找不到列 {text}")


def test_sheets_and_titles(wb):
    assert wb.sheetnames == ["表1", "表5-2", "表4"]
    assert wb["表5-2"]["A1"].value.startswith("表5-2  影響地價區域因素分析明細表（商業用地）")
    assert wb["表4"]["A1"].value.startswith("表4  比較法調查估價表")
    assert "1140901-99-001" in wb["表5-2"]["A2"].value


def test_table5_levels_subtotals_total(wb):
    ws = wb["表5-2"]
    rows = _rows(ws)
    from app.engine.rules import load_ruleset
    rs = load_ruleset("jinshan_commercial_regional")
    exp = FIX["expected"]["table5"]["subject_levels"]
    for r in rs.rules:
        if r.is_manual:
            continue
        row = _find_row(rows, 1, r.name)
        want = exp[r.id]
        if r.criteria["type"] == "boolean":
            want = "無" if want == r.criteria["false_level"] else "有"
        assert row[3] == want and row[5] == want, r.id     # 比準地 / 比較標的1 等級
        assert row[6] == 0.0, r.id                          # 修正百分比
    subtotals = [r for r in rows if r[1] == "百分比小計"]
    assert len(subtotals) == 8 and all(r[4] == 0.0 for r in subtotals)
    total = _find_row(rows, 1, "=(1)+(2)+(3)+(4)+(5)+(6)+(7)+(8)")
    assert total[4] == 0.0
    assert ws["G5"].number_format == "0.00"
    assert "不動產估價師：測試估價師" in [v for r in rows for v in r if isinstance(v, str)]


def test_table4_values(wb):
    ws = wb["表4"]
    rows = _rows(ws)
    exp = FIX["expected"]["table4"]["comparable_1"]
    for item, pct in exp["individual"].items():
        row = next(r for r in rows if isinstance(r[2], str) and r[2].startswith(f"{item}") and r[2][len(item):len(item) + 1] not in "0123456789")
        assert row[9] == pct, item
    r14 = _find_row(rows, 2, "14面前道路寬度")
    assert r14[3:6] == ["中山路", 18, "M"] and r14[6:9] == ["金包里街", 6, "M"]
    r20 = _find_row(rows, 2, "20嫌惡設施(類型)")
    assert r20[3:6] == ["金山第一公墓", 260, "M"] and r20[9] == 3.0
    r25 = _find_row(rows, 2, "25有無禁限建")
    assert r25[3] == "無" and r25[6] == "無"
    r6 = _find_row(rows, 1, "6其他")
    assert r6[3] == "-" and r6[6] == "-" and r6[9] == "-"          # 免修正是「-」不是 0
    assert _find_row(rows, 1, "合計")[6] == exp["individual_total_pct"]
    rabs = _find_row(rows, 1, "調整百分率絕對值加總")
    assert rabs[6] == exp["abs_sum_pct"] and rabs[8] == exp["similarity"]
    rtrial = _find_row(rows, 1, "試算價格")
    assert rtrial[6] == exp["trial_price"] and rtrial[8] == exp["weight_pct"]
    assert _find_row(rows, 1, "比準地比較價格")[6] == FIX["expected"]["table4"]["subject_comparison_price"]
    assert _find_row(rows, 0, "土地正常單價")[6] == 184763
    assert _find_row(rows, 0, "交易日期")[6] == "114.05.28" and _find_row(rows, 0, "交易日期")[9] == 2.0
    assert abs(_find_row(rows, 0, "調整至估價基準日單價(元/M2)")[6] - exp["price_at_valuation_date"]) <= 1   # 範本尾數疑點，容差 ±1
    assert _find_row(rows, 0, "地價區段")[9] == 0.0
    assert ws.page_setup.orientation == "landscape"
    assert _find_row(rows, 1, "全案")[3] == "全案備註"
    assert _find_row(rows, 1, "比準地或各比較標的")[6] == "比較標的1備註"


def test_table1_layout_and_levels(wb):
    ws = wb["表1"]
    rows = _rows(ws)
    assert ws["A1"].value.startswith("表1  地價區段勘查表") and ws["A2"].value == "新北市金山區"
    assert ws["C3"].value == "1140901" and ws["F3"].value == "P002-00" and "金包里街" in ws["J3"].value
    texts = [v for r in rows for v in r if isinstance(v, str)]
    assert "名稱：金山變電所 ○本區段內 ●本區段外(距 700 M)" in texts
    assert "名稱：金包里老街 ●本區段內 ○本區段外(距  M)" in texts
    assert any("● 國光客運金山站 ○本區段內 ●本區段外(距 300 M)" in t for t in texts)
    assert any("●密集" in t and "○非常密集" in t for t in texts)
    assert "90%以上作為店舖" in texts and "有排水系統不易淹水" in texts
    # 等級數字欄：都市計畫 1/2、排水 2/5、交流道 5/5（左半）、停車場地 2/5、金融機構 2/5（右半）
    def lv(label, left=True):
        col = 3 if left else 12
        row = _find_row(rows, col, label)
        return (row[col - 2], row[col - 1])
    assert lv("都市計畫(內外)") == (1, 2)
    assert lv("保（排）水之良否") == (2, 5)
    assert lv("交流道") == (5, 5)
    assert lv("停車場地", left=False) == (2, 5)
    assert lv("金融機構", left=False) == (2, 5)
    assert lv("百貨公司", left=False) == (5, 5)
    assert any("○住商混合" in t for t in texts)   # fixture 沒有土地利用現況 → 全部未勾


def test_workbook_figures_sheet():
    """圖說工作表：三張 PNG 各佔一段（不抓底圖）。"""
    from app.maps.render import case_figures
    from app.output.xlsx import build_workbook
    d = json.loads((ROOT / "fixtures" / "sample_case_P002-00.json").read_text(encoding="utf-8"))
    g = json.loads((ROOT / "fixtures" / "sample_geometry_P002-00.json").read_text(encoding="utf-8"))
    data = {k: d[k] for k in ("case", "sections", "subject_parcel", "comparables")}
    for sid, sec in g["sections"].items():
        data["sections"][sid] = {**data["sections"][sid], **sec}
    data["subject_parcel"] = {**data["subject_parcel"], **g["subject_parcel"]}
    figs = case_figures(data, basemap=False)
    assert [m for m, _, _ in figs] == ["sketch", "zoning", "section"] and all(png[:4] == b"\x89PNG" for _, _, png in figs)
    wb = build_workbook(data, figures=figs)
    assert wb.sheetnames[3:] == ["地價區段略圖", "地價使用分區圖", "地價區段圖"] and all(len(wb[n]._images) == 1 for n in wb.sheetnames[3:])
    from app.output.grid import grids_to_pdf, workbook_grids
    grids = workbook_grids(wb)
    assert [g["title"] for g in grids] == ["表1", "表5-2", "表4"] and grids[0]["orientation"] == "portrait" and grids[1]["orientation"] == "portrait" and grids[2]["orientation"] == "landscape" and any(c["v"] == "213,000" or c["v"] == "212,958" for c in grids[2]["cells"])
    pdf = grids_to_pdf(grids, figs, footer="測試")
    assert pdf[:5] == b"%PDF-" and pdf.count(b"/Type /Page\n") + pdf.count(b"/Type /Page ") >= 6 or pdf.count(b"/Page") >= 6
