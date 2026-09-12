"""API：/api/adapt 自動判斷四種檔案、/api/export/xlsx 回傳可讀的 Excel、/api/run 與 /api/verify 仍正常。"""
import io
import json
from pathlib import Path

from fastapi.testclient import TestClient
from openpyxl import load_workbook

from app.adapters.excel_comparables import write_comparables_xlsx
from app.adapters.excel_parcels import write_parcels_xlsx
from app.main import app

ROOT = Path(__file__).resolve().parents[2]
FIX = json.loads((ROOT / "fixtures" / "sample_case_P002-00.json").read_text(encoding="utf-8"))
CASE = {k: FIX[k] for k in ("case", "sections", "subject_parcel", "comparables")}
client = TestClient(app)


def test_health_reports_llm_status():
    r = client.get("/api/health")
    assert r.status_code == 200 and r.json()["ok"] is True and "llm" in r.json()


def test_run_and_verify():
    r = client.post("/api/run", json=CASE)
    assert r.status_code == 200
    assert abs(r.json()["table4"]["subject_comparison_price"] - 212958) <= 1
    exp4 = FIX["expected"]["table4"]
    r = client.post("/api/verify", json=CASE | {"submitted_table4": {"comparables": {"1": exp4["comparable_1"]}, "subject_comparison_price": exp4["subject_comparison_price"]}})
    assert r.status_code == 200 and [f for f in r.json()["findings"] if f["severity"] == "error"] == []


def _upload(name: str, content: bytes, **form):
    return client.post("/api/adapt", files={"file": (name, content)}, data=form)


def test_adapt_parcels_auto():
    buf = io.BytesIO()
    write_parcels_xlsx([FIX["subject_parcel"]], buf, case_no="1140901-99-001")
    r = _upload("清冊.xlsx", buf.getvalue())
    assert r.status_code == 200, r.text
    j = r.json()
    assert j["kind"] == "parcels" and j["data"]["parcels"][0]["parcel_id"] == "金美段489地號"
    assert j["missing_fields"] == [] and "stats" not in j


def test_adapt_comparables_auto():
    buf = io.BytesIO()
    write_comparables_xlsx(FIX["comparables"], buf)
    r = _upload("實例.xlsx", buf.getvalue())
    assert r.status_code == 200 and r.json()["kind"] == "comparables"
    assert r.json()["data"]["comparables"][0]["normal_unit_price"] == 184763


def test_adapt_rules_table_pdf_and_csv():
    pdf = (ROOT / "docs" / "reference" / "評價基準明細表範例.pdf").read_bytes()
    r = _upload("基準表.pdf", pdf, id_prefix="upload_test")
    assert r.status_code == 200 and r.json()["kind"] == "rules_table"
    assert set(r.json()["data"]["rulesets"]) == {"regional", "individual"}
    assert r.json()["data"]["rulesets"]["regional"]["id"] == "upload_test_regional"
    csv = (ROOT / "fixtures" / "jinshan_commercial_rules_table.csv").read_bytes()
    r = _upload("基準表.csv", csv, land_use="商業用地")
    assert r.status_code == 200 and len(r.json()["data"]["rulesets"]["individual"]["rules"]) == 19


def test_adapt_pdf_forms():
    pdf = (ROOT / "docs" / "reference" / "查估書表範本.pdf").read_bytes()
    r = _upload("查估書表.pdf", pdf, use_vision="never")
    assert r.status_code == 200, r.text
    j = r.json()
    assert j["kind"] == "pdf_forms" and j["data"]["case"]["case_no"] == "1140901-99-001"
    assert j["data"]["submitted"]["table4"]["subject_comparison_price"] == 212958
    # 抽出來的 schema 可以直接丟回 /api/verify
    body = {k: j["data"][k] for k in ("case", "sections", "subject_parcel", "comparables")}
    body["submitted_table5"] = {int(k): v for k, v in j["data"]["submitted"]["table5"].items()}
    body["submitted_table4"] = j["data"]["submitted"]["table4"]
    r2 = client.post("/api/verify", json=body)
    assert r2.status_code == 200, r2.text
    assert [f for f in r2.json()["findings"] if f["severity"] == "error"] == []


def test_adapt_bad_kind_and_unknown_file():
    assert _upload("x.xlsx", b"abc", kind="nope").status_code == 400
    assert _upload("x.txt", b"hello").status_code == 400


def test_export_xlsx():
    r = client.post("/api/export/xlsx", json=CASE | {"meta": {"appraiser": "王小明"}})
    assert r.status_code == 200, r.text
    assert r.headers["content-type"].startswith("application/vnd.openxmlformats")
    wb = load_workbook(io.BytesIO(r.content))
    assert wb.sheetnames[:3] == ["表1", "表5-2", "表4"]        # 有幾何時另附「圖說」工作表
    ws = wb["表4"]
    vals = [c.value for row in ws.iter_rows() for c in row if c.value is not None]
    assert 212958 in vals and "不動產估價師：王小明" in vals
