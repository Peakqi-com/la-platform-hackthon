"""收益法端點：搜尋收益實例 → 採用 → 核算 → 審查 → 表14 權重影響比準地地價 → 正式範本表2／表14。租賃資料用自製檔，不依賴 data/。"""
import io
import json
import zipfile
from pathlib import Path

from fastapi.testclient import TestClient
from openpyxl import load_workbook

from app.engine.tables import round_land_price
from app.main import app

FIX = Path(__file__).resolve().parents[2] / "fixtures" / "shulin_case_1110901.json"


def _rent_file(tmp_path, monkeypatch):
    def rec(i, date, unit, note=""):
        return {"id": f"RENT{i}", "season": "111S3", "district": "樹林區", "target": "土地", "date": date, "land_area": 50.0, "total_rent": unit * 50,
                "rent_ex_park": unit * 50, "unit_rent": float(unit), "zone": "住", "lots": [{"section": "東昇段", "zone": "都市：住", "area": 50.0, "lot_raw": "03660000"}], "note": note}
    f = tmp_path / "f_rent.json"
    f.write_text(json.dumps({"records": [rec(1, "1110630", 400), rec(2, "1110501", 300), rec(3, "1101001", 350, "土地上有未保存登記建物")], "seasons": ["111S3"]}), encoding="utf-8")
    monkeypatch.setenv("LVR_RENT_JSON", str(f))
    from app.market.rent import load_rent
    load_rent.cache_clear()


def test_income_endpoints_flow(tmp_path, monkeypatch):
    from app.market.rent import load_rent
    _rent_file(tmp_path, monkeypatch)
    client = TestClient(app)
    d = json.loads(FIX.read_text(encoding="utf-8"))
    data = {k: d[k] for k in ("case", "sections", "subject_parcel", "comparables")}
    data["subject_parcel"]["area_m2"] = 25.44
    cid = client.post("/api/cases", json={**data, "name": "收益法測試"}).json()["id"]
    try:
        v = client.get(f"/api/cases/{cid}/income").json()
        assert v["enabled"] is False and v["mode"] == "land" and v["result"] is None and v["defaults"]["deposit_months"]["value"] == 3
        s = client.post(f"/api/cases/{cid}/income/search", json={}).json()
        assert s["mode"] == "land" and [c["id"] for c in s["chosen"]] == ["RENT1", "RENT2", "RENT3"]
        assert s["chosen"][2]["flags"] and not s["chosen"][2]["excluded"]                         # 含地上建物：可採用但須情況調整
        a = client.post(f"/api/cases/{cid}/income/apply", json={"ids": ["RENT1", "RENT2", "RENT3"]}).json()
        view, res = a["view"], a["view"]["result"]
        assert res["mode"] == "land" and res["complete"] and res["income_price"] > 0
        assert res["examples"][0]["date_pct"] == 0.54                                             # 房租指數 111.06 → 111.09
        assert view["land_price"] == round_land_price(view["comparison_price"]) == view["decision"]["land_price"]   # 權重預設 0：比準地地價不變

        rec = a["rec"]
        fs = client.post("/api/verify", json={**rec["data"], "submitted_table5": None, "submitted_table4": None}).json()["findings"]
        assert any(f["table"] == "表2" and "情況調整" in f["location"] for f in fs)
        assert any(f["table"] == "表14" and "理由" in f["message"] for f in fs)

        data2 = rec["data"]
        data2["income"]["weights"] = {"comparison": 0.7, "income": 0.3}
        data2["income"]["reason"] = "收益實例為小面積土地租賃，可信度較低，比較價格權重較高"
        client.post("/api/cases", json={**data2, "id": cid})
        v2 = client.get(f"/api/cases/{cid}/income").json()
        assert v2["land_price"] == round_land_price(v2["comparison_price"] * 0.7 + v2["result"]["income_price"] * 0.3)

        t2 = load_workbook(io.BytesIO(client.get(f"/api/cases/{cid}/official/t2.xlsx").content))
        assert t2.sheetnames == ["94表2收益法 "]                                                  # 素地：附表成本法免填
        ws = t2.active
        assert ws["F4"].value == v2["result"]["est_monthly_rent"] and ws["F21"].value == v2["result"]["land_income_unit"] and ws["F16"].value == "-"
        t14 = load_workbook(io.BytesIO(client.get(f"/api/cases/{cid}/official/t14.xlsx").content)).active
        assert t14["J5"].value == v2["land_price"] and t14["H5"].value == v2["result"]["income_price"] and t14["I5"].value == 0.3
        names = zipfile.ZipFile(io.BytesIO(client.get(f"/api/cases/{cid}/official.zip").content)).namelist()
        assert len(names) == 5 and any("表2" in n for n in names) and any("表14" in n for n in names)
    finally:
        client.delete(f"/api/cases/{cid}")
        load_rent.cache_clear()


def test_official_t2_requires_income_enabled():
    client = TestClient(app)
    d = json.loads(FIX.read_text(encoding="utf-8"))
    cid = client.post("/api/cases", json={**{k: d[k] for k in ("case", "sections", "subject_parcel", "comparables")}, "name": "無收益法"}).json()["id"]
    try:
        assert client.get(f"/api/cases/{cid}/official/t2.xlsx").status_code == 422
        assert len(zipfile.ZipFile(io.BytesIO(client.get(f"/api/cases/{cid}/official.zip").content)).namelist()) == 3
    finally:
        client.delete(f"/api/cases/{cid}")
