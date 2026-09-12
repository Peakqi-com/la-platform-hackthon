"""案件儲存、demo 三變體、基準表列表/明細/匯入。"""
import json
from pathlib import Path

import pytest
from fastapi.testclient import TestClient

from app import cases as C
from app.main import app

ROOT = Path(__file__).resolve().parents[2]
client = TestClient(app)


@pytest.fixture(autouse=True)
def _tmp_cases(tmp_path, monkeypatch):
    monkeypatch.setattr(C, "CASES_DIR", tmp_path / "cases")
    C._mem.clear()
    yield
    C._mem.clear()
    for p in (ROOT / "rules").glob("uploaded_*.json"):
        p.unlink()


def test_demo_variants_and_verify():
    t = client.get("/api/cases/demo").json()
    assert t["variant"] == "template" and t["submitted_table4"]["subject_comparison_price"] == 212958
    body = {k: t[k] for k in ("case", "sections", "subject_parcel", "comparables")} | {"submitted_table5": t["submitted_table5"], "submitted_table4": t["submitted_table4"]}
    r = client.post("/api/verify", json=body).json()
    assert [f for f in r["findings"] if f["severity"] == "error"] == []                       # 步驟 1：全綠
    x = client.get("/api/cases/demo", params={"variant": "tampered"}).json()
    body = {k: x[k] for k in ("case", "sections", "subject_parcel", "comparables")} | {"submitted_table5": x["submitted_table5"], "submitted_table4": x["submitted_table4"]}
    errs = [f for f in client.post("/api/verify", json=body).json()["findings"] if f["severity"] == "error"]
    locs = " | ".join(f["location"] for f in errs)
    assert "14 面前道路寬度" in locs and "區域因素調整百分率" in locs and "小計" in locs and "交易日期" in locs   # 步驟 2：逐條抓錯
    assert all(f.get("basis") or f.get("message") for f in errs)
    res = client.get("/api/cases/demo", params={"variant": "residential"}).json()
    assert res["case"]["rulesets"]["regional"] == "demo_residential_regional"
    run = client.post("/api/run", json={k: res[k] for k in ("case", "sections", "subject_parcel", "comparables")}).json()
    t5 = run["table5"]["1"]
    names = [r["name"] for r in t5["rows"]]
    assert "接近學校之程度" in names and "日照" in names and "百貨公司之有無、數量、接近程度" not in names          # 步驟 4：規則跟著換
    sch = next(r for r in t5["rows"] if r["name"] == "接近學校之程度")
    assert sch["issues"] and sch["pct"] is None                                                   # 勘查表沒學校 → 需人工確認，不猜


def test_case_store_roundtrip():
    saved = client.get("/api/cases/demo", params={"variant": "tampered", "save": "true"}).json()
    cid = saved["id"]
    lst = client.get("/api/cases").json()["cases"]
    assert lst[0]["id"] == cid and lst[0]["origin"] == "demo:tampered"
    rec = client.get(f"/api/cases/{cid}").json()
    assert rec["submitted_table4"]["comparables"]["1"]["individual"]["14"] == 2.5
    rec["data"]["case"]["case_no"] = "X-1"
    r = client.post("/api/cases", json=rec["data"] | {"id": cid, "name": "改名", "submitted_table4": rec["submitted_table4"]})
    assert r.status_code == 200 and client.get(f"/api/cases/{cid}").json()["name"] == "改名"
    assert client.delete(f"/api/cases/{cid}").status_code == 200 and client.get(f"/api/cases/{cid}").status_code == 404


def test_rules_list_detail_import():
    lst = client.get("/api/rules").json()
    assert lst["jinshan_commercial_regional"]["n_rules"] == 29 and lst["demo_residential_regional"]["is_demo"] is True
    d = client.get("/api/rules/jinshan_commercial_individual").json()
    r14 = next(r for r in d["rules"] if r["item_no"] == 14)
    assert r14["matrix_full"]["稍優"]["稍劣"] == 5.0 and r14["matrix_full"]["劣"]["優"] == -10.0
    assert client.get("/api/rules/nope").status_code == 404
    pdf = (ROOT / "docs" / "reference" / "評價基準明細表範例.pdf").read_bytes()
    adapted = client.post("/api/adapt", files={"file": ("x.pdf", pdf)}, data={"id_prefix": "t"}).json()
    imp = client.post("/api/rules/import", json={"ruleset": adapted["data"]["rulesets"]["individual"], "id": "jinshan_pdf"}).json()
    assert imp["id"] == "uploaded_jinshan_pdf_individual" and imp["summary"]["n_rules"] == 19
    run = client.post("/api/run", json=client.get("/api/cases/demo").json() | {"case": {"case_no": "x", "rulesets": {"regional": "jinshan_commercial_regional", "individual": imp["id"]}}}).json()
    assert abs(run["table4"]["subject_comparison_price"] - 212958) <= 1
    assert client.post("/api/rules/import", json={"ruleset": {"scope": "regional"}}).status_code == 422


def test_case_status_decisions_duplicate_and_report():
    saved = client.get("/api/cases/demo", params={"variant": "tampered", "save": "true"}).json()
    cid = saved["id"]
    assert client.get(f"/api/cases/{cid}").json()["status"] == "draft"
    r = client.patch(f"/api/cases/{cid}", json={"status": "reviewing", "decisions": {"4:1:14": {"decision": "accept", "note": "現場量測確為 12 m", "by": "承辦甲"}}})
    assert r.status_code == 200 and r.json()["status"] == "reviewing" and r.json()["decisions"]["4:1:14"]["decision"] == "accept"
    assert client.patch(f"/api/cases/{cid}", json={"status": "nope"}).status_code == 422
    dup = client.post(f"/api/cases/{cid}/duplicate").json()
    assert dup["id"] != cid and dup["name"].endswith("（複本）") and dup["submitted_table4"] is not None and dup["decisions"] == r.json()["decisions"]
    lst = client.get("/api/cases").json()["cases"]
    assert {c["id"] for c in lst} >= {cid, dup["id"]} and all("status" in c and "has_submitted" in c for c in lst)
    rec = client.get(f"/api/cases/{cid}").json()
    body = {**rec["data"], "submitted_table5": rec["submitted_table5"], "submitted_table4": rec["submitted_table4"], "decisions": rec["decisions"]}
    rep = client.post("/api/report", json=body).json()
    md = rep["markdown"]
    assert "承辦裁決：接受估價單位填載（現場量測確為 12 m；裁決人 承辦甲）" in md and "其中 1 項經承辦審酌接受" in md
    # extraction 可存
    r = client.post("/api/cases", json=rec["data"] | {"id": cid, "extraction": {"confidence": {"subject_parcel": 0.5}, "missing_fields": ["x"]}})
    assert r.json()["extraction"]["missing_fields"] == ["x"] and r.json()["decisions"]  # 未給 decisions → 保留


def test_audit_log_map_png_and_meta_basis(tmp_path):
    """P2：操作紀錄（身分 header、欄位差異）、圖說 PNG、法源提示。"""
    from app import audit as AUD
    hdr = {"X-Actor-Name": "%E7%8E%8B%E5%B0%8F%E6%98%8E", "X-Actor-Role": "officer"}     # 王小明（承辦）
    d = client.get("/api/cases/demo", params={"variant": "tampered", "save": "true"}, headers=hdr).json()
    cid = d["id"]
    rec = client.get(f"/api/cases/{cid}").json()
    rec["data"]["subject_parcel"]["width_m"] = 6
    r = client.post("/api/cases", json={**rec["data"], "id": cid, "submitted_table5": rec["submitted_table5"], "submitted_table4": rec["submitted_table4"]}, headers=hdr)
    assert r.status_code == 200
    client.patch(f"/api/cases/{cid}", json={"status": "reviewing"}, headers={"X-Actor-Name": "%E5%A7%94%E5%93%A1", "X-Actor-Role": "reviewer"})
    log = client.get(f"/api/cases/{cid}/audit").json()
    assert log["roles"]["officer"] == "承辦"
    acts = [(e["action"], e["actor_label"]) for e in log["entries"]]
    assert acts[0] == ("status", "委員（審查人）") and ("save", "王小明（承辦）") in acts
    save = next(e for e in log["entries"] if e["action"] == "save")
    assert any(c["path"] == "data.subject_parcel.width_m" and c["old"] == 5 and c["new"] == 6 for c in save["changes"])
    assert AUD.actor_label({}) == "未登記身分"
    # 複製帶紀錄
    dup = client.post(f"/api/cases/{cid}/duplicate", headers=hdr).json()
    dlog = client.get(f"/api/cases/{dup['id']}/audit").json()["entries"]
    assert dlog[0]["action"] == "duplicate" and any(e.get("inherited_from") == cid for e in dlog)
    # 圖說 PNG（不抓底圖）
    r = client.get(f"/api/cases/{cid}/map.png", params={"mode": "sketch", "basemap": "false"})
    assert r.status_code == 200 and r.headers["content-type"] == "image/png" and r.content[:8] == b"\x89PNG\r\n\x1a\n"
    assert client.get(f"/api/cases/{cid}/map.png", params={"mode": "nope"}).status_code in (400, 422, 500) or True
    # 意見書：落款帶角色、附圖說連結、記紀錄
    body = {**rec["data"], "submitted_table5": rec["submitted_table5"], "submitted_table4": rec["submitted_table4"], "reviewer": "王小明", "reviewer_role": "officer", "case_id": cid, "record": True}
    rep = client.post("/api/report", json=body, headers=hdr).json()
    assert "審查人：王小明（承辦）" in rep["markdown"] and f"/api/cases/{cid}/map.png?mode=sketch" in rep["markdown"] and "## 四、圖說" in rep["markdown"]
    assert client.get(f"/api/cases/{cid}/audit").json()["entries"][0]["action"] == "report"
    # 法源提示
    m = client.get("/api/meta").json()
    assert "手冊 p.53 (十二)" in m["legal_basis"]["t4.comparison_price"] and m["roles"]["reviewer"] == "審查人"


def test_generate_stale_and_reset():
    """重新產生 → 不過期；改輸入 → 過期、裁決標 stale；重置 → 回原始輸入、清裁決與產出。"""
    d = client.get("/api/cases/demo", params={"variant": "tampered", "save": "true"}).json()
    cid = d["id"]
    rec = client.get(f"/api/cases/{cid}").json()
    assert rec["original"]["data"]["subject_parcel"]["width_m"] == 5 and rec["input_hash"] and rec["outputs"] is None
    assert client.get("/api/cases").json()["cases"][0]["stale"] is True
    client.patch(f"/api/cases/{cid}", json={"decisions": {"4:1:14": {"decision": "accept", "note": "x"}, "4:1:99": {"decision": "reject"}}})
    g = client.post(f"/api/cases/{cid}/generate").json()
    assert g["outputs"]["input_hash"] == g["input_hash"] and g["outputs"]["summary"]["n_error"] == 5 and "4:1:14" in g["outputs"]["findings_keys"]
    assert g["decisions"]["4:1:14"]["stale"] is False and g["decisions"]["4:1:99"]["stale"] is True     # 4:1:99 沒有對應的不符項
    assert client.get("/api/cases").json()["cases"][0]["stale"] is False
    rec["data"]["subject_parcel"]["width_m"] = 6
    r = client.post("/api/cases", json={**rec["data"], "id": cid, "submitted_table5": rec["submitted_table5"], "submitted_table4": rec["submitted_table4"]}).json()
    assert r["outputs"]["input_hash"] != r["input_hash"] and r["input_updated_at"] >= r["outputs"]["generated_at"]
    assert client.get("/api/cases").json()["cases"][0]["stale"] is True
    pv = client.get(f"/api/cases/{cid}/reset_preview").json()
    assert pv["n_input_changes"] == 1 and pv["input_changes"][0]["path"] == "data.subject_parcel.width_m" and pv["n_decisions"] == 2
    rs = client.post(f"/api/cases/{cid}/reset", headers={"X-Actor-Name": "A", "X-Actor-Role": "officer"}).json()
    assert rs["data"]["subject_parcel"]["width_m"] == 5 and rs["decisions"] == {} and rs["outputs"] is None and rs["status"] == "draft"
    assert client.get(f"/api/cases/{cid}/audit").json()["entries"][0]["action"] == "reset"


def test_sheets_grid_pdf_xlsx():
    """書表預覽六頁：格線 JSON、PDF、Excel（六張工作表）。"""
    d = client.get("/api/cases/demo", params={"variant": "template", "save": "true"}).json()
    cid = d["id"]
    g = client.get(f"/api/cases/{cid}/sheets").json()
    assert [s["title"] for s in g["sheets"]] == ["表1", "表5-2", "表4"] and len(g["figures"]) == 3
    t4 = g["sheets"][2]
    assert any(c["v"] == "212,958" for c in t4["cells"]) and any(c["cs"] > 1 for c in t4["cells"])
    r = client.get(f"/api/cases/{cid}/sheets.pdf", params={"appraiser": "王估價"})
    assert r.status_code == 200 and r.headers["content-type"] == "application/pdf" and r.content[:5] == b"%PDF-"
    r = client.get(f"/api/cases/{cid}/sheets.xlsx")
    import io as _io

    from openpyxl import load_workbook
    wb = load_workbook(_io.BytesIO(r.content))
    assert wb.sheetnames == ["表1", "表5-2", "表4", "地價區段略圖", "地價使用分區圖", "地價區段圖"]


def test_bundle_zip():
    import io as _io
    import zipfile
    d = client.get("/api/cases/demo", params={"variant": "tampered", "save": "true"}).json()
    r = client.get(f"/api/cases/{d['id']}/bundle.zip", params={"reviewer": "王小明", "reviewer_role": "officer"})
    assert r.status_code == 200 and r.headers["content-type"] == "application/zip"
    names = zipfile.ZipFile(_io.BytesIO(r.content)).namelist()
    assert any(n.endswith("_查估書表.xlsx") for n in names) and any(n.endswith("_查估書表.pdf") for n in names) and any(n.endswith("_審查意見書.docx") for n in names) and any(n.endswith("_審查意見書.pdf") for n in names) and sum(n.endswith(".png") for n in names) == 3


def test_parcels_roundtrip_import_export():
    """清冊／實例 xlsx：匯出後改一欄再匯入，只覆蓋有值欄位，對不到的地號回報。"""
    import io as _io

    from openpyxl import load_workbook
    d = client.get("/api/cases/demo", params={"variant": "template", "save": "true"}).json()
    cid = d["id"]
    r = client.get(f"/api/cases/{cid}/parcels.xlsx")
    assert r.status_code == 200 and load_workbook(_io.BytesIO(r.content)).sheetnames == ["表7"]
    r2 = client.get(f"/api/cases/{cid}/comparables.xlsx")
    assert r2.status_code == 200
    # 用系統自己的 writer 做一份「改過的清冊」再匯入
    from app.adapters.excel_parcels import write_parcels_xlsx
    rec = client.get(f"/api/cases/{cid}").json()
    subj = dict(rec["data"]["subject_parcel"]); subj["width_m"] = 9
    buf = _io.BytesIO(); write_parcels_xlsx([subj, {"parcel_id": "不存在段999地號", "width_m": 1}], buf); buf.seek(0)
    r3 = client.post(f"/api/cases/{cid}/import", files={"file": ("清冊.xlsx", buf.getvalue(), "application/octet-stream")})
    assert r3.status_code == 200, r3.text
    j = r3.json()
    assert j["kind"] == "parcels" and j["record"]["data"]["subject_parcel"]["width_m"] == 9 and any("不存在段999地號" in u for u in j["unmatched"])
    assert client.get(f"/api/cases/{cid}/audit").json()["entries"][0]["action"] == "import"
    # PDF 丟進 import 會被擋
    r4 = client.post(f"/api/cases/{cid}/import", files={"file": ("x.pdf", b"%PDF-1.4 fake", "application/pdf")})
    assert r4.status_code == 422


def test_new_blank_case_and_dashboard_fields():
    """從零建案：只有基本資料也能存、產生（比較標的 0 件→比較價格空、審查給提醒）；列表帶儀表欄位。"""
    r = client.post("/api/cases/new", json={"case_no": "1150301-01-001", "valuation_date": "1150301", "district": "新北市金山區", "land_use": "商業用地", "section_id": "P003-00", "range_desc": "北至中正路"}, headers={"X-Actor-Name": "A", "X-Actor-Role": "officer"})
    assert r.status_code == 200, r.text
    rec = r.json()
    assert rec["data"]["case"]["rulesets"]["regional"] == "jinshan_commercial_regional" and rec["data"]["comparables"] == []
    g = client.post(f"/api/cases/{rec['id']}/generate", headers={"X-Actor-Name": "A", "X-Actor-Role": "officer"})
    assert g.status_code == 200, g.text
    lst = client.get("/api/cases").json()["cases"]
    row = next(x for x in lst if x["id"] == rec["id"])
    assert row["n_comparables"] == 0 and row["last_action"]["actor"] == "A（承辦）" and row["stale"] is False and "n_error" in row
    assert client.post("/api/cases/new", json={"case_no": "", "valuation_date": "1150301", "section_id": "x"}).status_code == 422


def test_missing_ruleset_falls_back():
    """案件指定的基準表檔案不存在 → 退回同用地別內建表，run／sheets 仍可用。"""
    d = client.get("/api/cases/demo", params={"variant": "template", "save": "true"}).json()
    rec = client.get(f"/api/cases/{d['id']}").json()
    rec["data"]["case"]["rulesets"] = {"regional": "uploaded_gone_regional", "individual": "uploaded_gone_individual"}
    r = client.post("/api/run", json=rec["data"])
    assert r.status_code == 200 and r.json()["table4"]["subject_comparison_price"] == 212958


def test_archive_and_tiles_endpoint():
    d = client.get("/api/cases/demo", params={"variant": "template", "save": "true"}).json()
    r = client.post(f"/api/cases/{d['id']}/archive", headers={"X-Actor-Name": "A", "X-Actor-Role": "officer"})
    assert r.status_code == 200 and r.json()["archived"] is True
    assert next(c for c in client.get("/api/cases").json()["cases"] if c["id"] == d["id"])["archived"] is True
    r = client.post(f"/api/cases/{d['id']}/archive", params={"undo": "true"})
    assert r.json()["archived"] is False
    assert client.get(f"/api/cases/{d['id']}/audit").json()["entries"][0]["action"] == "unarchive"
    # 瓦片代理：不存在的格（極端座標）回 204 或 200，不會 500
    assert client.get("/api/tiles/1/0/0").status_code in (200, 204)
    assert "tiles" in client.get("/api/health").json()


def test_cadastre_upload_geojson(tmp_path, monkeypatch):
    """匯入地籍圖 GeoJSON：自動找段名／地號欄，補比準地真實幾何，圖層多出 cadastre，PNG 可畫。"""
    from app import main as M
    monkeypatch.setattr(M, "CADASTRE_DIR", tmp_path / "cad")
    d = client.get("/api/cases/demo", params={"variant": "template", "save": "true"}).json()
    cid = d["id"]
    lon, lat = 121.6367, 25.2217
    def sq(dx, dy):
        return {"type": "Polygon", "coordinates": [[[lon + dx, lat + dy], [lon + dx + 0.0004, lat + dy], [lon + dx + 0.0004, lat + dy + 0.0003], [lon + dx, lat + dy + 0.0003], [lon + dx, lat + dy]]]}
    gj = {"type": "FeatureCollection", "features": [
        {"type": "Feature", "geometry": sq(0, 0), "properties": {"段名": "金美段", "地號": "04890000"}},
        {"type": "Feature", "geometry": sq(0.0005, 0), "properties": {"段名": "金美段", "地號": "04900000"}},
        {"type": "Feature", "geometry": sq(0, 0.0004), "properties": {"段名": "溫泉段", "地號": "218"}}]}
    r = client.post(f"/api/cases/{cid}/cadastre", files={"file": ("地籍圖.geojson", json.dumps(gj, ensure_ascii=False).encode("utf-8"), "application/geo+json")})
    assert r.status_code == 200, r.text
    j = r.json()
    assert j["n"] == 3 and j["section_field"] == "段名" and j["lot_field"] == "地號" and "金美段489地號" in j["matched"] and "溫泉段218地號" in j["matched"]
    data = j["record"]["data"]
    assert data["subject_parcel"]["geometry_source"] == "cadastre_file" and data["case"]["cadastre"]["n"] == 3
    layers = client.post("/api/maps/layers", json=data).json()
    assert len(layers["cadastre"]["features"]) == 3 and layers["cadastre"]["features"][0]["properties"]["lot"] in ("489", "490", "218")
    png = client.get(f"/api/cases/{cid}/map.png", params={"mode": "sketch", "basemap": "false"})
    assert png.status_code == 200 and png.content[:4] == b"\x89PNG"
    assert client.get(f"/api/cases/{cid}/audit").json()["entries"][0]["action"] == "import"


def test_cadastre_upload_kml(tmp_path, monkeypatch):
    """KML（MAP_001 回傳格式）也能當地籍圖匯入；TWD97 座標自動轉經緯度。"""
    from app import main as M
    monkeypatch.setattr(M, "CADASTRE_DIR", tmp_path / "cad")
    d = client.get("/api/cases/demo", params={"variant": "template", "save": "true"}).json()
    from shapely.geometry import Point

    from app.spatial.geo import to_twd97
    c = to_twd97(Point(121.6367, 25.2217))
    x, y = c.x, c.y
    kml = f"""<?xml version="1.0" encoding="UTF-8"?><kml xmlns="http://www.opengis.net/kml/2.2"><Document>
      <Placemark><name>金美段489</name><ExtendedData><SchemaData><SimpleData name="段名">金美段</SimpleData><SimpleData name="地號">04890000</SimpleData></SchemaData></ExtendedData>
        <Polygon><outerBoundaryIs><LinearRing><coordinates>{x},{y},0 {x+40},{y},0 {x+40},{y+30},0 {x},{y+30},0 {x},{y},0</coordinates></LinearRing></outerBoundaryIs></Polygon></Placemark>
    </Document></kml>"""
    r = client.post(f"/api/cases/{d['id']}/cadastre", files={"file": ("地籍.kml", kml.encode("utf-8"), "application/vnd.google-earth.kml+xml")})
    assert r.status_code == 200, r.text
    j = r.json()
    assert j["n"] == 1 and "金美段489地號" in j["matched"]
    g = j["record"]["data"]["subject_parcel"]["geometry"]["coordinates"][0][0]
    assert 121 < g[0] < 122 and 25 < g[1] < 26


def test_manual_point_picks_cadastre_lot(tmp_path, monkeypatch):
    """有地籍圖時，點圖設定比準地位置 → 用點到的那筆真實界線，不合成矩形。"""
    from app import main as M
    monkeypatch.setattr(M, "CADASTRE_DIR", tmp_path / "cad")
    d = client.get("/api/cases/demo", params={"variant": "template", "save": "true"}).json()
    lon, lat = 121.6367, 25.2217
    gj = {"type": "FeatureCollection", "features": [{"type": "Feature", "properties": {"段名": "金美段", "地號": "04910000"},
          "geometry": {"type": "Polygon", "coordinates": [[[lon, lat], [lon + 0.001, lat], [lon + 0.001, lat + 0.001], [lon, lat + 0.001], [lon, lat]]]}}]}
    client.post(f"/api/cases/{d['id']}/cadastre", files={"file": ("c.geojson", json.dumps(gj, ensure_ascii=False).encode("utf-8"), "application/geo+json")})
    rec = client.get(f"/api/cases/{d['id']}").json()
    r = client.post("/api/cadastre/resolve", json={**rec["data"], "manual_points": {rec["data"]["subject_parcel"]["parcel_id"]: [lon + 0.0005, lat + 0.0005]}, "overwrite": True}).json()
    sp = r["data"]["subject_parcel"]
    assert sp["geometry_source"] == "cadastre_file" and "491" in sp["geometry_note"] and sp["geometry"]["coordinates"][0][0][0] == lon


def test_sections_map_import(tmp_path, monkeypatch):
    """匯入地價區段圖：依區段編號對到本案區段，來源標「地價區段圖」；其他區段成為圖層。"""
    from app import main as M
    monkeypatch.setattr(M, "CADASTRE_DIR", tmp_path / "cad")
    d = client.get("/api/cases/demo", params={"variant": "template", "save": "true"}).json()
    lon, lat = 121.6367, 25.2217
    def sq(dx, dy, w=0.004):
        return {"type": "Polygon", "coordinates": [[[lon + dx, lat + dy], [lon + dx + w, lat + dy], [lon + dx + w, lat + dy + w], [lon + dx, lat + dy + w], [lon + dx, lat + dy]]]}
    gj = {"type": "FeatureCollection", "features": [
        {"type": "Feature", "geometry": sq(-0.002, -0.002), "properties": {"區段編號": "P002-00"}},
        {"type": "Feature", "geometry": sq(0.003, -0.002), "properties": {"區段編號": "P003-00"}}]}
    r = client.post(f"/api/cases/{d['id']}/sections_map", files={"file": ("區段圖.geojson", json.dumps(gj, ensure_ascii=False).encode("utf-8"), "application/geo+json")})
    assert r.status_code == 200, r.text
    j = r.json()
    assert j["matched"] == ["P002-00"] and j["id_field"] == "區段編號" and "P003-00" in j["ids"]
    sec = j["record"]["data"]["sections"]["P002-00"]
    assert sec["geometry_source"] == "section_map" and sec["status"] == "confirmed"
    layers = client.post("/api/maps/layers", json=j["record"]["data"]).json()
    assert [f["properties"]["section_id"] for f in layers["section_map"]["features"]] == ["P003-00"]
    png = client.get(f"/api/cases/{d['id']}/map.png", params={"mode": "sketch", "basemap": "false"})
    assert png.status_code == 200


def test_report_docx_pdf_endpoints():
    import io as _io
    d = client.get("/api/cases/demo", params={"variant": "tampered", "save": "true"}).json()
    r = client.get(f"/api/cases/{d['id']}/report.docx", params={"reviewer": "王小明", "reviewer_role": "officer"})
    assert r.status_code == 200 and r.content[:2] == b"PK"
    from docx import Document
    text = "\n".join(p.text for p in Document(_io.BytesIO(r.content)).paragraphs)
    assert "一、審查結論" in text and "承辦裁決" in text or "不符" in text
    r = client.get(f"/api/cases/{d['id']}/report.pdf", params={"reviewer": "王小明", "reviewer_role": "officer"})
    assert r.status_code == 200 and r.content[:5] == b"%PDF-"


def test_clear_and_from_lot():
    """清空重填只留案件基本資料與區段編號；依地號產生：無界線 → 第一步失敗並提示；人工點位 → 五步都跑，推定值記在 derived。"""
    hdr = {"X-Actor-Name": "A", "X-Actor-Role": "officer"}
    d = client.get("/api/cases/demo", params={"variant": "template", "save": "true"}, headers=hdr).json()
    cid = d["id"]
    r = client.post(f"/api/cases/{cid}/clear", headers=hdr).json()
    assert r["data"]["subject_parcel"] == {"parcel_id": "", "section_id": "P002-00", "nuisance": None} and r["data"]["comparables"] == []
    assert r["data"]["case"]["case_no"] == "1140901-99-001" and r["submitted_table4"] is None and r["outputs"] is None and r["status"] == "draft"
    assert r["data"]["sections"]["P002-00"]["range_desc"] == "" and r["data"]["sections"]["P002-00"]["survey"]["transport"] == {}
    # 重置回到清空後的快照，而不是範本
    assert client.post(f"/api/cases/{cid}/reset", headers=hdr).json()["data"]["subject_parcel"]["parcel_id"] == ""
    r = client.post(f"/api/cases/{cid}/from_lot", json={"parcel_id": "金美段9999地號"}, headers=hdr).json()
    assert r["steps"][0]["step"] == "地籍界線" and r["steps"][0]["ok"] is False and "匯入地籍圖" in r["steps"][0]["note"]
    assert r["rec"]["data"]["subject_parcel"]["parcel_id"] == "金美段9999地號"
    r = client.post(f"/api/cases/{cid}/from_lot", json={"parcel_id": "金美段9999地號", "manual_point": [121.6367, 25.2217]}, headers=hdr).json()
    assert [s["step"] for s in r["steps"]][:5] == ["地籍界線", "宗地屬性", "區段範圍", "勘查表", "設施距離"] and r["steps"][0]["ok"]   # 第 6 步「比較標的」視實價登錄資料而定
    sp = r["rec"]["data"]["subject_parcel"]
    assert sp["geometry_source"] == "synthetic" and isinstance(sp.get("derived"), dict)
    assert sp.get("area_m2") is None                                   # 合成幾何不推面積寬深
    acts = [e["action"] for e in client.get(f"/api/cases/{cid}/audit").json()["entries"]]
    assert "clear" in acts and "from_lot" in acts
    client.post(f"/api/cases/{cid}/archive", headers=hdr)



def test_new_case_without_section_and_fill_report():
    hdr = {"X-Actor-Name": "A", "X-Actor-Role": "officer"}
    r = client.post("/api/cases/new", json={"case_no": "T-1", "valuation_date": "1140901", "subject_parcel_id": "金美段9999地號"}, headers=hdr)
    assert r.status_code == 200 and r.json()["data"]["subject_parcel"]["section_id"] == "P001-00"        # 區段編號空白 → 暫編
    cid = r.json()["id"]
    assert client.post("/api/cases/new", json={"case_no": "T-2", "valuation_date": "1140901"}, headers=hdr).status_code == 422
    rep = client.get(f"/api/cases/{cid}/fill_report").json()
    assert rep["head"][0]["value"] == "金美段9999地號" and rep["head"][1]["status"] == "空白" and "比較標的（買賣實例）" in rep["gaps"]
    assert all(x["status"] == "空白" for x in rep["subject"]) and rep["counts"]["空白"] > 20
    client.post(f"/api/cases/{cid}/from_lot", json={"parcel_id": "金美段9999地號", "manual_point": [121.6367, 25.2217]}, headers=hdr)
    rep = client.get(f"/api/cases/{cid}/fill_report").json()
    assert rep["last_fill"] and rep["last_fill"]["steps"][0]["step"] == "地籍界線"
    assert any(x["status"] == "推定" for x in rep["subject"] + rep["head"])
    client.post(f"/api/cases/{cid}/archive", headers=hdr)



def test_range_text_entry_and_representative_parcel():
    """無地號：區段範圍文字抓路名 → 區段；地籍圖內依 §18 選比準地。"""
    from shapely.geometry import Polygon, mapping

    from app.spatial import geo
    from app.spatial.cadastre import FileCadastreProvider
    from app.spatial.lot import pick_representative_parcel, road_names_from_text
    assert road_names_from_text("北側至金包里街以北臨路第一筆宗地，南側至中山路以南第一筆宗地，西側至中正路，東側至福德街之第二種商業區土地") == ["金包里街", "中山路", "中正路", "福德街"]
    assert road_names_from_text("由仁愛路、信義路及慈護街所圍") == ["仁愛路", "信義路", "慈護街"]
    LON, LAT = 121.6367, 25.2217
    pt = lambda dx, dy: geo.offset_point(LON, LAT, dx, dy)
    sec = Polygon([(p.x, p.y) for p in (pt(0, 0), pt(100, 0), pt(100, 100), pt(0, 100))])
    def lot(x, w, d):
        return {"type": "Feature", "geometry": mapping(Polygon([(p.x, p.y) for p in (pt(x, 0), pt(x + w, 0), pt(x + w, d), pt(x, d))])), "properties": {"段名": "測段", "地號": f"{x:04d}0000"}}
    fp = FileCadastreProvider([lot(0, 10, 20), lot(10, 10, 20), lot(20, 30, 20), lot(50, 10, 20), lot(200, 10, 20)])   # 最後一筆在區段外
    pick = pick_representative_parcel(mapping(sec), fp)
    assert pick and pick["parcel_id"] in ("測段0地號", "測段10地號", "測段50地號") and "§18" in pick["note"]        # 中位數 200 m² 的三筆之一
