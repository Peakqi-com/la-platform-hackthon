"""
審查意見書 → Word（.docx）與 PDF。內容與 opinion.to_markdown 同一份資料：標題、案件資訊、一、審查結論；二、核算摘要；三、逐項意見；四、圖說（附圖）。
Word 用 python-docx（標楷體，沒裝字型只影響外觀）；PDF 用 reportlab（內建 TTC 中文字型，見 output/grid.pdf_font）。
"""
from __future__ import annotations

import io
from typing import Any

from .opinion import CHECKLIST_TITLES, OpinionReport

FONT_ZH = "標楷體"


def _summary_lines(r: OpinionReport) -> list[str]:
    s = r.computed_summary or {}
    out = []
    for c in s.get("comparables", []):
        tp = f"{c['trial_price']:,}" if c.get("trial_price") is not None else "—"
        out.append(f"比較標的{c['comp_no']}：區域因素調整 {c['regional_adjustment_pct']}%、個別因素合計 {c['individual_total_pct']}%、試算價格 {tp} 元/m²、權重 {c['weight_pct']}%")
    if s.get("subject_comparison_price") is not None:
        out.append(f"比準地比較價格 {int(s['subject_comparison_price']):,} 元/m²（四捨五入至個位，作業手冊 p.53（十二））")
    if s.get("subject_land_price") is not None:
        out.append(f"比準地地價（僅比較法，依查估辦法 §21 尾數進位）{int(s['subject_land_price']):,} 元/m²")
    return out


def _meta_lines(r: OpinionReport) -> list[tuple[str, str]]:
    return [("案號", r.case_no), ("比準地", r.subject_parcel_id or "—"), ("估價基準日", r.valuation_date), ("審查日期", r.review_date), ("審查人", r.reviewer or "—")]


def _counts_line(r: OpinionReport) -> str:
    return f"不符 {r.counts.get('error', 0)} 項、需確認 {r.counts.get('warn', 0)} 項、備註 {r.counts.get('info', 0)} 項。"


FOOT = "本意見書由規則引擎依作業手冊與評價基準明細表核算產生，各條意見標籤 [F-xxx] 對應系統審查項目；數字未經人工修改。"


# ------------------------------------------------------------------ Word


def report_to_docx(r: OpinionReport, figures: list[tuple[str, str, bytes]] | None = None, use_polished: bool = True) -> bytes:
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml.ns import qn
    from docx.shared import Cm, Pt

    doc = Document()
    st = doc.styles["Normal"]
    st.font.name = FONT_ZH
    st.font.size = Pt(12)
    st.element.rPr.rFonts.set(qn("w:eastAsia"), FONT_ZH)
    for sec in doc.sections:
        sec.left_margin = sec.right_margin = Cm(2.2)
        sec.top_margin = sec.bottom_margin = Cm(2)

    def para(text: str, *, size: int = 12, bold: bool = False, align=None, space_after: int = 4):
        p = doc.add_paragraph()
        run = p.add_run(text)
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.name = FONT_ZH
        run._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_ZH)
        if align is not None:
            p.alignment = align
        p.paragraph_format.space_after = Pt(space_after)
        return p

    para("土地徵收補償市價查估案件審查意見書", size=18, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=10)
    tbl = doc.add_table(rows=0, cols=2)
    tbl.style = "Table Grid"
    for k, v in _meta_lines(r):
        row = tbl.add_row().cells
        row[0].text, row[1].text = k, str(v)
        for cell in row:
            for p in cell.paragraphs:
                for run in p.runs:
                    run.font.name = FONT_ZH
                    run._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_ZH)
                    run.font.size = Pt(11)
        row[0].width = Cm(3.5)
    doc.add_paragraph()

    para("一、審查結論", size=14, bold=True, space_after=6)
    para(r.conclusion)
    para(_counts_line(r), space_after=10)

    lines = _summary_lines(r)
    if lines:
        para("二、核算摘要", size=14, bold=True, space_after=6)
        for ln in lines:
            para("• " + ln)
        doc.add_paragraph()

    para("三、逐項意見", size=14, bold=True, space_after=6)
    if not r.sections:
        para("（無）")
    for key, items in r.sections.items():
        para(CHECKLIST_TITLES.get(key, key), size=12, bold=True, space_after=2)
        for it in items:
            para("• " + (it.polished if use_polished and it.polished else it.sentence))
        doc.add_paragraph()

    if figures:
        para("四、圖說", size=14, bold=True, space_after=6)
        for _mode, title, png in figures:
            para(title, bold=True, space_after=2)
            doc.add_picture(io.BytesIO(png), width=Cm(16))
            doc.add_paragraph()
        para("（區段略圖、使用分區圖、地價區段圖由系統依區段範圍、宗地位置、設施量測與使用分區圖資產生；底圖 © 國土測繪中心。）", size=10)

    para(FOOT, size=9, space_after=0)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


# ------------------------------------------------------------------ PDF


def report_to_pdf(r: OpinionReport, figures: list[tuple[str, str, bytes]] | None = None, use_polished: bool = True) -> bytes:
    from reportlab.lib import colors
    from reportlab.lib.enums import TA_CENTER
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import mm
    from reportlab.lib.utils import ImageReader
    from reportlab.platypus import Image, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    from app.output.grid import pdf_font
    font = pdf_font()
    title_st = ParagraphStyle("t", fontName=font, fontSize=18, leading=26, alignment=TA_CENTER, spaceAfter=10)
    h_st = ParagraphStyle("h", fontName=font, fontSize=14, leading=20, spaceBefore=10, spaceAfter=6)
    h2_st = ParagraphStyle("h2", fontName=font, fontSize=12, leading=18, spaceBefore=6, spaceAfter=2)
    body = ParagraphStyle("b", fontName=font, fontSize=11, leading=17, spaceAfter=4)
    bullet = ParagraphStyle("bl", parent=body, leftIndent=12, firstLineIndent=-12)
    small = ParagraphStyle("s", fontName=font, fontSize=9, leading=13, textColor=colors.HexColor("#555555"))

    def esc(t: Any) -> str:
        return str(t).replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")

    story: list[Any] = [Paragraph("土地徵收補償市價查估案件審查意見書", title_st)]
    meta = Table([[k, str(v)] for k, v in _meta_lines(r)], colWidths=[35 * mm, 125 * mm])
    meta.setStyle(TableStyle([("FONTNAME", (0, 0), (-1, -1), font), ("FONTSIZE", (0, 0), (-1, -1), 11), ("GRID", (0, 0), (-1, -1), 0.5, colors.black),
                              ("BACKGROUND", (0, 0), (0, -1), colors.HexColor("#ffedd5")), ("VALIGN", (0, 0), (-1, -1), "MIDDLE"), ("TOPPADDING", (0, 0), (-1, -1), 3), ("BOTTOMPADDING", (0, 0), (-1, -1), 3)]))
    story += [meta, Spacer(1, 8), Paragraph("一、審查結論", h_st), Paragraph(esc(r.conclusion), body), Paragraph(esc(_counts_line(r)), body)]
    lines = _summary_lines(r)
    if lines:
        story.append(Paragraph("二、核算摘要", h_st))
        story += [Paragraph("• " + esc(ln), bullet) for ln in lines]
    story.append(Paragraph("三、逐項意見", h_st))
    if not r.sections:
        story.append(Paragraph("（無）", body))
    for key, items in r.sections.items():
        story.append(Paragraph(esc(CHECKLIST_TITLES.get(key, key)), h2_st))
        story += [Paragraph("• " + esc(it.polished if use_polished and it.polished else it.sentence), bullet) for it in items]
    if figures:
        story.append(Paragraph("四、圖說", h_st))
        for _mode, title, png in figures:
            iw, ih = ImageReader(io.BytesIO(png)).getSize()
            w = 165 * mm
            story += [Paragraph(esc(title), h2_st), Image(io.BytesIO(png), width=w, height=w * ih / iw), Spacer(1, 6)]
        story.append(Paragraph("（區段略圖、使用分區圖、地價區段圖由系統依區段範圍、宗地位置、設施量測與使用分區圖資產生；底圖 © 國土測繪中心。）", small))
    story += [Spacer(1, 10), Paragraph(esc(FOOT), small)]
    buf = io.BytesIO()
    SimpleDocTemplate(buf, pagesize=A4, leftMargin=22 * mm, rightMargin=22 * mm, topMargin=18 * mm, bottomMargin=18 * mm,
                      title="審查意見書", author="估價案件審查輔助系統").build(story)
    return buf.getvalue()
