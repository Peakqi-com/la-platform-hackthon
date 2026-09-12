"""
宗地個別因素清冊（表7）Excel → Parcel[]。

支援兩種版面：
  transposed  內政部官方版（作業手冊 p.85 表7）：欄位是列、宗地是欄；15~20 各佔兩列（名稱／距離）、25 佔兩列（有無／情形）
  long        一列一宗地，表頭是欄位編號或欄名（當天若拿到整理過的清單）

同一支模組也提供 write_parcels_xlsx()：依官方轉置版面把 Parcel[] 寫回 Excel（round-trip 測試、當天反向產清冊）。
"""
from __future__ import annotations

import io
import re
from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, Side
from openpyxl.utils import get_column_letter

from .common import (
    MISSING,
    AdapterResult,
    FieldSpec,
    annotate_facility_defaults,
    is_blank,
    make_facility,
    norm_label,
    norm_text,
    parse_by_kind,
    read_records,
)

SOURCE_TABLE7 = "表7 宗地個別因素清冊（需用土地人填）"

# 依表7 順序。group = 表7 左欄主要項目；item_no = 欄位編號。
PARCEL_SPECS: list[FieldSpec] = [
    FieldSpec("serial_no", ("宗地流水號", "流水號"), "text", group="0基本資料"),
    FieldSpec("district", ("鄉鎮市區",), "text", group="0基本資料"),
    FieldSpec("section_name", ("段小段名稱", "段小段", "段名", "地段"), "text", group="0基本資料"),
    FieldSpec("lot_no", ("地號",), "text", group="0基本資料"),
    FieldSpec("owner", ("土地所有權人或管理人姓名", "所有權人", "土地所有權人"), "text", group="0基本資料"),
    FieldSpec("parcel_id", ("宗地標示", "parcel_id"), "text"),
    FieldSpec("address", ("地址", "address"), "text"),
    FieldSpec("section_id", ("地價區段", "地價區段號", "區段編號", "section_id"), "text"),
    FieldSpec("area_m2", ("面積(M2)", "面積", "area_m2"), "num", 7, "1.宗地條件", required=True),
    FieldSpec("width_m", ("寬度(M)", "寬度", "width_m"), "num", 8, "1.宗地條件", required=True),
    FieldSpec("depth_m", ("深度(M)", "深度", "depth_m"), "num", 9, "1.宗地條件", required=True),
    FieldSpec("shape", ("形狀", "shape"), "text", 10, "1.宗地條件", required=True),
    FieldSpec("frontage", ("臨街情形", "frontage"), "text", 11, "1.宗地條件", required=True),
    FieldSpec("terrain", ("地勢", "terrain"), "text", 12, "1.宗地條件", required=True),
    FieldSpec("road_type", ("道路種類", "road_type"), "text", 13, "2.道路條件", required=True),
    FieldSpec("front_road_name", ("面前道路名稱", "front_road_name"), "text", 14, "2.道路條件"),
    FieldSpec("front_road_width_m", ("面前道路寬度", "面前道路寬度(M)", "front_road_width_m"), "num", 14, "2.道路條件", required=True),
    FieldSpec("school_name", ("學校名稱", "school_name"), "text", 15, "3.接近條件"),
    FieldSpec("school_distance_m", ("接近學校之程度(M)", "接近學校之程度", "school_distance_m"), "num", 15, "3.接近條件", required=True),
    FieldSpec("market_name", ("市場名稱", "market_name"), "text", 16, "3.接近條件"),
    FieldSpec("market_distance_m", ("接近市場之程度(M)", "接近市場之程度", "market_distance_m"), "num", 16, "3.接近條件", required=True),
    FieldSpec("park_name", ("公園、廣場名稱", "公園廣場名稱", "公園名稱", "park_name"), "text", 17, "3.接近條件"),
    FieldSpec("park_distance_m", ("接近公園、廣場之程度(M)", "接近公園廣場之程度", "park_distance_m"), "num", 17, "3.接近條件", required=True),
    FieldSpec("station_name", ("車站名稱", "station_name"), "text", 18, "3.接近條件"),
    FieldSpec("station_distance_m", ("接近車站之程度(M)", "接近車站之程度", "station_distance_m"), "num", 18, "3.接近條件", required=True),
    FieldSpec("commercial_district_name", ("商圈名稱", "commercial_district_name"), "text", 19, "3.接近條件"),
    FieldSpec("commercial_district_distance_m", ("接近商圈之程度(M)", "接近商圈之程度", "commercial_district_distance_m"), "num", 19, "3.接近條件", required=True),
    FieldSpec("nuisance_name", ("嫌惡設施名稱", "嫌惡設施(類型)", "嫌惡設施", "nuisance_name"), "text", 20, "4.周邊環境條件"),
    FieldSpec("nuisance_distance_m", ("接近嫌惡設施之程度(M)", "接近嫌惡設施之程度", "nuisance_distance_m"), "text", 20, "4.周邊環境條件", required=True),
    FieldSpec("street_parking", ("停車方便性", "street_parking"), "text", 21, "4.周邊環境條件", required=True),
    FieldSpec("zoning", ("使用分區或編定用地", "使用分區", "zoning"), "text", 22, "5.行政條件", required=True),
    FieldSpec("bcr_pct", ("建蔽率(%)", "建蔽率", "bcr_pct"), "num", 23, "5.行政條件", required=True),
    FieldSpec("far_pct", ("容積率(%)", "容積率", "far_pct"), "num", 24, "5.行政條件", required=True),
    FieldSpec("building_restricted", ("有無禁、限建", "有無禁限建", "有無禁限建築", "building_restricted"), "bool", 25, "5.行政條件", required=True),
    FieldSpec("restriction_note", ("禁、限建情形", "禁限建情形", "restriction_note"), "text", 25, "5.行政條件"),
    FieldSpec("other", ("其他", "other"), "text", 6, "6.其他"),
    FieldSpec("negotiation", ("協議價購程序", "協議價購", "negotiation"), "text", group="協議價購程序"),
]

_FACILITY_FIELDS = [  # (schema key, name spec key, distance spec key, 預設 type)
    ("school", "school_name", "school_distance_m", "school"),
    ("market", "market_name", "market_distance_m", "traditional_market"),
    ("park", "park_name", "park_distance_m", "park"),
    ("station", "station_name", "station_distance_m", "bus_stop"),
    ("commercial_district", "commercial_district_name", "commercial_district_distance_m", "commercial_district"),
]

_SPLIT_RE = re.compile(r"[、;；,，/\n]+")


def _split_multi(v: Any) -> list[str]:
    if v is MISSING or is_blank(v):
        return []
    return [x.strip() for x in _SPLIT_RE.split(str(v)) if x.strip()]


def assemble_parcel(rec: dict[str, Any], result: AdapterResult, path: str, source: str = SOURCE_TABLE7) -> dict:
    """平面 record（key → raw/MISSING）→ Parcel（docs/03）。缺欄位記到 result.missing_fields。"""
    g: dict[str, Any] = {s.key: parse_by_kind(s.kind, rec.get(s.key, MISSING)) for s in PARCEL_SPECS}

    def val(k: str) -> Any:
        return None if g[k] is MISSING else g[k]

    def need(k: str, schema_key: str) -> None:
        if g[k] is MISSING:
            result.missing(f"{path}.{schema_key}")

    p: dict[str, Any] = {}
    sec, lot = val("section_name"), val("lot_no")
    lot_s = None
    if lot is not None:
        lot_s = str(lot)
        if isinstance(lot, float) and lot.is_integer():
            lot_s = str(int(lot))
        if not lot_s.endswith("地號"):
            lot_s += "地號"
    p["parcel_id"] = val("parcel_id") or (f"{sec}{lot_s}" if sec and lot_s else lot_s)
    if p["parcel_id"] is None:
        result.missing(f"{path}.parcel_id")
    p["address"] = val("address") or (f"{val('district') or ''}{sec or ''}{lot_s or ''}" or None)
    if val("serial_no") is not None:
        p["serial_no"] = str(val("serial_no"))
    for k in ("district", "owner"):
        if val(k) is not None:
            p[k] = val(k)
    p["section_id"] = val("section_id")

    for k in ("area_m2", "width_m", "depth_m", "shape", "frontage", "terrain", "road_type"):
        p[k] = val(k)
        need(k, k)

    # 面前道路：表7 只有寬度；非臨街地填「無」→ None（引擎依 bands none_level 或視為缺）
    fr_name, fr_w = val("front_road_name"), val("front_road_width_m")
    if fr_name is not None or fr_w is not None:
        p["front_road"] = {"name": fr_name, "width_m": fr_w}
    else:
        p["front_road"] = None
    if g["front_road_width_m"] is MISSING:
        result.missing(f"{path}.front_road.width_m")
    elif fr_w is None and p.get("frontage") not in (None, "非臨街地"):
        result.warn(f"{path}.front_road.width_m: 填「無/-」但臨街情形不是非臨街地，請確認")

    for key, nk, dk, ftype in _FACILITY_FIELDS:
        name, dist = val(nk), val(dk)
        raw_name = rec.get(nk, MISSING)
        if name is not None and norm_text(name).startswith("無"):
            name, raw_name = None, "無"   # 「無」= 沒有該設施（明示），不是缺欄
        if name is None and dist is None:
            p[key] = None
            if g[dk] is MISSING and (raw_name is MISSING or is_blank(raw_name)):
                result.missing(f"{path}.{key}")
            # 名稱「無」= 該設施不存在 → None，不列 missing
        else:
            f = make_facility(name, dist, ftype=None, source=source)
            if f and "type" not in f:
                f["type"] = ftype
            p[key] = annotate_facility_defaults(f, scope="individual", result=result, path=f"{path}.{key}")

    names = _split_multi(rec.get("nuisance_name", MISSING))
    dists = _split_multi(rec.get("nuisance_distance_m", MISSING))
    names = [n for n in names if n not in ("無", "-", "－")]
    dists_n = [parse_by_kind("num", d) for d in dists]
    dists_n = [d for d in dists_n if d is not None and d is not MISSING]
    nuis: list[dict] = []
    for i, n in enumerate(names):
        d = dists_n[i] if i < len(dists_n) else None
        f = make_facility(n, d, source=source)
        if f:
            if "type" not in f:
                f["type"] = "nuisance"
                result.warn(f"{path}.nuisance[{i}]: 無法由名稱「{n}」判斷設施類型，請人工指定")
            nuis.append(annotate_facility_defaults(f, scope="individual", result=result, path=f"{path}.nuisance[{i}]"))
    if not names and not dists_n and g["nuisance_distance_m"] is MISSING and is_blank(rec.get("nuisance_name", MISSING)):
        result.missing(f"{path}.nuisance")
    if len(dists_n) > len(names) and names:
        result.warn(f"{path}.nuisance: 距離數量({len(dists_n)})多於名稱數量({len(names)})")
    p["nuisance"] = nuis

    for k in ("street_parking", "zoning", "bcr_pct", "far_pct", "building_restricted"):
        p[k] = val(k)
        need(k, k)
    if val("restriction_note") is not None:
        p["restriction_note"] = val("restriction_note")
    p["other"] = val("other")
    if val("negotiation") is not None:
        p["negotiation"] = val("negotiation")
    return p


def read_parcels(src: str | Path | bytes | io.BytesIO, *, sheet: str | None = None) -> AdapterResult:
    """Excel（xlsx）→ AdapterResult(kind='parcels', data={'parcels': [...], 'meta': {...}})。"""
    result = AdapterResult(kind="parcels", data={"parcels": [], "meta": {}})
    wb = load_workbook(io.BytesIO(src) if isinstance(src, (bytes, bytearray)) else src, data_only=True)
    ws = wb[sheet] if sheet else wb.active
    try:
        records, layout, unknown = read_records(ws, PARCEL_SPECS, "serial_no")
    except ValueError as e:
        result.warn(str(e))
        return result
    result.data["meta"]["layout"] = layout
    result.data["meta"]["sheet"] = ws.title
    groups = {norm_label(s.group) for s in PARCEL_SPECS if s.group}
    for lbl in unknown:
        if any(k in lbl for k in ("興辦事業", "案號", "填寫日期", "核章", "宗地個別因素清冊", "共有土地")):
            _harvest_meta(lbl, result.data["meta"])
            continue
        if norm_label(lbl) in groups:
            continue
        result.warn(f"未對應到 schema 的欄位「{lbl}」已略過")
    # 表頭文字（標題列）也掃一次 meta
    for r in range(1, min(ws.max_row, 6) + 1):
        for c in range(1, min(ws.max_column, 30) + 1):
            v = ws.cell(row=r, column=c).value
            if isinstance(v, str):
                _harvest_meta(v, result.data["meta"])
    parcels = []
    for i, rec in enumerate(records):
        if all(v is MISSING for v in rec.values()):
            continue
        parcels.append(assemble_parcel(rec, result, f"parcels[{i}]"))
    result.data["parcels"] = parcels
    if not parcels:
        result.warn("清冊裡沒有讀到任何宗地")
    return result.finalize()


def _harvest_meta(text: str, meta: dict) -> None:
    m = re.search(r"興辦事業計畫名稱\s*[:：]\s*(\S.*)$", text)
    if m and m.group(1).strip():
        meta["project_name"] = m.group(1).strip()
    m = re.search(r"案號\s*[:：]\s*([0-9A-Za-z\-–]+)", text)
    if m:
        meta["case_no"] = m.group(1).replace("–", "-")
    m = re.search(r"填寫日期\s*[:：]\s*(\S+)", text)
    if m:
        meta["fill_date"] = m.group(1)


# ------------------------------------------------------------------ writer（官方轉置版面）

_thin = Side(style="thin", color="000000")
_border = Border(left=_thin, right=_thin, top=_thin, bottom=_thin)
_center = Alignment(horizontal="center", vertical="center", wrap_text=True)
_FONT = Font(name="標楷體", size=10)

# 寫出時每列的 (group, item_no, label, extractor)
_WRITE_ROWS: list[tuple[str, int | None, str, Any]] = [
    ("0基本資料", None, "宗地流水號", lambda p: p.get("serial_no")),
    ("0基本資料", None, "鄉鎮市區", lambda p: p.get("district")),
    ("0基本資料", None, "段小段名稱", lambda p: _section_name(p)),
    ("0基本資料", None, "地號", lambda p: _lot_no(p)),
    ("0基本資料", None, "土地所有權人或管理人姓名", lambda p: p.get("owner")),
    ("0基本資料", None, "地址", lambda p: p.get("address")),
    ("0基本資料", None, "地價區段", lambda p: p.get("section_id")),
    ("1.宗地條件", 7, "面積(M2)", lambda p: p.get("area_m2")),
    ("1.宗地條件", 8, "寬度(M)", lambda p: p.get("width_m")),
    ("1.宗地條件", 9, "深度(M)", lambda p: p.get("depth_m")),
    ("1.宗地條件", 10, "形狀", lambda p: p.get("shape")),
    ("1.宗地條件", 11, "臨街情形", lambda p: p.get("frontage")),
    ("1.宗地條件", 12, "地勢", lambda p: p.get("terrain")),
    ("2.道路條件", 13, "道路種類", lambda p: p.get("road_type")),
    ("2.道路條件", 14, "面前道路名稱", lambda p: (p.get("front_road") or {}).get("name")),
    ("2.道路條件", 14, "面前道路寬度", lambda p: (p.get("front_road") or {}).get("width_m")),
    ("3.接近條件", 15, "學校名稱", lambda p: _fname(p.get("school"))),
    ("3.接近條件", 15, "接近學校之程度(M)", lambda p: _fdist(p.get("school"))),
    ("3.接近條件", 16, "市場名稱", lambda p: _fname(p.get("market"))),
    ("3.接近條件", 16, "接近市場之程度(M)", lambda p: _fdist(p.get("market"))),
    ("3.接近條件", 17, "公園、廣場名稱", lambda p: _fname(p.get("park"))),
    ("3.接近條件", 17, "接近公園、廣場之程度(M)", lambda p: _fdist(p.get("park"))),
    ("3.接近條件", 18, "車站名稱", lambda p: _fname(p.get("station"))),
    ("3.接近條件", 18, "接近車站之程度(M)", lambda p: _fdist(p.get("station"))),
    ("3.接近條件", 19, "商圈名稱", lambda p: _fname(p.get("commercial_district"))),
    ("3.接近條件", 19, "接近商圈之程度(M)", lambda p: _fdist(p.get("commercial_district"))),
    ("4.周邊環境條件", 20, "嫌惡設施名稱", lambda p: _join([f.get("name") for f in (p.get("nuisance") or [])])),
    ("4.周邊環境條件", 20, "接近嫌惡設施之程度(M)", lambda p: _join([f.get("distance_m") for f in (p.get("nuisance") or [])])),
    ("4.周邊環境條件", 21, "停車方便性", lambda p: p.get("street_parking")),
    ("5.行政條件", 22, "使用分區或編定用地", lambda p: p.get("zoning")),
    ("5.行政條件", 23, "建蔽率(%)", lambda p: p.get("bcr_pct")),
    ("5.行政條件", 24, "容積率(%)", lambda p: p.get("far_pct")),
    ("5.行政條件", 25, "有無禁、限建", lambda p: _yesno(p.get("building_restricted"), p)),
    ("5.行政條件", 25, "禁、限建情形", lambda p: p.get("restriction_note", "無" if p.get("building_restricted") is False else None)),
    ("6.其他", 6, "其他", lambda p: p.get("other")),
    ("協議價購程序", None, "協議價購程序", lambda p: p.get("negotiation")),
]


def _fname(f):
    if f is None:
        return "無"
    return f.get("name") if isinstance(f, dict) else None


def _fdist(f):
    if not isinstance(f, dict):
        return None
    return f.get("distance_m")


def _join(vals):
    vals = [v for v in vals if v is not None]
    if not vals:
        return "無"
    return "、".join(str(int(v)) if isinstance(v, float) and v.is_integer() else str(v) for v in vals)


def _yesno(b, p):
    if b is None:
        return None if "building_restricted" not in p else "-"
    return "有" if b else "無"


def _section_name(p):
    pid = p.get("parcel_id") or ""
    m = re.match(r"^(.*?段)(\d+(?:-\d+)?)地號$", pid)
    return m.group(1) if m else None


def _lot_no(p):
    pid = p.get("parcel_id") or ""
    m = re.match(r"^(.*?段)(\d+(?:-\d+)?)地號$", pid)
    return m.group(2) if m else pid or None


def write_parcels_xlsx(parcels: list[dict], dest: str | Path | io.BytesIO, *, project_name: str = "",
                       case_no: str = "", fill_date: str = "") -> Workbook:
    """Parcel[] → 官方表7 轉置版面 xlsx。dict 沒有的 key 留空白；值為 None 寫「-」；設施為 None 寫「無」。"""
    wb = Workbook()
    ws = wb.active
    ws.title = "表7"
    ws.cell(row=1, column=1, value="表7　宗地個別因素清冊").font = Font(name="標楷體", size=14, bold=True)
    ws.cell(row=2, column=1, value=f"興辦事業計畫名稱：{project_name}")
    ws.cell(row=2, column=6, value=f"案號：{case_no}")
    r0 = 3
    for i, (group, item_no, label, fn) in enumerate(_WRITE_ROWS):
        r = r0 + i
        ws.cell(row=r, column=1, value=group)
        ws.cell(row=r, column=2, value=item_no)
        ws.cell(row=r, column=3, value=label)
        for j, p in enumerate(parcels):
            v = fn(p)
            if v is None:
                v = "-" if _has_key(p, label) else None
            c = ws.cell(row=r, column=4 + j, value=v)
            c.border = _border
            c.alignment = _center
            c.font = _FONT
        for cc in range(1, 4):
            ws.cell(row=r, column=cc).border = _border
            ws.cell(row=r, column=cc).font = _FONT
            ws.cell(row=r, column=cc).alignment = _center
    # 合併主要項目欄
    start = r0
    for i in range(1, len(_WRITE_ROWS) + 1):
        if i == len(_WRITE_ROWS) or _WRITE_ROWS[i][0] != _WRITE_ROWS[start - r0][0]:
            if r0 + i - 1 > start:
                ws.merge_cells(start_row=start, start_column=1, end_row=r0 + i - 1, end_column=1)
            start = r0 + i
    ws.cell(row=r0 + len(_WRITE_ROWS) + 1, column=1, value=f"填寫日期：{fill_date}")
    ws.cell(row=r0 + len(_WRITE_ROWS) + 1, column=6, value="○○機關（即需用土地人）核章：承辦人：　　單位主管：　　首長：")
    ws.column_dimensions["A"].width = 12
    ws.column_dimensions["B"].width = 5
    ws.column_dimensions["C"].width = 26
    for j in range(len(parcels)):
        ws.column_dimensions[get_column_letter(4 + j)].width = 16
    wb.save(dest)
    return wb


_LABEL_TO_KEY = {
    "面積(M2)": "area_m2", "寬度(M)": "width_m", "深度(M)": "depth_m", "形狀": "shape", "臨街情形": "frontage",
    "地勢": "terrain", "道路種類": "road_type", "面前道路寬度": "front_road", "面前道路名稱": "front_road",
    "學校名稱": "school", "接近學校之程度(M)": "school", "市場名稱": "market", "接近市場之程度(M)": "market",
    "公園、廣場名稱": "park", "接近公園、廣場之程度(M)": "park", "車站名稱": "station", "接近車站之程度(M)": "station",
    "商圈名稱": "commercial_district", "接近商圈之程度(M)": "commercial_district",
    "嫌惡設施名稱": "nuisance", "接近嫌惡設施之程度(M)": "nuisance", "停車方便性": "street_parking",
    "使用分區或編定用地": "zoning", "建蔽率(%)": "bcr_pct", "容積率(%)": "far_pct", "有無禁、限建": "building_restricted",
    "禁、限建情形": "restriction_note", "其他": "other", "協議價購程序": "negotiation", "地價區段": "section_id",
    "宗地流水號": "serial_no", "鄉鎮市區": "district", "土地所有權人或管理人姓名": "owner",
    "段小段名稱": "parcel_id", "地號": "parcel_id", "地址": "address",
}


def _has_key(p: dict, label: str) -> bool:
    k = _LABEL_TO_KEY.get(label)
    return k is not None and k in p
