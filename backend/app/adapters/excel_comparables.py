"""
買賣實例（比較標的）Excel → Comparable[]。

Comparable = Parcel 欄位 + 交易資料（實例編號、土地正常單價、交易日期、期日調整、權重、實價登錄來源）。
版面：預設一列一筆（long），也接受表7 式轉置版面。欄名見 COMPARABLE_SPECS。
write_comparables_xlsx() 依 long 版面寫出，供 round-trip 測試與當天反向產表。
"""
from __future__ import annotations

import io
from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, Side

from .common import MISSING, AdapterResult, FieldSpec, parse_by_kind, read_records
from .excel_parcels import _WRITE_ROWS, PARCEL_SPECS, assemble_parcel

SOURCE_COMPARABLES = "買賣實例調查估價表"

TRADE_SPECS: list[FieldSpec] = [
    FieldSpec("comp_no", ("實例編號", "比較標的編號", "比較標的", "comp_no"), "int", required=True),
    FieldSpec("normal_unit_price", ("土地正常單價", "土地正常單價(元/M2)", "正常單價", "normal_unit_price"), "num", required=True),
    FieldSpec("transaction_date", ("交易日期", "transaction_date"), "date", required=True),
    FieldSpec("price_total", ("交易總價", "總價", "price_total"), "num"),
    FieldSpec("date_adjustment_pct", ("期日調整率", "期日調整率(%)", "調整百分率", "date_adjustment_pct"), "num"),
    FieldSpec("index_at_valuation", ("估價基準日地價指數", "基準日地價指數", "index_at_valuation"), "num"),
    FieldSpec("index_at_transaction", ("交易日地價指數", "交易期地價指數", "index_at_transaction"), "num"),
    FieldSpec("date_note", ("期日調整說明", "date_note"), "text"),
    FieldSpec("weight_pct", ("權重", "權重(%)", "比較標的權重", "weight_pct"), "num"),
    FieldSpec("lvr_id", ("實價登錄編號", "實價登錄序號", "lvr_id"), "text"),
    FieldSpec("source_note", ("資料來源", "備註", "source_note"), "text"),
]
COMPARABLE_SPECS: list[FieldSpec] = TRADE_SPECS + PARCEL_SPECS


def read_comparables(src: str | Path | bytes | io.BytesIO, *, sheet: str | None = None) -> AdapterResult:
    result = AdapterResult(kind="comparables", data={"comparables": [], "meta": {}})
    wb = load_workbook(io.BytesIO(src) if isinstance(src, (bytes, bytearray)) else src, data_only=True)
    ws = wb[sheet] if sheet else wb.active
    try:
        records, layout, unknown = read_records(ws, COMPARABLE_SPECS, "comp_no")
    except ValueError as e:
        result.warn(str(e))
        return result
    result.data["meta"] = {"layout": layout, "sheet": ws.title}
    for lbl in unknown:
        result.warn(f"未對應到 schema 的欄位「{lbl}」已略過")
    comps = []
    for i, rec in enumerate(records):
        if all(v is MISSING for v in rec.values()):
            continue
        path = f"comparables[{i}]"
        c = assemble_parcel(rec, result, path, source=SOURCE_COMPARABLES)
        t = {s.key: parse_by_kind(s.kind, rec.get(s.key, MISSING)) for s in TRADE_SPECS}

        def v(k: str, _t=t):
            return None if _t[k] is MISSING else _t[k]
        c["comp_no"] = v("comp_no") if v("comp_no") is not None else i + 1
        if v("comp_no") is None:
            result.warn(f"{path}.comp_no 缺，依順序給 {i + 1}")
        c["normal_unit_price"] = v("normal_unit_price")
        if c["normal_unit_price"] is None:
            result.missing(f"{path}.normal_unit_price")
        c["transaction_date"] = v("transaction_date")
        if c["transaction_date"] is None:
            result.missing(f"{path}.transaction_date")
        da: dict[str, Any] = {}
        if v("date_adjustment_pct") is not None:
            da["pct"] = v("date_adjustment_pct")
        if v("index_at_valuation") is not None:
            da["index_at_valuation"] = v("index_at_valuation")
        if v("index_at_transaction") is not None:
            da["index_at_transaction"] = v("index_at_transaction")
        if v("date_note") is not None:
            da["note"] = v("date_note")
        if "pct" not in da and not ("index_at_valuation" in da and "index_at_transaction" in da):
            result.missing(f"{path}.date_adjustment.pct")
        c["date_adjustment"] = da
        if v("weight_pct") is not None:
            c["weight_pct"] = v("weight_pct")
        src_d: dict[str, Any] = {}
        if v("lvr_id") is not None:
            src_d["lvr_id"] = v("lvr_id")
        if v("price_total") is not None:
            src_d["price_total"] = v("price_total")
        if v("source_note") is not None:
            src_d["note"] = v("source_note")
        if src_d:
            c["source"] = src_d
        comps.append(c)
    comps.sort(key=lambda c: c["comp_no"])
    result.data["comparables"] = comps
    if not comps:
        result.warn("沒有讀到任何買賣實例")
    return result.finalize()


# ------------------------------------------------------------------ writer（long 版面）

_thin = Side(style="thin")
_border = Border(left=_thin, right=_thin, top=_thin, bottom=_thin)

_TRADE_COLS: list[tuple[str, Any]] = [
    ("實例編號", lambda c: c.get("comp_no")),
    ("土地正常單價(元/M2)", lambda c: c.get("normal_unit_price")),
    ("交易日期", lambda c: c.get("transaction_date")),
    ("交易總價", lambda c: (c.get("source") or {}).get("price_total")),
    ("期日調整率(%)", lambda c: (c.get("date_adjustment") or {}).get("pct")),
    ("估價基準日地價指數", lambda c: (c.get("date_adjustment") or {}).get("index_at_valuation")),
    ("交易日地價指數", lambda c: (c.get("date_adjustment") or {}).get("index_at_transaction")),
    ("期日調整說明", lambda c: (c.get("date_adjustment") or {}).get("note")),
    ("權重(%)", lambda c: c.get("weight_pct")),
    ("實價登錄編號", lambda c: (c.get("source") or {}).get("lvr_id")),
    ("資料來源", lambda c: (c.get("source") or {}).get("note")),
]


def write_comparables_xlsx(comps: list[dict], dest: str | Path | io.BytesIO) -> Workbook:
    wb = Workbook()
    ws = wb.active
    ws.title = "買賣實例"
    cols = _TRADE_COLS + [(label, fn) for (_g, _no, label, fn) in _WRITE_ROWS]
    for j, (label, _fn) in enumerate(cols, start=1):
        c = ws.cell(row=1, column=j, value=label)
        c.font = Font(name="標楷體", bold=True)
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = _border
    for i, comp in enumerate(comps, start=2):
        for j, (label, fn) in enumerate(cols, start=1):
            v = fn(comp)
            if v is None and label in ("其他", "停車方便性", "建蔽率(%)", "容積率(%)") and _has(comp, label):
                v = "-"
            cell = ws.cell(row=i, column=j, value=v)
            cell.border = _border
            cell.font = Font(name="標楷體")
    return _save(wb, dest)


def _has(comp: dict, label: str) -> bool:
    from .excel_parcels import _LABEL_TO_KEY
    k = _LABEL_TO_KEY.get(label)
    return k is not None and k in comp


def _save(wb: Workbook, dest) -> Workbook:
    wb.save(dest)
    return wb
