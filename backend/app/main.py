"""
FastAPI 入口。
  POST /api/run          生成模式：案件 JSON → 表5 + 表4
  POST /api/verify       審查模式：案件 JSON + 估價師填的表 → findings
  GET  /api/rules        目前載入的基準表（給前端畫判定條件用）
  POST /api/adapt        上傳檔案（表7 清冊 / 買賣實例 / 基準明細表 / 書表 PDF）→ 內部 schema + missing_fields
  POST /api/export/xlsx  案件 JSON → 表5 + 表4 Excel（照範本版面）
擴充方向見 docs/04_architecture.md。
"""
from __future__ import annotations

import copy
import hashlib
import io
import json
import re
import time
import urllib.parse
import uuid
from pathlib import Path
from typing import Any

from fastapi import FastAPI, File, Form, HTTPException, Request, UploadFile
from fastapi.exceptions import RequestValidationError
from fastapi.responses import HTMLResponse, JSONResponse, Response, StreamingResponse
from pydantic import BaseModel

from app import audit as AUD
from app.engine.rules import RULES_DIR, load_ruleset
from app.engine.tables import InputError, run_case
from app.engine.verify import collect_findings

app = FastAPI(title="土地徵收補償市價查估 估價案件審查輔助系統", version="0.2.0")


@app.on_event("startup")
def _warm_caches() -> None:
    """啟動時在背景把設施庫、分區庫、實價登錄、門牌索引與本區步行圖載進來，第一個請求不必等十幾秒。"""
    import threading

    def run():
        try:
            from app.maps.zoning import get_zoning_store
            from app.market.lvr import load_lvr
            from app.spatial import service
            from app.spatial.admin import _lvr_zone_index
            from app.spatial.locate import address_db
            service.get_poi_store(); get_zoning_store(); load_lvr(); address_db(); _lvr_zone_index()
        except Exception:  # noqa: BLE001, S110 - 暖機失敗不影響服務
            pass
    threading.Thread(target=run, daemon=True).start()

_PYD_MSG = {
    "Field required": "為必填", "Input should be a valid integer": "須為整數", "Input should be a valid number": "須為數字",
    "Input should be a valid string": "須為文字", "Input should be a valid dictionary": "格式須為物件", "Input should be a valid list": "格式須為清單",
    "Input should be a valid boolean": "須為是／否", "Input should be a valid dictionary or object to extract fields from": "格式須為物件",
}
_FIELD_ZH = {"case": "案件", "sections": "地價區段", "subject_parcel": "比準地", "comparables": "比較標的", "submitted_table5": "區域因素分析明細表填載值",
             "submitted_table4": "比較法調查估價表填載值", "file": "檔案", "kind": "檔案種類", "hints": "點選位置", "parcels": "宗地", "ruleset": "基準表"}


@app.exception_handler(RequestValidationError)
async def _validation_zh(request: Request, exc: RequestValidationError):
    """把 pydantic 的英文驗證錯誤翻成承辦看得懂的中文。"""
    msgs = []
    for e in exc.errors():
        loc = [str(x) for x in e.get("loc", []) if x not in ("body", "query", "path")]
        field = "／".join(_FIELD_ZH.get(x, x) for x in loc) or "資料"
        m = e.get("msg", "")
        zh = next((v for k, v in _PYD_MSG.items() if m.startswith(k)), m)
        msgs.append(f"{field}{zh}")
    return JSONResponse(status_code=422, content={"detail": "資料格式有誤：" + "；".join(msgs[:6]) + ("…" if len(msgs) > 6 else "")})


@app.exception_handler(InputError)
async def _input_error(request: Request, exc: InputError):
    """案件資料不足以核算（比較標的缺正常單價等）：回 422 中文，不是 500。"""
    return JSONResponse(status_code=422, content={"detail": str(exc)})


class CasePayload(BaseModel):
    case: dict[str, Any]
    sections: dict[str, Any]
    subject_parcel: dict[str, Any]
    comparables: list[dict[str, Any]]


class VerifyPayload(CasePayload):
    submitted_table5: dict[int, dict[str, Any]] | None = None   # {comp_no: {...}}
    submitted_table4: dict[str, Any] | None = None


def _rulesets(case: dict[str, Any]):
    """案件指定的基準表；若檔案已不存在（例如匯入的基準表被刪），退回同用地別的內建表，不讓整個案件打不開。"""
    from app.cases import pick_rulesets
    rs = case.get("rulesets") or {}
    fallback = pick_rulesets(case.get("land_use"))
    out = []
    for scope in ("regional", "individual"):
        rid = rs.get(scope) or fallback[scope]
        try:
            out.append(load_ruleset(rid))
        except FileNotFoundError:
            try:
                out.append(load_ruleset(fallback[scope]))
                rs[scope] = fallback[scope]      # 回寫到案件資料，下次列表與頁面顯示的就是實際使用的表
            except FileNotFoundError as e:
                raise HTTPException(400, f"找不到基準表：{e}") from e
    return out[0], out[1]


@app.get("/api/health")
def health():
    from app.llm import provider_status
    from app.maps.render import tile_cache_status
    from app.market.lvr import lvr_status
    from app.spatial import service
    from app.spatial.terrain import status as terrain_status
    return {"ok": True, "llm": provider_status(), "spatial": service.status(), "tiles": tile_cache_status(), "lvr": lvr_status(), "terrain": terrain_status(),
            "data_files": service.data_file_times()}


@app.post("/api/reload")
def reload_data(request: Request):
    """換了 data/ 或 rules/ 檔案後清掉快取重新載入（不必重啟服務）；記操作紀錄。"""
    from app.spatial import service
    out = service.reload_all()
    AUD.log("_system", AUD.actor_from_headers(request.headers), "reload", "重新載入資料檔：" + "、".join(out["cleared"]))
    return out


def _ruleset_summary(p: Path) -> dict[str, Any]:
    d = json.loads(p.read_text(encoding="utf-8"))
    return {"id": d["id"], "scope": d["scope"], "land_use": d.get("land_use"), "table": d.get("table"), "source": d.get("source"),
            "note": d.get("$schema_note"), "n_rules": len(d.get("rules", [])), "groups": d.get("groups", []),
            "is_demo": str(d.get("source", "")).startswith("示範") or d["id"].startswith("demo_")}


@app.get("/api/rules")
def rules(detail: bool = False):
    """目前可用的基準表。detail=true 回完整規則（含 criteria／矩陣，給表格點格子看依據用）。"""
    out = {}
    for p in sorted(RULES_DIR.glob("*.json")):
        if "_regional" not in p.name and "_individual" not in p.name:
            continue
        if detail:
            out[p.stem] = json.loads(p.read_text(encoding="utf-8"))
        else:
            out[p.stem] = _ruleset_summary(p)
    return out


@app.get("/api/rules/{rid}")
def rule_detail(rid: str):
    p = RULES_DIR / f"{rid}.json"
    if not p.exists():
        raise HTTPException(404, "沒有這份基準表")
    d = json.loads(p.read_text(encoding="utf-8"))
    from app.engine.rules import Rule, adjustment
    for r in d["rules"]:
        rule = Rule(id=r["id"], name=r["name"], group=r["group"], levels=r["levels"], max_pct=r["max_pct"], criteria=r["criteria"], explicit_matrix=r.get("matrix"))
        r["matrix_full"] = {s_: {c_: adjustment(rule, s_, c_) for c_ in r["levels"]} for s_ in r["levels"]} if r["criteria"].get("type") != "manual" else None
    return d


class RulesImportPayload(BaseModel):
    ruleset: dict[str, Any]
    id: str | None = None


@app.post("/api/rules/import")
def rules_import(payload: RulesImportPayload):
    """匿存 /api/adapt(rules_table) 轉出或使用者上傳的基準表 JSON → rules/uploaded_<id>.json，之後 case.rulesets 可直接引用。"""
    d = payload.ruleset
    for k in ("scope", "rules", "land_use"):
        if k not in d:
            raise HTTPException(422, f"基準表 JSON 缺 {k}")
    rid = _store_ruleset(d, payload.id)
    return {"id": rid, "summary": _ruleset_summary(RULES_DIR / f"{rid}.json")}


def _store_ruleset(d: dict[str, Any], rid: str | None = None) -> str:
    """基準表 JSON → rules/uploaded_<id>_<scope>.json（先試載入，載不進就不存）。回傳 id。"""
    import tempfile
    rid = re.sub(r"[^A-Za-z0-9_]", "_", rid or d.get("id") or "uploaded")
    if not rid.endswith(f"_{d['scope']}"):
        rid = f"{rid}_{d['scope']}"
    if not rid.startswith("uploaded_"):
        rid = "uploaded_" + rid
    d["id"] = rid
    with tempfile.NamedTemporaryFile("w", suffix=".json", delete=False, encoding="utf-8") as f:
        json.dump(d, f, ensure_ascii=False)
        tmp = f.name
    try:
        load_ruleset(tmp)
    except Exception as e:
        raise HTTPException(422, f"基準表 JSON 無法載入：{e}") from e
    (RULES_DIR / f"{rid}.json").write_text(json.dumps(d, ensure_ascii=False, indent=2), encoding="utf-8")
    return rid


@app.post("/api/run")
def run(payload: CasePayload):
    reg, ind = _rulesets(payload.case)
    result = run_case(reg, ind, payload.model_dump())
    return {"table5": {k: v.to_dict() for k, v in result["table5"].items()}, "table4": result["table4"].to_dict()}


@app.post("/api/verify")
def verify(payload: VerifyPayload):
    reg, ind = _rulesets(payload.case)
    data = payload.model_dump()
    result = run_case(reg, ind, data)
    findings = collect_findings(reg, ind, data, result, payload.submitted_table5, payload.submitted_table4)
    return {"findings": findings, "computed": {"table5": {k: v.to_dict() for k, v in result["table5"].items()},
                                              "table4": result["table4"].to_dict()}}


# ------------------------------------------------------------------ adapter / export

ADAPT_KINDS = ("auto", "parcels", "comparables", "rules_table", "pdf_forms")


def detect_kind(filename: str, content: bytes) -> str:
    """依副檔名 + 內容判斷是哪種輸入。"""
    name = (filename or "").lower()
    if name.endswith(".pdf") or content[:4] == b"%PDF":
        import fitz
        doc = fitz.open(stream=content, filetype="pdf")
        text = "".join(doc[i].get_text("text") for i in range(min(len(doc), 3)))
        return "rules_table" if "評價基準明細表" in text else "pdf_forms"
    if name.endswith(".csv"):
        return "rules_table"
    if name.endswith((".xlsx", ".xlsm", ".xls")):
        from openpyxl import load_workbook

        from app.adapters.common import norm_label
        wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        labels = set()
        for ws in wb.worksheets:
            for row in ws.iter_rows(min_row=1, max_row=60, max_col=60, values_only=True):
                labels.update(norm_label(v) for v in row if isinstance(v, str))
        if "pct_優" in labels or "pct優" in labels:
            return "rules_table"
        if labels & {"實例編號", "土地正常單價", "交易日期"}:
            return "comparables"
        if labels & {"宗地流水號", "面積", "地號"}:
            return "parcels"
    raise HTTPException(400, f"無法判斷檔案種類：{filename}；請以 kind 指定 {ADAPT_KINDS[1:]}")


@app.post("/api/adapt")
async def adapt(file: UploadFile = File(...), kind: str = Form("auto"), land_use: str | None = Form(None),  # noqa: B008
                use_vision: str = Form("auto"), id_prefix: str | None = Form(None)):
    """上傳 → AdapterResult {kind, data, missing_fields, warnings, confidence}。缺欄位不猜，列在 missing_fields 讓 UI 補。"""
    if kind not in ADAPT_KINDS:
        raise HTTPException(400, f"kind 須為 {ADAPT_KINDS}")
    content = await file.read()
    if not content:
        raise HTTPException(400, "空檔案")
    k = detect_kind(file.filename, content) if kind == "auto" else kind
    try:
        if k == "parcels":
            from app.adapters.excel_parcels import read_parcels
            res = read_parcels(content)
        elif k == "comparables":
            from app.adapters.excel_comparables import read_comparables
            res = read_comparables(content)
        elif k == "rules_table":
            from app.adapters.rules_table import read_rules_table
            res = read_rules_table(content, filename=file.filename, land_use=land_use, id_prefix=id_prefix)
        else:
            res = _read_pdf_forms_cached(content, file.filename, use_vision)
    except HTTPException:
        raise
    except Exception as e:
        raise HTTPException(422, f"檔案解析失敗（{k}）：{e}") from e
    out = res.to_dict()
    out["filename"] = file.filename
    return out


_ADAPT_CACHE: dict[str, tuple[float, Any]] = {}     # sha1|use_vision → (time, AdapterResult)；首頁先預覽再建案，同一份 PDF 不重跑影像辨識
_ADAPT_TTL = 1800


def _read_pdf_forms_cached(content: bytes, filename: str | None, use_vision: str):
    from app.adapters.pdf_forms import read_pdf_forms
    key = f"{hashlib.sha1(content).hexdigest()}|{use_vision}"
    hit = _ADAPT_CACHE.get(key)
    if hit and time.time() - hit[0] < _ADAPT_TTL:
        return copy.deepcopy(hit[1])
    res = read_pdf_forms(content, filename=filename, use_vision=use_vision)
    if len(_ADAPT_CACHE) >= 32:
        _ADAPT_CACHE.pop(min(_ADAPT_CACHE, key=lambda k: _ADAPT_CACHE[k][0]))
    _ADAPT_CACHE[key] = (time.time(), copy.deepcopy(res))
    return res


class ExportPayload(CasePayload):
    meta: dict[str, Any] | None = None   # appraiser, fill_date, notes{subject, comparables{comp_no: text}, case}, comp_labels
    figures: bool = True                  # 附「圖說」工作表（三張 PNG）；沒有幾何就略過
    case_id: str | None = None            # 有給就記操作紀錄


@app.post("/api/export/xlsx")
def export_xlsx(payload: ExportPayload, request: Request):
    from app.output.xlsx import build_workbook
    reg, ind = _rulesets(payload.case)
    data = payload.model_dump(exclude={"meta", "figures", "case_id"})
    figs = None
    if payload.figures:
        try:
            from app.maps.render import case_figures
            figs = case_figures(data, appraiser=(payload.meta or {}).get("appraiser") or data["case"].get("appraiser", ""))
        except ValueError:
            figs = None   # 沒有幾何：不附圖
    try:
        meta = {"appraiser": data["case"].get("appraiser", ""), "fill_date": data["case"].get("fill_date", ""), **{k: v for k, v in (payload.meta or {}).items() if v}}
        wb = build_workbook(data, meta=meta, regional=reg, individual=ind, figures=figs)
    except KeyError as e:
        raise HTTPException(422, f"案件資料缺欄位：{e}") from e
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    fname = f"{payload.case.get('case_no', 'case')}_審查書表.xlsx"
    if payload.case_id:
        AUD.log(payload.case_id, AUD.actor_from_headers(request.headers), "export", fname + ("（含圖說）" if figs else ""))
    headers = {"Content-Disposition": f"attachment; filename*=UTF-8''{urllib.parse.quote(fname)}"}
    return StreamingResponse(buf, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", headers=headers)


# ------------------------------------------------------------------ spatial（第 4 步）


class DistancesPayload(BaseModel):
    origin: dict[str, Any]                       # GeoJSON Point/Polygon（WGS84）
    types: list[str]                             # facility_measurement.json 的 key
    k: int = 3
    max_m: float | None = 5000
    mode_overrides: dict[str, str] | None = None
    section: dict[str, Any] | None = None        # 給了就判 in_section
    origin_label: str = "parcel_centroid"


@app.get("/api/spatial/status")
def spatial_status():
    from app.spatial import service
    return service.status()


@app.post("/api/spatial/distances")
def spatial_distances(payload: DistancesPayload):
    """點/多邊形 → 各類設施最近 k 筆的 Facility（含 measure/origin/source/佐證）。"""
    from app.spatial import service
    from app.spatial.distance import facility_from_poi
    from app.spatial.reference import mode_for_type
    store, osrm = service.get_poi_store(), service.get_osrm()
    out: dict[str, list] = {}
    for t in payload.types:
        facs = []
        for poi, _d in store.nearest(payload.origin, [t], k=payload.k, max_m=payload.max_m):
            facs.append(facility_from_poi(poi, payload.origin, mode_for_type(t, payload.mode_overrides), origin_label=payload.origin_label,
                                          osrm=osrm, section=payload.section))
        out[t] = facs
    return {"facilities": out, "osrm": osrm.status(), "poi_count": len(store)}


class FillPayload(CasePayload):
    origin_mode_individual: str = "parcel_centroid"     # parcel_centroid | parcel_frontage
    origin_mode_regional: str = "section_boundary"      # section_boundary | subject_parcel
    mode_overrides: dict[str, str] | None = None        # {facility_type: "walking"|"straight"}
    overwrite: bool = False


@app.post("/api/spatial/fill")
def spatial_fill(payload: FillPayload):
    """有 geometry 的案件 → 自動填距離設施（案件層級同一參照設施），回填好的案件 + provenance + warnings。"""
    from app.spatial import service
    reg, ind = _rulesets(payload.case)
    data = payload.model_dump(exclude={"origin_mode_individual", "origin_mode_regional", "mode_overrides", "overwrite"})
    try:
        return service.fill_case(data, regional=reg, individual=ind, origin_mode_individual=payload.origin_mode_individual,
                                 origin_mode_regional=payload.origin_mode_regional, mode_overrides=payload.mode_overrides,
                                 overwrite=payload.overwrite)
    except ValueError as e:
        raise HTTPException(422, str(e)) from e


# ------------------------------------------------------------------ 地籍幾何 / 區段 bootstrap / 三張圖（第 4b 步）

FIXTURES_DIR = Path(__file__).resolve().parents[2] / "fixtures"


@app.get("/api/cases/demo")
def demo_case_endpoint(request: Request, variant: str = "template", save: bool = False):
    """demo 變體：template | tampered | residential。save=true 存成案件並回 id。"""
    from app import cases as C
    if variant not in ("template", "tampered", "residential", "blank_survey", "shulin"):
        raise HTTPException(400, "variant 須為 template | tampered | residential | blank_survey | shulin")
    d = C.demo_case(variant)
    if save:
        rec = C.save_case(d["data"], name=d["name"], origin=f"demo:{variant}", submitted_table5=d["submitted_table5"], submitted_table4=d["submitted_table4"])
        d["id"] = rec["id"]
        if variant == "shulin":                                        # 決賽題目：存檔後直接依地號產生（界線、區段、勘查表、宗地屬性），不然表4 個別因素會是空的
            try:
                out = cases_from_lot(rec["id"], FromLotPayload(parcel_id=d["data"]["subject_parcel"].get("parcel_id") or "", overwrite=False, with_comparables=False), request)
                d["steps"] = out.get("steps")
            except Exception as e:  # noqa: BLE001 - 圖資或外部查詢失敗不擋範例載入
                d["steps"] = [{"step": "依地號產生", "ok": False, "note": str(e)}]
    # 舊介面相容：頂層直接是 case 資料
    return {**d["data"], **{k: v for k, v in d.items() if k != "data"}}


class SaveCasePayload(CasePayload):
    name: str | None = None
    id: str | None = None
    submitted_table5: dict[int, dict[str, Any]] | None = None
    submitted_table4: dict[str, Any] | None = None
    extraction: dict[str, Any] | None = None      # PDF 抽取的信心值／缺漏／提醒（{confidence:{path:0~1}, missing_fields:[], warnings:[], pages:[]}）
    status: str | None = None                      # draft | reviewing | done
    decisions: dict[str, Any] | None = None        # {finding_key: {decision: accept|reject|pending, note, by, at}}


class PatchCasePayload(BaseModel):
    name: str | None = None
    status: str | None = None
    decisions: dict[str, Any] | None = None
    extraction: dict[str, Any] | None = None


@app.get("/api/cases")
def cases_list():
    from app import cases as C
    return {"cases": C.list_cases()}



_ORIGIN_LABELS = {"manual": "送審書表或手動建立", "demo:template": "範例", "demo:tampered": "範例（含填載錯誤）", "demo:residential": "範例（住宅用地）", "demo:shulin": "決賽題目（樹林區）",
                  "demo:blank_survey": "範例（僅勘查表）", "from_lot": "依地號產生", "import": "檔案匯入", "inputs": "輸入檔建立"}


def _origin_label(origin: str | None) -> str:
    return _ORIGIN_LABELS.get(origin or "manual", origin or "手動建立")


def _attach_geometry(data: dict[str, Any]) -> list[str]:
    """送審書表建案時補幾何（只補缺的，不動填載值）：比準地與比較標的界線（本案地籍圖→預載地籍圖→門牌定位）、
    區段範圍（路網推估街廓，草稿）。有幾何三張圖說與地圖才畫得出來；設施距離仍用書表抽取值。回傳做了什麼。"""
    from app.spatial.bootstrap import block_from_roads, describe_range
    from app.spatial.geo import centroid
    from app.spatial.lot import _major_zone
    from app.spatial.roads_store import get_roads
    notes: list[str] = []
    subject = data.get("subject_parcel") or {}
    try:
        _resolve_case_geometries(data, None, False)
    except Exception as e:  # noqa: BLE001 - 幾何補不到不影響建案
        notes.append(f"地籍界線：{e}")
    if subject.get("geometry") is not None:
        notes.append(f"比準地界線：{subject.get('geometry_note') or subject.get('geometry_source')}")
    for c in data.get("comparables") or []:
        if c.get("geometry") is None:
            try:
                if _locate_by_address(c, data):
                    notes.append(f"比較標的{c.get('comp_no')} 位置依門牌推定")
            except Exception as e:  # noqa: BLE001 - 對不到就留人工點圖
                notes.append(f"比較標的{c.get('comp_no')} 位置未推定：{e}")
    sid = subject.get("section_id") or next(iter(data.get("sections") or {}), None)
    sec = (data.get("sections") or {}).get(sid) if sid else None
    try:                                                                   # 有四至文字的區段先由路名圍面（比準地與比較標的所在區段都做）
        from app.maps.zoning import get_zoning_store
        from app.spatial.lot import section_from_range_text
        roads_ = get_roads(data=data)
        for osid, osec in (data.get("sections") or {}).items():
            if osec.get("geometry") is None and (osec.get("range_desc") or "").strip():
                r_ = section_from_range_text(osec, roads_, get_zoning_store())
                notes.append(f"區段 {osid}：{'依四至圍出' if r_['ok'] else '四至圍不出'}")
    except Exception as e:  # noqa: BLE001
        notes.append(f"四至圍面：{e}")
    if sec is not None and sec.get("geometry") is None and subject.get("geometry") is not None:
        try:
            roads = get_roads(data=data)
            c = centroid(subject["geometry"])
            res = block_from_roads([(c.x, c.y)], roads) if roads is not None and getattr(roads, "items", None) else None
        except Exception as e:  # noqa: BLE001
            res, roads = None, None
            notes.append(f"區段範圍：{e}")
        if res:
            from app.maps.zoning import get_zoning_store
            try:
                z = _major_zone(res["geometry"], get_zoning_store())
            except Exception:  # noqa: BLE001
                z = None
            sec["geometry"], sec["geometry_source"], sec["status"] = res["geometry"], "estimate_osm_block", "draft"
            sec["geometry_note"] = f"依比準地位置由路網推估之街廓（邊界道路：{'、'.join(res['bounding_roads'])}），需估價人員確認；書表填載之區段範圍文字保留"
            if not (sec.get("range_desc") or "").strip():
                sec["range_desc"] = describe_range(res["geometry"], roads, zoning=z, section_id=sid)["range_desc"]
            notes.append(f"區段範圍：路網推估街廓 {res['area_m2']:,.0f} m²（草稿）")
    return notes

@app.post("/api/cases")
def cases_save(payload: SaveCasePayload, request: Request):
    from app import cases as C
    data = payload.model_dump(exclude={"name", "id", "submitted_table5", "submitted_table4", "extraction", "status", "decisions"})
    prev = copy.deepcopy(C.get_case(payload.id)) if payload.id else None   # 存檔會就地改 _mem，先複製才比得出差異
    geom_notes: list[str] = []
    if prev is None and (payload.submitted_table5 or payload.submitted_table4) and (data.get("subject_parcel") or {}).get("geometry") is None:
        geom_notes = _attach_geometry(data)                                 # 送審書表建案：補界線與區段範圍，圖說才畫得出來
    try:
        rec = C.save_case(data, name=payload.name, cid=payload.id, submitted_table5=payload.submitted_table5, submitted_table4=payload.submitted_table4,
                          extraction=payload.extraction, status=payload.status, decisions=payload.decisions)
    except ValueError as e:
        raise HTTPException(422, str(e)) from e
    actor = AUD.actor_from_headers(request.headers)
    if prev is None:
        AUD.log(rec["id"], actor, "create", f"{rec.get('name') or ''}（來源 {_origin_label(rec.get('origin'))}）" + (f"；{'；'.join(geom_notes)}" if geom_notes else ""))
    else:
        changes = AUD.diff_paths({"data": prev.get("data"), "submitted_table5": prev.get("submitted_table5"), "submitted_table4": prev.get("submitted_table4"), "name": prev.get("name")},
                                 {"data": rec.get("data"), "submitted_table5": rec.get("submitted_table5"), "submitted_table4": rec.get("submitted_table4"), "name": rec.get("name")})
        if changes:
            AUD.log(rec["id"], actor, "save", f"變更 {len(changes)} 處", changes)
    return rec



@app.post("/api/cases/{cid}/geometry/attach")
def cases_geometry_attach(cid: str, request: Request):
    """既有案件補幾何（比準地／比較標的界線、區段範圍草稿），不動填載值。"""
    from app import cases as C
    rec = C.get_case(cid)
    if not rec:
        raise HTTPException(404, "沒有這個案件")
    data = copy.deepcopy(rec["data"])
    notes = _attach_geometry(data)
    rec = C.save_case(data, cid=cid, submitted_table5=rec.get("submitted_table5"), submitted_table4=rec.get("submitted_table4"))
    AUD.log(cid, AUD.actor_from_headers(request.headers), "geometry_attach", "；".join(notes) or "無變更")
    return {"rec": rec, "notes": notes}

@app.patch("/api/cases/{cid}")
def cases_patch(cid: str, payload: PatchCasePayload, request: Request):
    from app import cases as C
    prev = copy.deepcopy(C.get_case(cid))
    try:
        rec = C.patch_case(cid, name=payload.name, status=payload.status, decisions=payload.decisions, extraction=payload.extraction)
    except ValueError as e:
        raise HTTPException(422, str(e)) from e
    if not rec:
        raise HTTPException(404, "沒有這個案件")
    actor = AUD.actor_from_headers(request.headers)
    if payload.status is not None and payload.status != (prev or {}).get("status"):
        AUD.log(cid, actor, "status", f"{AUD.STATUS_ZH.get((prev or {}).get('status') or '', '—')} → {AUD.STATUS_ZH.get(payload.status, payload.status)}")
    if payload.decisions is not None:
        changes = AUD.diff_paths((prev or {}).get("decisions") or {}, payload.decisions, ignore_prefixes=("at",))
        if changes:
            AUD.log(cid, actor, "decisions", f"變更 {len(changes)} 處", changes)
    if payload.name is not None and payload.name != (prev or {}).get("name"):
        AUD.log(cid, actor, "patch", f"案件名稱：{(prev or {}).get('name')} → {payload.name}")
    return rec


@app.post("/api/cases/{cid}/generate")
def cases_generate(cid: str, request: Request):
    """重新產生書表：跑引擎與審查，把輸入指紋、時間、摘要記在案件上；之後輸入再改動就標「已過期」。"""
    from app import cases as C
    from app.report.opinion import finding_key
    rec = C.get_case(cid)
    if not rec:
        raise HTTPException(404, "沒有這個案件")
    data = rec["data"]
    from app.market.index import fill_date_adjustments
    if fill_date_adjustments(data):          # 比較標的缺期日調整率者依都市地價指數補上（手冊 p.50 (四)），存回案件
        rec = C.save_case(data, cid=cid, submitted_table5=rec.get("submitted_table5"), submitted_table4=rec.get("submitted_table4"))
        data = rec["data"]
    reg, ind = _rulesets(data["case"])
    result = run_case(reg, ind, data)
    findings = collect_findings(reg, ind, data, result, rec.get("submitted_table5"), rec["submitted_table4"])
    t4 = result["table4"]
    outputs = {"generated_at": C._now(), "input_hash": rec.get("input_hash") or C.input_hash(data, rec.get("submitted_table5"), rec.get("submitted_table4")),
               "summary": {"subject_comparison_price": t4.subject_comparison_price, "subject_land_price": t4.subject_land_price,
                           "n_error": sum(1 for f in findings if f["severity"] == "error"), "n_warn": sum(1 for f in findings if f["severity"] == "warn")},
               "findings_keys": sorted({finding_key(f) or f"loc:{f.get('table')}:{f.get('location')}" for f in findings})}
    rec = C.set_outputs(cid, outputs)
    AUD.log(cid, AUD.actor_from_headers(request.headers), "generate", f"比較價格 {t4.subject_comparison_price}；不符 {outputs['summary']['n_error']}、需確認 {outputs['summary']['n_warn']}")
    return rec


@app.get("/api/cases/{cid}/reset_preview")
def cases_reset_preview(cid: str):
    from app import cases as C
    r = C.reset_preview(cid)
    if r is None:
        raise HTTPException(404, "沒有這個案件")
    return r


@app.post("/api/cases/{cid}/reset")
def cases_reset(cid: str, request: Request):
    """重置案件：回到載入時的原始輸入，清掉裁決與產出，狀態回草稿。"""
    from app import cases as C
    before = C.reset_preview(cid)
    rec = C.reset_case(cid)
    if not rec:
        raise HTTPException(404, "沒有這個案件或沒有原始快照")
    AUD.log(cid, AUD.actor_from_headers(request.headers), "reset", f"回到 {before['original_at']} 的原始輸入；清除 {before['n_input_changes']} 處輸入變更、{before['n_decisions']} 筆裁決與產出")
    return rec


@app.get("/api/cases/{cid}/fill_report")
def cases_fill_report(cid: str):
    """填寫結果清單：每欄的值、狀態（資料／量測／推定／空白）、來源與說明；供一鍵產生後檢視。"""
    from app import cases as C
    from app.report.fill_report import build_fill_report
    rec = C.get_case(cid)
    if not rec:
        raise HTTPException(404, "沒有這個案件")
    reg, ind = _rulesets(rec["data"]["case"])
    return build_fill_report(rec, reg, ind)


class PolishPayload(BaseModel):
    targets: list[str] = ["range_desc", "notes"]    # range_desc：區段範圍描述；notes：表4／表5 備註三欄


@app.get("/api/cases/{cid}/notes")
def cases_notes(cid: str):
    """書表備註模板句（比準地、各比較標的、全案）與區段範圍描述；人工填的 case.notes 優先。"""
    from app import cases as C
    from app.report.notes import build_notes
    rec = C.get_case(cid)
    if not rec:
        raise HTTPException(404, "沒有這個案件")
    from app.market.index import fill_date_adjustments
    data = copy.deepcopy(rec["data"])
    fill_date_adjustments(data)                                   # 預覽用：缺期日調整率者先依指數補（不存檔，重新產生書表時才存）
    reg, ind = _rulesets(data["case"])
    sid = data["subject_parcel"].get("section_id")
    return {"notes": build_notes(data, run_case(reg, ind, data)), "manual": data["case"].get("notes") or {},
            "range_desc": (data["sections"].get(sid) or {}).get("range_desc") or "", "section_id": sid}


@app.post("/api/cases/{cid}/polish_text")
def cases_polish_text(cid: str, payload: PolishPayload, request: Request):
    """語言模型潤飾文字欄位（區段範圍描述、備註），數字與條號不得改變；通過者寫回案件（case.notes／section.range_desc），原文保留。"""
    from app import cases as C
    from app.llm import LLMNotConfigured, get_provider
    from app.report.notes import build_notes, polish_text
    rec = C.get_case(cid)
    if not rec:
        raise HTTPException(404, "沒有這個案件")
    try:
        provider = get_provider()
        if not provider.status().get("configured", True):
            raise LLMNotConfigured("語言模型未設定")
    except LLMNotConfigured as e:
        raise HTTPException(422, f"語言模型未設定（{e}），無法潤飾；書表使用模板句") from e
    data = copy.deepcopy(rec["data"])
    reg, ind = _rulesets(data["case"])
    results: dict[str, Any] = {}
    if "range_desc" in payload.targets:
        sid = data["subject_parcel"].get("section_id")
        sec = data["sections"].get(sid) or {}
        r = polish_text(sec.get("range_desc") or "", provider)
        results["range_desc"] = r
        if r["status"] == "ok":
            sec.setdefault("range_desc_template", sec.get("range_desc"))
            sec["range_desc"] = r["text"]
    if "notes" in payload.targets:
        notes = build_notes(data, run_case(reg, ind, data))
        manual = data["case"].setdefault("notes", {})
        for key in ("subject", "case"):
            r = polish_text(notes[key], provider)
            results[f"notes.{key}"] = r
            if r["status"] == "ok":
                manual[key] = r["text"]
        manual.setdefault("comparables", {})
        for no, txt in notes["comparables"].items():
            r = polish_text(txt, provider)
            results[f"notes.comparables.{no}"] = r
            if r["status"] == "ok":
                manual["comparables"][no] = r["text"]
    rec = C.save_case(data, cid=cid, submitted_table5=rec.get("submitted_table5"), submitted_table4=rec.get("submitted_table4"))
    n_ok = sum(1 for r in results.values() if r["status"] == "ok")
    AUD.log(cid, AUD.actor_from_headers(request.headers), "polish", f"語言模型潤飾 {n_ok}/{len(results)} 段通過守門（{getattr(provider, 'name', '')}）")
    return {"rec": rec, "results": results}


_TODO_CACHE: dict[str, tuple[str, dict]] = {}


@app.get("/api/cases/{cid}/todo")
def cases_todo(cid: str):
    """待辦清單（同一案件、同一版本輸入與裁決只算一次；重新產生或存檔會換鍵）。"""
    from app import cases as C
    rec0 = C.get_case(cid)
    if not rec0:
        raise HTTPException(404, "沒有這個案件")
    key = f"{rec0.get('updated_at')}|{(rec0.get('outputs') or {}).get('generated_at')}|{json.dumps(rec0.get('decisions') or {}, sort_keys=True)}|{rec0.get('status')}"
    hit = _TODO_CACHE.get(cid)
    if hit and hit[0] == key:
        return hit[1]
    out = _cases_todo(cid)
    _TODO_CACHE[cid] = (key, out)
    return out


def _cases_todo(cid: str):
    """這一案還缺什麼：每項帶數量與要去的頁面（案件列的待辦清單用）。"""
    from app import cases as C
    from app.report.fill_report import build_fill_report
    from app.report.opinion import finding_key
    rec = C.get_case(cid)
    if not rec:
        raise HTTPException(404, "沒有這個案件")
    data = rec["data"]
    reg, ind = _rulesets(data["case"])
    review = bool(rec.get("submitted_table4") or rec.get("submitted_table5"))
    q = f"case={cid}"
    items: list[dict[str, Any]] = []
    stale = not rec.get("outputs") or rec["outputs"].get("input_hash") != rec.get("input_hash")
    rep = build_fill_report(rec, reg, ind)
    n_comp = len(data.get("comparables") or [])
    if not review and not data["subject_parcel"].get("parcel_id"):
        items.append({"key": "parcel", "label": "比準地地號未填", "count": None, "href": f"/input?tab=case&{q}", "level": "error"})
    if not review and data["subject_parcel"].get("geometry") is None and data["subject_parcel"].get("parcel_id"):
        items.append({"key": "geometry", "label": "比準地無界線：匯入地籍圖或點圖", "count": None, "href": f"/input?tab=case&{q}", "level": "error"})
    if n_comp == 0:
        items.append({"key": "comps", "label": "比較標的 0 件", "count": None, "href": f"/input?tab=parcels&{q}", "level": "error"})
    n_sub_blank = sum(1 for r in rep["subject"] if r["status"] == "空白")
    if n_sub_blank:
        items.append({"key": "subject_blank", "label": "比準地空白欄", "count": n_sub_blank, "href": f"/input?tab=parcels&{q}", "level": "warn"})
    n_sec_blank = sum(1 for r in rep["section"] if r["status"] == "空白")
    if n_sec_blank:
        items.append({"key": "section_blank", "label": "勘查表需人工填載", "count": n_sec_blank, "href": f"/table1?{q}", "level": "warn"})
    n_inferred = sum(1 for r in rep["head"] + rep["subject"] + rep["section"] if r["status"] == "推定")
    if n_inferred and not review:
        items.append({"key": "inferred", "label": "推定值待確認", "count": n_inferred, "href": f"/input?tab=case&report=1&{q}", "level": "info"})
    missing = (rec.get("extraction") or {}).get("missing_fields") or []
    if review and missing:
        items.append({"key": "missing", "label": "送審書表抽取缺漏欄位", "count": len(missing), "href": f"/input?tab=parcels&{q}", "level": "warn"})
    if stale:
        items.append({"key": "stale", "label": "產出已過期，請重新產生書表", "count": None, "href": f"/sheets?{q}", "level": "warn"})
    else:
        try:
            result = run_case(reg, ind, data)
            findings = collect_findings(reg, ind, data, result, rec.get("submitted_table5"), rec["submitted_table4"])
            decisions = rec.get("decisions") or {}
            undecided = [f for f in findings if f["severity"] == "error" and (decisions.get(finding_key(f) or f"loc:{f.get('table')}:{f.get('location')}") or {}).get("decision") not in ("accept", "reject")]
            gaps = [f for f in undecided if f.get("kind") == "gap"]
            if gaps and not review:
                items.append({"key": "gaps", "label": "資料缺口待補", "count": len(gaps), "href": f"/input?tab=parcels&{q}", "level": "error"})
            elif undecided:
                items.append({"key": "undecided", "label": "不符項待承辦裁決", "count": len(undecided), "href": f"/review?{q}", "level": "error"})
            n_warn = sum(1 for f in findings if f["severity"] == "warn")
            if n_warn:
                items.append({"key": "warn", "label": "需確認事項", "count": n_warn, "href": f"/review?{q}", "level": "info"})
        except Exception:  # noqa: BLE001, S110 - 核算失敗不影響待辦清單其餘項目
            pass
    return {"mode": "review" if review else "generate", "items": items, "ready": not any(i["level"] == "error" for i in items) and not stale}


@app.post("/api/cases/{cid}/clear")
def cases_clear(cid: str, request: Request):
    """清空輸入重填：保留案件基本資料與區段編號（正式區段圖範圍也保留），其餘輸入、送審書表、裁決、產出全部清空。"""
    from app import cases as C
    rec = C.clear_case(cid)
    if not rec:
        raise HTTPException(404, "沒有這個案件")
    AUD.log(cid, AUD.actor_from_headers(request.headers), "clear", "清空宗地、比較標的、勘查表、送審書表、裁決與產出；保留案件基本資料與區段編號")
    return rec


def _case_workbook(rec: dict, *, figures: bool, appraiser: str = "", fill_date: str = ""):
    from app.output.xlsx import build_workbook
    data = rec["data"]
    appraiser = appraiser or data["case"].get("appraiser") or ""       # 簽章欄與填寫日期存在案件基本資料，全部輸出共用
    fill_date = fill_date or data["case"].get("fill_date") or ""
    reg, ind = _rulesets(data["case"])
    figs = None
    if figures:
        try:
            from app.maps.render import case_figures
            figs = case_figures(data, appraiser=appraiser, page=True)
        except ValueError:
            figs = None
    from app.report.notes import build_notes
    try:
        notes = build_notes(data, run_case(reg, ind, data))          # 備註欄模板句（人工填的 case.notes 優先）
    except Exception:  # noqa: BLE001 - 備註產生失敗不擋書表
        notes = build_notes(data, None)
    return build_workbook(data, meta={"appraiser": appraiser, "fill_date": fill_date, "notes": notes}, regional=reg, individual=ind, figures=figs), figs


@app.get("/api/cases/{cid}/sheets")
def cases_sheets(cid: str):
    """書表預覽：三張表的格線（與 Excel 同一份版面）＋三張圖說連結，依範本頁序。"""
    from app import cases as C
    from app.maps.render import MODE_TITLE
    from app.output.grid import workbook_grids
    rec = C.get_case(cid)
    if not rec:
        raise HTTPException(404, "沒有這個案件")
    wb, _ = _case_workbook(rec, figures=False)
    data = rec["data"]
    has_geom = bool(data["subject_parcel"].get("geometry") or any((s or {}).get("geometry") for s in (data.get("sections") or {}).values()))
    figures = [{"mode": m, "title": t, "url": f"/api/cases/{cid}/map.png?mode={m}&page=1"} for m, t in MODE_TITLE.items()] if has_geom else []
    return {"sheets": workbook_grids(wb), "figures": figures, "case_no": data["case"].get("case_no"), "stale": bool(C.list_cases() and next((x["stale"] for x in C.list_cases() if x["id"] == cid), True))}


@app.get("/api/cases/{cid}/sheets.pdf")
def cases_sheets_pdf(cid: str, request: Request, appraiser: str = "", fill_date: str = "", inline: bool = False):
    """完整書表 PDF：勘查表每個區段一頁（比準地區段在前）、表5、表4 各一頁＋三張圖說各一頁，照範本順序。inline=1 → 瀏覽器直接開啟（列印用），否則下載。"""
    from app import cases as C
    from app.output.grid import grids_to_pdf, workbook_grids
    rec = C.get_case(cid)
    if not rec:
        raise HTTPException(404, "沒有這個案件")
    wb, figs = _case_workbook(rec, figures=True, appraiser=appraiser, fill_date=fill_date)
    case = rec["data"]["case"]
    pdf = grids_to_pdf(workbook_grids(wb), figs, footer=f"案號 {case.get('case_no', '')}　估價基準日 {case.get('valuation_date', '')}　系統依查估辦法與作業手冊核算產出")
    fname = f"{case.get('case_no', 'case')}_查估書表.pdf"
    AUD.log(cid, AUD.actor_from_headers(request.headers), "print" if inline else "export", fname)
    return StreamingResponse(io.BytesIO(pdf), media_type="application/pdf",
                             headers={"Content-Disposition": f"{'inline' if inline else 'attachment'}; filename*=UTF-8''{urllib.parse.quote(fname)}"})


@app.get("/api/cases/{cid}/sheets.xlsx")
def cases_sheets_xlsx(cid: str, request: Request, appraiser: str = "", fill_date: str = ""):
    """完整書表 Excel：各區段勘查表、表5、表4 各一張工作表＋三張圖說工作表。"""
    from app import cases as C
    rec = C.get_case(cid)
    if not rec:
        raise HTTPException(404, "沒有這個案件")
    wb, figs = _case_workbook(rec, figures=True, appraiser=appraiser, fill_date=fill_date)
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    fname = f"{rec['data']['case'].get('case_no', 'case')}_查估書表.xlsx"
    AUD.log(cid, AUD.actor_from_headers(request.headers), "export", fname + ("（含圖說）" if figs else ""))
    return StreamingResponse(buf, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                             headers={"Content-Disposition": f"attachment; filename*=UTF-8''{urllib.parse.quote(fname)}"})


def _case_report(rec: dict, reviewer: str = "", reviewer_role: str = "", figs: list | None = None):
    """案件 → 審查意見書（含裁決、落款角色、圖說連結）。bundle 與 report.md 共用。"""
    from app.report.opinion import build_report
    data = rec["data"]
    case = data["case"]
    reg, ind = _rulesets(case)
    result = run_case(reg, ind, data)
    findings = collect_findings(reg, ind, data, result, rec.get("submitted_table5"), rec["submitted_table4"])
    computed = {"table5": {k: v.to_dict() for k, v in result["table5"].items()}, "table4": result["table4"].to_dict()}
    rv = (reviewer or "").strip()
    if rv and reviewer_role in AUD.ROLES:
        rv = f"{rv}（{AUD.ROLES[reviewer_role]}）"
    return build_report(case, findings, computed, subject_parcel_id=data["subject_parcel"].get("parcel_id", ""), reviewer=rv, decisions=rec.get("decisions") or {},
                        has_submitted=bool(rec.get("submitted_table4") or rec.get("submitted_table5")),
                        figures=[{"mode": m, "title": t, "url": f"圖說_{t}.png"} for m, t, _ in (figs or [])])


def _report_figs(rec: dict) -> list | None:
    """意見書附圖（三張圖說）；沒有幾何就不附。"""
    try:
        from app.maps.render import case_figures
        return case_figures(rec["data"], appraiser=rec["data"]["case"].get("appraiser", ""))
    except ValueError:
        return None


@app.get("/api/cases/{cid}/report.docx")
def cases_report_docx(cid: str, request: Request, reviewer: str = "", reviewer_role: str = ""):
    """審查意見書 Word 檔（含三張圖說附圖；落款用查詢參數的姓名與角色）。"""
    from app import cases as C
    from app.report.render import report_to_docx
    rec = C.get_case(cid)
    if not rec:
        raise HTTPException(404, "沒有這個案件")
    figs = _report_figs(rec)
    data = report_to_docx(_case_report(rec, reviewer, reviewer_role, figs), figs)
    fname = f"{rec['data']['case'].get('case_no', 'case')}_審查意見書.docx"
    AUD.log(cid, AUD.actor_from_headers(request.headers), "export", fname)
    return StreamingResponse(io.BytesIO(data), media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                             headers={"Content-Disposition": f"attachment; filename*=UTF-8''{urllib.parse.quote(fname)}"})


@app.get("/api/cases/{cid}/report.pdf")
def cases_report_pdf(cid: str, request: Request, reviewer: str = "", reviewer_role: str = ""):
    """審查意見書 PDF（含三張圖說附圖）。"""
    from app import cases as C
    from app.report.render import report_to_pdf
    rec = C.get_case(cid)
    if not rec:
        raise HTTPException(404, "沒有這個案件")
    figs = _report_figs(rec)
    data = report_to_pdf(_case_report(rec, reviewer, reviewer_role, figs), figs)
    fname = f"{rec['data']['case'].get('case_no', 'case')}_審查意見書.pdf"
    AUD.log(cid, AUD.actor_from_headers(request.headers), "export", fname)
    return StreamingResponse(io.BytesIO(data), media_type="application/pdf",
                             headers={"Content-Disposition": f"attachment; filename*=UTF-8''{urllib.parse.quote(fname)}"})


@app.get("/api/cases/{cid}/bundle.zip")
def cases_bundle(cid: str, request: Request, appraiser: str = "", reviewer: str = "", reviewer_role: str = ""):
    """下載全部：查估書表 Excel、PDF、審查意見書 Word 與 PDF、三張圖說 PNG，一個 zip。"""
    import zipfile

    from app import cases as C
    from app.output.grid import grids_to_pdf, workbook_grids
    from app.report.render import report_to_docx, report_to_pdf
    rec = C.get_case(cid)
    if not rec:
        raise HTTPException(404, "沒有這個案件")
    data = rec["data"]
    case = data["case"]
    wb, figs = _case_workbook(rec, figures=True, appraiser=appraiser)
    xbuf = io.BytesIO()
    wb.save(xbuf)
    pdf = grids_to_pdf(workbook_grids(wb), figs, footer=f"案號 {case.get('case_no', '')}　估價基準日 {case.get('valuation_date', '')}　系統依查估辦法與作業手冊核算產出")
    rep = _case_report(rec, reviewer, reviewer_role, figs)
    no = case.get("case_no", "case")
    zbuf = io.BytesIO()
    with zipfile.ZipFile(zbuf, "w", zipfile.ZIP_DEFLATED) as z:
        z.writestr(f"{no}_查估書表.xlsx", xbuf.getvalue())
        z.writestr(f"{no}_查估書表.pdf", pdf)
        z.writestr(f"{no}_審查意見書.docx", report_to_docx(rep, figs))
        z.writestr(f"{no}_審查意見書.pdf", report_to_pdf(rep, figs))
        for _m, t, png in figs or []:
            z.writestr(f"圖說_{t}.png", png)
    zbuf.seek(0)
    fname = f"{no}_全部輸出.zip"
    AUD.log(cid, AUD.actor_from_headers(request.headers), "export", fname)
    return StreamingResponse(zbuf, media_type="application/zip", headers={"Content-Disposition": f"attachment; filename*=UTF-8''{urllib.parse.quote(fname)}"})



def _official_inputs(rec: dict, appraiser: str = "", fill_date: str = ""):
    """地政局正式範本填值用：引擎結果＋備註＋簽章欄。"""
    from app.report.notes import build_notes
    data = rec["data"]
    reg, ind = _rulesets(data["case"])
    result = run_case(reg, ind, data)
    try:
        notes = build_notes(data, result)
    except Exception:  # noqa: BLE001 - 備註產生失敗不擋書表
        notes = build_notes(data, None)
    meta = {"appraiser": appraiser or data["case"].get("appraiser") or "", "fill_date": fill_date or data["case"].get("fill_date") or "", "notes": notes}
    return data, result, reg, ind, meta


@app.get("/api/cases/{cid}/official.zip")
def cases_official_zip(cid: str, request: Request, appraiser: str = "", fill_date: str = ""):
    """地政局正式範本三份 Excel（表3 每區段一張工作表、表5-1、表4）打包。"""
    from app import cases as C
    from app.output.official_xlsx import official_zip
    rec = C.get_case(cid)
    if not rec:
        raise HTTPException(404, "沒有這個案件")
    try:
        blob = official_zip(*_official_inputs(rec, appraiser, fill_date))
    except InputError as e:
        raise HTTPException(422, str(e)) from e
    fname = f"{rec['data']['case'].get('case_no', 'case')}_正式範本書表.zip"
    AUD.log(cid, AUD.actor_from_headers(request.headers), "export", fname)
    return StreamingResponse(io.BytesIO(blob), media_type="application/zip",
                             headers={"Content-Disposition": f"attachment; filename*=UTF-8''{urllib.parse.quote(fname)}"})


@app.get("/api/cases/{cid}/official/{key}.xlsx")
def cases_official_xlsx(cid: str, key: str, request: Request, appraiser: str = "", fill_date: str = ""):
    """地政局正式範本單張：key = t3（勘查表）| t5（區域因素分析明細表）| t4（比較法調查估價表）。"""
    from app import cases as C
    from app.output.official_xlsx import TEMPLATES, fill_table3, fill_table4, fill_table5, workbook_bytes
    if key not in TEMPLATES:
        raise HTTPException(404, "key 須為 t3、t5 或 t4")
    rec = C.get_case(cid)
    if not rec:
        raise HTTPException(404, "沒有這個案件")
    try:
        data, result, reg, ind, meta = _official_inputs(rec, appraiser, fill_date)
        wb = fill_table3(data, reg, meta) if key == "t3" else fill_table5(data, result["table5"], reg, meta) if key == "t5" else fill_table4(data, result["table4"], ind, meta)
    except InputError as e:
        raise HTTPException(422, str(e)) from e
    fname = f"{data['case'].get('case_no', 'case')}_{TEMPLATES[key][2]}.xlsx"
    AUD.log(cid, AUD.actor_from_headers(request.headers), "export", fname)
    return StreamingResponse(io.BytesIO(workbook_bytes(wb)), media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                             headers={"Content-Disposition": f"attachment; filename*=UTF-8''{urllib.parse.quote(fname)}"})

@app.get("/api/cases/{cid}/parcels.xlsx")
def cases_parcels_xlsx(cid: str, request: Request):
    """匯出宗地個別因素清冊 xlsx（表7 版面；填好可再匯入）。"""
    from app import cases as C
    from app.adapters.excel_parcels import write_parcels_xlsx
    rec = C.get_case(cid)
    if not rec:
        raise HTTPException(404, "沒有這個案件")
    d = rec["data"]
    buf = io.BytesIO()
    write_parcels_xlsx([d["subject_parcel"], *d.get("comparables", [])], buf, case_no=d["case"].get("case_no", ""))
    buf.seek(0)
    fname = f"{d['case'].get('case_no', 'case')}_宗地個別因素清冊.xlsx"
    AUD.log(cid, AUD.actor_from_headers(request.headers), "export", fname)
    return StreamingResponse(buf, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                             headers={"Content-Disposition": f"attachment; filename*=UTF-8''{urllib.parse.quote(fname)}"})


@app.get("/api/cases/{cid}/comparables.xlsx")
def cases_comparables_xlsx(cid: str, request: Request):
    """匯出買賣實例 xlsx（填好可再匯入）。"""
    from app import cases as C
    from app.adapters.excel_comparables import write_comparables_xlsx
    rec = C.get_case(cid)
    if not rec:
        raise HTTPException(404, "沒有這個案件")
    d = rec["data"]
    buf = io.BytesIO()
    write_comparables_xlsx(d.get("comparables", []), buf)
    buf.seek(0)
    fname = f"{d['case'].get('case_no', 'case')}_買賣實例.xlsx"
    AUD.log(cid, AUD.actor_from_headers(request.headers), "export", fname)
    return StreamingResponse(buf, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                             headers={"Content-Disposition": f"attachment; filename*=UTF-8''{urllib.parse.quote(fname)}"})


# ------------------------------------------------------------------ 輸入檔（一個案件、多份檔）


def _geo_kind(feats: list[dict]) -> str:
    props = (feats[0].get("properties") or {}) if feats else {}
    if _pick_field(props, _LOT_KEYS):
        return "cadastre"
    if _pick_field(props, _SECID_KEYS):
        return "section_map"
    raise HTTPException(422, "圖檔屬性裡找不到地號欄位（地籍圖）或區段編號欄位（地價區段圖）")


def _detect_input_kind(filename: str, content: bytes) -> str:
    """輸入檔種類：書表 PDF／清冊／實例／基準表（PDF、CSV、xlsx、JSON）／地籍圖／地價區段圖（GeoJSON、KML、GML、SHP zip）。"""
    name = (filename or "").lower()
    if name.endswith(".json"):
        try:
            d = json.loads(content.decode("utf-8"))
        except Exception as e:
            raise HTTPException(422, f"JSON 無法解析：{e}") from e
        if isinstance(d, dict) and d.get("type") == "FeatureCollection":
            return _geo_kind(d.get("features") or [])
        if isinstance(d, dict) and "scope" in d and "rules" in d:
            return "rules_json"
        raise HTTPException(422, "JSON 不是地籍圖／區段圖（GeoJSON），也不是基準表 JSON")
    if name.endswith((".geojson", ".kml", ".gml", ".xml", ".zip")):
        try:
            feats = _read_cadastre_upload(content, name, 3826)
        except Exception as e:
            raise HTTPException(422, f"圖檔解析失敗：{e}") from e
        return _geo_kind(feats)
    try:
        return detect_kind(filename, content)
    except HTTPException:
        raise
    except Exception as e:
        raise HTTPException(422, f"無法讀取檔案：{e}") from e


_INPUT_ORDER = {"pdf_forms": 0, "rules_table": 1, "rules_json": 1, "parcels": 2, "comparables": 3, "cadastre": 4, "section_map": 5}


def _apply_one_input(cid: str, content: bytes, filename: str, *, kind: str = "auto", use_vision: str = "auto",
                     actor: dict | None = None, log: bool = True) -> dict:
    """一份輸入檔併入案件：辨識種類 → 依 app/inputs.py 規則合併 → 原檔存到 data/cases/<id>/inputs/ → 記在 rec["inputs"] 與操作紀錄。回傳輸入檔紀錄。"""
    from app import cases as C
    from app.inputs import (
        KIND_LABELS,
        aggregate_extraction,
        list_summary,
        merge_parcel_list,
        merge_pdf_forms,
        pdf_summary,
    )
    rec = C.get_case(cid)
    if not rec:
        raise HTTPException(404, "沒有這個案件")
    if not content:
        raise HTTPException(400, f"空檔案：{filename}")
    sha = hashlib.sha1(content).hexdigest()
    dup = next((i for i in rec.get("inputs") or [] if i.get("sha1") == sha), None)
    if dup:
        return {"skipped": True, "filename": filename, "kind": dup.get("kind"), "kind_label": dup.get("kind_label"),
                "summary": f"同一份檔已加入（{dup.get('filename')}，{str(dup.get('at') or '')[:16]}），未重複併入"}
    k = _detect_input_kind(filename, content) if kind == "auto" else kind
    if not rec.get("inputs") and not rec.get("inputs_base"):          # 第一份輸入檔併入前的快照：移除輸入檔時從這裡重新併入其餘檔案
        rec["inputs_base"] = {"data": copy.deepcopy(rec["data"]), "submitted_table5": copy.deepcopy(rec.get("submitted_table5")),
                              "submitted_table4": copy.deepcopy(rec.get("submitted_table4")), "at": C._now()}
    iid = uuid.uuid4().hex[:8]
    kind_out = "rules_table" if k == "rules_json" else k
    entry: dict[str, Any] = {"id": iid, "filename": filename, "kind": kind_out, "kind_label": KIND_LABELS.get(kind_out, kind_out), "size": len(content), "sha1": sha,
                             "at": C._now(), "actor": AUD.actor_label(actor), "pages": [], "summary": "", "missing": [], "warnings": [],
                             "conflicts": [], "overrides": [], "filled": [], "matched": [], "unmatched": [], "notes": []}
    status_patch: str | None = None
    try:
        if k == "pdf_forms":
            res = _read_pdf_forms_cached(content, filename, use_vision)
            data, t5, t4, rep = merge_pdf_forms(rec["data"], rec.get("submitted_table5"), rec.get("submitted_table4"), res.data)
            subj = data.get("subject_parcel") or {}
            if subj.get("geometry") is None and subj.get("parcel_id"):
                rep.notes.extend(_attach_geometry(data))                       # 補界線與區段範圍，圖說才畫得出來
            rec = C.save_case(data, cid=cid, submitted_table5=t5, submitted_table4=t4)
            entry.update(pages=res.data.get("pages") or [], missing=list(res.missing_fields), warnings=list(res.warnings), confidence=dict(res.confidence), **rep.to_dict())
            entry["summary"] = pdf_summary(entry["pages"], rep, len(res.data.get("comparables") or []))
            if (t5 or t4) and rec.get("status") == "draft":
                status_patch = "reviewing"
        elif k in ("parcels", "comparables"):
            if k == "parcels":
                from app.adapters.excel_parcels import read_parcels
                res = read_parcels(content)
                items = res.data.get("parcels")
            else:
                from app.adapters.excel_comparables import read_comparables
                res = read_comparables(content)
                items = res.data.get("comparables")
            data = copy.deepcopy(rec["data"])
            rep = merge_parcel_list(data, items or [], k)
            rec = C.save_case(data, cid=cid, submitted_table5=rec.get("submitted_table5"), submitted_table4=rec.get("submitted_table4"))
            entry.update(missing=list(res.missing_fields), warnings=list(res.warnings), **rep.to_dict())
            entry["summary"] = list_summary(rep)
        elif k in ("rules_table", "rules_json"):
            if k == "rules_json":
                d = json.loads(content.decode("utf-8"))
                rulesets = {d["scope"]: d}
                warnings: list[str] = []
            else:
                from app.adapters.rules_table import read_rules_table
                res = read_rules_table(content, filename=filename)
                rulesets = res.data.get("rulesets") or {}
                warnings = list(res.warnings)
            if not rulesets:
                raise HTTPException(422, "檔案裡沒有可匯入的基準表")
            data = copy.deepcopy(rec["data"])
            ids = []
            for scope, rs in rulesets.items():
                rid = _store_ruleset(copy.deepcopy(rs))
                data["case"].setdefault("rulesets", {})[scope] = rid
                if rs.get("land_use"):
                    data["case"]["land_use"] = rs["land_use"]
                ids.append(rid)
            rec = C.save_case(data, cid=cid, submitted_table5=rec.get("submitted_table5"), submitted_table4=rec.get("submitted_table4"))
            entry.update(warnings=warnings, rulesets=ids)
            entry["summary"] = f"已匯入並套用至本案：{'、'.join(ids)}"
        elif k == "cadastre":
            r = _apply_cadastre(rec, content, filename)
            rec = r["record"]
            entry.update(matched=r["matched"], unmatched=r["unmatched"])
            entry["summary"] = f"地籍圖 {r['n']} 筆；對到真實界線 {len(r['matched'])} 筆宗地" + (f"；對不到 {'、'.join(r['unmatched'])}" if r["unmatched"] else "")
        elif k == "section_map":
            r = _apply_section_map(rec, content, filename)
            rec = r["record"]
            entry.update(matched=r["matched"], unmatched=r["unmatched"])
            entry["summary"] = f"地價區段圖 {r['n']} 區段；對到本案區段 {'、'.join(r['matched']) or '無'}" + (f"；對不到 {'、'.join(r['unmatched'])}" if r["unmatched"] else "")
        else:
            raise HTTPException(422, f"不支援的輸入檔種類：{k}")
    except HTTPException:
        raise
    except Exception as e:
        raise HTTPException(422, f"檔案解析失敗（{KIND_LABELS.get(kind_out, k)}）：{e}") from e
    d = C.inputs_dir(cid)
    d.mkdir(parents=True, exist_ok=True)
    safe = re.sub(r'[\\/:*?"<>|]', "_", filename or "file")
    (d / f"{iid}_{safe}").write_bytes(content)
    entry["path"] = f"{iid}_{safe}"
    inputs = [*(rec.get("inputs") or []), entry]
    C.patch_case(cid, inputs=inputs, extraction=aggregate_extraction(inputs), status=status_patch)
    if log:
        AUD.log(cid, actor, "input", f"{filename}（{entry['kind_label']}）：{entry['summary']}")
    return entry


def _remove_input(cid: str, iid: str, actor: dict | None) -> dict:
    """移除一份輸入檔：回到第一份輸入檔併入前的快照，再依序重新併入其餘檔案（結果確定，不做反向刪欄位）。"""
    from app import cases as C
    from app.inputs import aggregate_extraction
    rec = C.get_case(cid)
    if not rec:
        raise HTTPException(404, "沒有這個案件")
    inputs = list(rec.get("inputs") or [])
    entry = next((i for i in inputs if i.get("id") == iid), None)
    if not entry:
        raise HTTPException(404, "沒有這份輸入檔")
    if entry.get("legacy"):
        raise HTTPException(422, "這筆是舊版建案留下的紀錄，沒有原檔可重新併入；要拿掉送審書表請用「清空重填」後重新加入檔案")
    base = rec.get("inputs_base")
    if not base:
        raise HTTPException(422, "找不到輸入檔併入前的快照，無法移除重併")
    remaining = []
    for i in inputs:
        if i.get("id") == iid or i.get("legacy"):
            continue
        p = C.inputs_dir(cid) / (i.get("path") or "")
        if not i.get("path") or not p.exists():
            continue
        remaining.append((i, p.read_bytes()))
    rec = C.save_case(copy.deepcopy(base["data"]), cid=cid, submitted_table5=copy.deepcopy(base.get("submitted_table5")), submitted_table4=copy.deepcopy(base.get("submitted_table4")))
    import shutil
    shutil.rmtree(C.inputs_dir(cid), ignore_errors=True)
    C.patch_case(cid, inputs=[], extraction=aggregate_extraction([]) or {})
    C.get_case(cid)["extraction"] = None
    replayed, failed = [], []
    for i, content in remaining:
        try:
            replayed.append(_apply_one_input(cid, content, i.get("filename") or "", kind=i.get("kind") or "auto", use_vision="auto", actor=actor, log=False))
        except HTTPException as e:
            failed.append(f"{i.get('filename')}：{e.detail}")
    AUD.log(cid, actor, "input_remove", f"{entry.get('filename')}（{entry.get('kind_label')}）；已重新併入其餘 {len(replayed)} 份" + (f"；失敗：{'；'.join(failed)}" if failed else ""))
    return {"record": C.get_case(cid), "removed": entry, "replayed": replayed, "failed": failed}


@app.post("/api/cases/{cid}/inputs")
async def cases_inputs_add(cid: str, request: Request, files: list[UploadFile] = File(...), kind: str = Form("auto"), use_vision: str = Form("auto")):   # noqa: B008
    """一次加入多份輸入檔到同一案件（書表 PDF、清冊、實例、基準表、地籍圖、區段圖）；依種類順序併入：書表 → 基準表 → 清冊 → 實例 → 地籍圖 → 區段圖。每份各回一筆結果，失敗的不影響其他份。"""
    from app import cases as C
    if not C.get_case(cid):
        raise HTTPException(404, "沒有這個案件")
    actor = AUD.actor_from_headers(request.headers)
    items = []
    for f in files:
        content = await f.read()
        try:
            k = _detect_input_kind(f.filename or "", content) if kind == "auto" else kind
        except HTTPException as e:
            items.append((99, f.filename or "", content, None, e.detail))
            continue
        items.append((_INPUT_ORDER.get(k, 50), f.filename or "", content, k, None))
    results = []
    for _o, fn, content, k, err in sorted(items, key=lambda x: x[0]):
        if err:
            results.append({"filename": fn, "error": err})
            continue
        try:
            results.append(_apply_one_input(cid, content, fn, kind=k or "auto", use_vision=use_vision, actor=actor))
        except HTTPException as e:
            results.append({"filename": fn, "kind": k, "error": e.detail})
    rec = C.get_case(cid)
    return {"record": rec, "results": results, "inputs_summary": C.inputs_status(rec)}


@app.post("/api/cases/from_inputs")
async def cases_from_inputs(request: Request, files: list[UploadFile] = File(...), name: str | None = Form(None), case_no: str = Form(""),   # noqa: B008
                            valuation_date: str = Form(""), district: str = Form("新北市金山區"), land_use: str = Form(""), section_id: str = Form(""),
                            use_vision: str = Form("auto")):
    """從多份輸入檔建立一個案件：先建空案（案號等以表單值起頭，書表 PDF 會補上），再依種類順序把每份檔併入。沒有書表 PDF 時必須給案號與估價基準日。"""
    from app import cases as C
    actor = AUD.actor_from_headers(request.headers)
    items = []
    for f in files:
        content = await f.read()
        try:
            k = _detect_input_kind(f.filename or "", content)
        except HTTPException as e:
            items.append((99, f.filename or "", content, None, e.detail))
            continue
        items.append((_INPUT_ORDER.get(k, 50), f.filename or "", content, k, None))
    items.sort(key=lambda x: x[0])
    has_pdf = any(k == "pdf_forms" for _o, _fn, _c, k, _e in items)
    if not has_pdf and not (case_no.strip() and valuation_date.strip()):
        raise HTTPException(422, "沒有送審書表 PDF 時，請填案號與估價基準日再建立")
    sid = section_id.strip() or "P001-00"
    if not case_no.strip() and has_pdf:                                   # 案件 id 由案號衍生：先從第一份書表讀案號（結果有快取，併入時不重跑）
        for _o, fn, content, k, _e in items:
            if k == "pdf_forms":
                try:
                    pc = _read_pdf_forms_cached(content, fn, use_vision).data.get("case") or {}
                    case_no = pc.get("case_no") or ""
                    valuation_date = valuation_date or pc.get("valuation_date") or ""
                except Exception:  # noqa: BLE001, S110 - 讀不到就用空案號，併入時再報錯
                    pass
                break
    data = {"case": {"case_no": case_no.strip(), "valuation_date": valuation_date.strip(), "district": district.strip(), "land_use": land_use.strip() or None},
            "sections": {sid: {"section_id": sid, "range_desc": "", "survey": C.blank_survey()}},
            "subject_parcel": {"parcel_id": "", "section_id": sid, "nuisance": None}, "comparables": []}
    rec = C.save_case(data, name=name or (case_no.strip() or "新案件"), origin="inputs")
    cid = rec["id"]
    AUD.log(cid, actor, "create", f"{rec.get('name')}（來源 檔案匯入，{len(items)} 份輸入檔）")
    results = []
    for _o, fn, content, k, err in items:
        if err:
            results.append({"filename": fn, "error": err})
            continue
        try:
            results.append(_apply_one_input(cid, content, fn, kind=k or "auto", use_vision=use_vision, actor=actor))
        except HTTPException as e:
            results.append({"filename": fn, "kind": k, "error": e.detail})
    rec = C.get_case(cid)
    c = rec["data"]["case"]
    if not c.get("land_use"):
        c["land_use"] = "商業用地"
    if not c.get("rulesets"):
        c["rulesets"] = C.pick_rulesets(c.get("land_use"))
    rec = C.save_case(rec["data"], cid=cid, submitted_table5=rec.get("submitted_table5"), submitted_table4=rec.get("submitted_table4"))
    if not name:
        cn = c.get("case_no") or (rec["inputs"][0]["filename"] if rec.get("inputs") else cid)
        C.patch_case(cid, name=f"{cn}（送審書表）" if has_pdf else f"{cn} {c.get('district') or ''} {next(iter(rec['data'].get('sections') or {}), '')}".strip())
    rec = C.get_case(cid)
    return {"case": rec, "results": results, "inputs_summary": C.inputs_status(rec)}


@app.delete("/api/cases/{cid}/inputs/{iid}")
def cases_inputs_remove(cid: str, iid: str, request: Request):
    return _remove_input(cid, iid, AUD.actor_from_headers(request.headers))


@app.get("/api/cases/{cid}/inputs/{iid}/file")
def cases_inputs_file(cid: str, iid: str):
    """下載原始輸入檔。"""
    from app import cases as C
    rec = C.get_case(cid)
    if not rec:
        raise HTTPException(404, "沒有這個案件")
    entry = next((i for i in rec.get("inputs") or [] if i.get("id") == iid), None)
    if not entry or not entry.get("path"):
        raise HTTPException(404, "沒有這份輸入檔的原檔")
    p = C.inputs_dir(cid) / entry["path"]
    if not p.exists():
        raise HTTPException(404, "原檔已不存在")
    return Response(content=p.read_bytes(), media_type="application/octet-stream",
                    headers={"Content-Disposition": f"attachment; filename*=UTF-8''{urllib.parse.quote(entry.get('filename') or p.name)}"})


@app.post("/api/cases/{cid}/import")
async def cases_import(cid: str, request: Request, file: UploadFile = File(...), kind: str = Form("auto")):   # noqa: B008
    """單檔匯入（宗地與實例分頁、補上送審書表用）：走同一套輸入檔流程，回傳格式相容舊版（matched／unmatched／changes）。"""
    from app import cases as C
    if not C.get_case(cid):
        raise HTTPException(404, "沒有這個案件")
    content = await file.read()
    entry = _apply_one_input(cid, content, file.filename or "", kind=kind, use_vision="auto", actor=AUD.actor_from_headers(request.headers))
    rec = C.get_case(cid)
    if entry.get("skipped"):
        return {"record": rec, "case": rec, "kind": entry.get("kind"), "matched": [], "unmatched": [], "changes": 0, "warnings": [entry["summary"]], "missing_fields": [], "input": entry}
    return {"record": rec, "case": rec, "kind": entry.get("kind"), "matched": entry.get("matched") or [], "unmatched": entry.get("unmatched") or [],
            "changes": len(entry.get("filled") or []) + len(entry.get("overrides") or []), "warnings": entry.get("warnings") or [],
            "missing_fields": entry.get("missing") or [], "input": entry}


class NewCasePayload(BaseModel):
    case_no: str
    valuation_date: str
    district: str = "新北市金山區"
    land_use: str = "商業用地"
    section_id: str = ""            # 空白 → 系統暫編 P001-00（區段號由查估單位編定，需確認）
    range_desc: str = ""
    subject_parcel_id: str = ""
    name: str | None = None


@app.post("/api/cases/new")
def cases_new(payload: NewCasePayload, request: Request):
    """從零建案（沒有送審書表時）：只需案號、估價基準日、鄉鎮市區、用地別、區段編號。"""
    from app import cases as C
    if not payload.case_no.strip() or not payload.valuation_date.strip():
        raise HTTPException(422, "案號與估價基準日為必填")
    if not payload.section_id.strip() and not payload.subject_parcel_id.strip():
        raise HTTPException(422, "區段編號與比準地地號至少填一項")
    rec = C.new_case(case_no=payload.case_no.strip(), valuation_date=payload.valuation_date.strip(), district=payload.district.strip(),
                     land_use=payload.land_use, section_id=payload.section_id.strip() or "P001-00", range_desc=payload.range_desc, subject_parcel_id=payload.subject_parcel_id.strip(), name=payload.name)
    AUD.log(rec["id"], AUD.actor_from_headers(request.headers), "create", f"{rec['name']}（從零建案）")
    return rec


@app.get("/api/tiles/{z}/{x}/{y}")
def tiles(z: int, x: int, y: int, layer: str = "EMAP"):
    """瓦片代理（layer=EMAP 電子地圖｜LANDSECT 段籍圖）：先讀 data/tiles 快取（scripts/prefetch_tiles.py 預抓），沒有才向國土測繪中心取；離線時回 204。"""
    from fastapi.responses import Response

    from app.maps.render import tile_bytes
    data = tile_bytes(z, x, y, timeout=4.0, layer=layer)
    if data is None:
        return Response(status_code=204)
    return Response(content=data, media_type="image/png", headers={"Cache-Control": "public, max-age=86400"})


@app.post("/api/cases/{cid}/archive")
def cases_archive(cid: str, request: Request, undo: bool = False):
    """封存（或復原）案件：清單預設不顯示封存案件；資料與紀錄都保留。"""
    from app import cases as C
    rec = C.set_archived(cid, not undo)
    if not rec:
        raise HTTPException(404, "沒有這個案件")
    AUD.log(cid, AUD.actor_from_headers(request.headers), "unarchive" if undo else "archive", rec.get("name") or "")
    return rec


CADASTRE_DIR = Path(__file__).resolve().parents[2] / "data" / "cadastre" / "cases"
_SEC_KEYS = ("段名", "段", "SECTION", "SECT", "sect_name", "section", "段小段", "地段")
_LOT_KEYS = ("地號", "LANDNO", "LOTNO", "lot", "lot_no", "LAND_NO", "地號碼", "PARCEL", "parcel_no")


def _pick_field(props: dict, cands: tuple[str, ...]) -> str | None:
    keys = {k.lower(): k for k in props}
    for c in cands:
        if c.lower() in keys:
            return keys[c.lower()]
    return None


def _read_cadastre_upload(content: bytes, name: str, src_epsg: int) -> list[dict]:
    """GeoJSON 或含 .shp 的 zip → WGS84 GeoJSON features（屬性原樣）。"""
    import tempfile
    import zipfile
    if name.endswith((".kml", ".gml", ".xml")):
        from app.spatial.cadastre import parse_kml_gml
        return parse_kml_gml(content)
    if not name.endswith(".zip"):
        return json.loads(content.decode("utf-8")).get("features") or []
    import shapefile
    from pyproj import Transformer
    from shapely.geometry import mapping as _mapping
    from shapely.geometry import shape as _shape
    from shapely.ops import transform as _transform
    feats: list[dict] = []
    with tempfile.TemporaryDirectory() as td:
        zp = Path(td) / "in.zip"
        zp.write_bytes(content)
        with zipfile.ZipFile(zp) as z:
            z.extractall(td)
        shp = next(iter(Path(td).rglob("*.shp")), None)
        if not shp:
            raise ValueError("zip 裡沒有 .shp")
        tr = Transformer.from_crs(f"EPSG:{src_epsg}", "EPSG:4326", always_xy=True)
        with shapefile.Reader(str(shp), encoding="utf-8") as sf:
            for sr in sf.shapeRecords():
                g = _shape(sr.shape.__geo_interface__)
                if src_epsg != 4326:
                    g = _transform(tr.transform, g)
                feats.append({"type": "Feature", "geometry": _mapping(g), "properties": sr.record.as_dict()})
    return feats


def _apply_cadastre(rec: dict, content: bytes, filename: str, section_field: str = "", lot_field: str = "", src_epsg: int = 3826) -> dict:
    """地籍圖檔併入案件：存成本案地籍圖層，並用它補比準地與比較標的的真實幾何。回傳 {record, n, section_field, lot_field, matched, unmatched}。"""
    from app import cases as C
    from app.spatial.cadastre import FileCadastreProvider, _lot_from_props, resolve_parcel_geometry
    cid = rec["id"]
    try:
        feats = _read_cadastre_upload(content, (filename or "").lower(), src_epsg)
    except Exception as e:
        raise HTTPException(422, f"地籍圖檔解析失敗：{e}") from e
    if not feats:
        raise HTTPException(422, "地籍圖檔沒有任何圖徵")
    props0 = feats[0].get("properties") or {}
    sf_ = section_field or _pick_field(props0, _SEC_KEYS)
    lf_ = lot_field or _pick_field(props0, _LOT_KEYS)
    if not lf_:
        raise HTTPException(422, f"找不到地號欄位，請指定 lot_field（檔案欄位有：{', '.join(list(props0)[:12])}）")
    norm = []
    for f in feats:
        pr = f.get("properties") or {}
        lot = _lot_from_props(pr.get(lf_))
        sec = str(pr.get(sf_) or "").strip() if sf_ else ""
        if not lot:
            continue
        norm.append({"type": "Feature", "geometry": f["geometry"],
                     "properties": {"section": sec, "lot": lot.removesuffix("-0"), "lot_key": lot, **{k: v for k, v in pr.items() if isinstance(v, (str, int, float))}}})
    CADASTRE_DIR.mkdir(parents=True, exist_ok=True)
    out = CADASTRE_DIR / f"{cid}.geojson"
    out.write_text(json.dumps({"type": "FeatureCollection", "features": norm}, ensure_ascii=False), encoding="utf-8")
    data = copy.deepcopy(rec["data"])
    fp = FileCadastreProvider(norm, section_field="section", lot_field="lot_key", source="需用土地人地籍圖")
    matched, unmatched = [], []
    for p in [data["subject_parcel"], *data.get("comparables", [])]:
        if not p.get("parcel_id"):
            continue
        r = resolve_parcel_geometry(p, file_provider=fp, overwrite=True)
        (matched if r.get("source") == "cadastre_file" else unmatched).append(p.get("parcel_id") or "")
    data["case"]["cadastre"] = {"file": str(out), "n": len(norm), "section_field": sf_, "lot_field": lf_, "filename": filename, "at": C._now()}
    rec2 = C.save_case(data, cid=cid, submitted_table5=rec.get("submitted_table5"), submitted_table4=rec.get("submitted_table4"))
    return {"record": rec2, "n": len(norm), "section_field": sf_, "lot_field": lf_, "matched": matched, "unmatched": unmatched}


@app.post("/api/cases/{cid}/cadastre")
async def cases_cadastre(cid: str, request: Request, file: UploadFile = File(...), section_field: str = Form(""), lot_field: str = Form(""),   # noqa: B008
                         src_epsg: int = Form(3826)):
    """匯入地籍圖檔（GeoJSON、KML／GML（國土測繪中心地籍圖 API MAP_001／MAP_002 回傳格式）、或含 .shp/.dbf/.shx 的 zip，預設 TWD97 EPSG:3826）：存成本案地籍圖層（三張圖畫出每筆界線與地號），並用它補比準地與比較標的的真實幾何。"""
    from app import cases as C
    rec = C.get_case(cid)
    if not rec:
        raise HTTPException(404, "沒有這個案件")
    content = await file.read()
    r = _apply_cadastre(rec, content, file.filename or "", section_field, lot_field, src_epsg)
    AUD.log(cid, AUD.actor_from_headers(request.headers), "import", f"{file.filename}（地籍圖 {r['n']} 筆）：對到 {len(r['matched'])} 筆宗地" + (f"；對不到 {'、'.join(r['unmatched'])}" if r["unmatched"] else ""))
    return r


_SECID_KEYS = ("區段編號", "區段", "地價區段", "區段代碼", "SECTION_NO", "SECT_NO", "SECTNO", "section_id", "ZONE", "ZONE_NO", "編號", "NAME", "name")


def _norm_secid(v: Any) -> str:
    return str(v or "").strip().upper().replace(" ", "").replace("－", "-").replace("—", "-")


def _apply_section_map(rec: dict, content: bytes, filename: str, id_field: str = "", src_epsg: int = 3826) -> dict:
    """地價區段圖併入案件：依區段編號對到本案區段寫入正式範圍；整份圖存為圖層。回傳 {record, n, id_field, matched, unmatched, ids}。"""
    from app import cases as C
    cid = rec["id"]
    try:
        feats = _read_cadastre_upload(content, (filename or "").lower(), src_epsg)
    except Exception as e:
        raise HTTPException(422, f"區段圖解析失敗：{e}") from e
    if not feats:
        raise HTTPException(422, "區段圖沒有任何圖徵")
    props0 = feats[0].get("properties") or {}
    fld = id_field or _pick_field(props0, _SECID_KEYS)
    norm = []
    for f in feats:
        pr = f.get("properties") or {}
        sid = _norm_secid(pr.get(fld)) if fld else ""
        norm.append({"type": "Feature", "geometry": f["geometry"], "properties": {"section_id": sid, **{k: v for k, v in pr.items() if isinstance(v, (str, int, float))}}})
    CADASTRE_DIR.mkdir(parents=True, exist_ok=True)
    out = CADASTRE_DIR / f"{cid}_sections.geojson"
    out.write_text(json.dumps({"type": "FeatureCollection", "features": norm}, ensure_ascii=False), encoding="utf-8")
    data = copy.deepcopy(rec["data"])
    matched, unmatched = [], []
    by_id = {f["properties"]["section_id"]: f for f in norm if f["properties"]["section_id"]}
    for sid, sec in (data.get("sections") or {}).items():
        hit = by_id.get(_norm_secid(sid))
        if hit is None and len(norm) == 1:      # 檔案只有一個區段：就是本案區段
            hit = norm[0]
        if hit is None:
            unmatched.append(sid)
            continue
        sec["geometry"] = hit["geometry"]
        sec["geometry_source"] = "section_map"
        sec["status"] = "confirmed"
        sec["geometry_note"] = f"地價區段圖：{filename}（區段編號欄「{fld or '—'}」）"
        matched.append(sid)
    data["case"]["section_map"] = {"file": str(out), "n": len(norm), "id_field": fld, "filename": filename, "at": C._now(),
                                   "ids": sorted(by_id)[:50]}
    rec2 = C.save_case(data, cid=cid, submitted_table5=rec.get("submitted_table5"), submitted_table4=rec.get("submitted_table4"))
    return {"record": rec2, "n": len(norm), "id_field": fld, "matched": matched, "unmatched": unmatched, "ids": sorted(by_id)[:50]}


@app.post("/api/cases/{cid}/sections_map")
async def cases_sections_map(cid: str, request: Request, file: UploadFile = File(...), id_field: str = Form(""), src_epsg: int = Form(3826)):   # noqa: B008
    """匯入地價區段圖（GeoJSON／KML／GML／Shapefile zip）：依區段編號對到本案區段，寫入正式範圍多邊形（來源「地價區段圖」）；整份圖存為圖層，鄰近區段也畫在圖上。"""
    from app import cases as C
    rec = C.get_case(cid)
    if not rec:
        raise HTTPException(404, "沒有這個案件")
    content = await file.read()
    r = _apply_section_map(rec, content, file.filename or "", id_field, src_epsg)
    AUD.log(cid, AUD.actor_from_headers(request.headers), "import", f"{file.filename}（地價區段圖 {r['n']} 區段）：對到 {'、'.join(r['matched']) or '無'}" + (f"；對不到 {'、'.join(r['unmatched'])}" if r["unmatched"] else ""))
    return r


@app.get("/api/cases/{cid}/audit")
def cases_audit(cid: str, limit: int = 200):
    from app import cases as C
    if not C.get_case(cid):
        raise HTTPException(404, "沒有這個案件")
    return {"entries": AUD.read(cid, limit), "roles": AUD.ROLES, "actions": AUD.ACTIONS}


@app.get("/api/cases/{cid}/map.png")
def cases_map_png(cid: str, mode: str = "section", highlight: str | None = None, basemap: bool = True, appraiser: str = "", page: bool = False):
    """案件圖說 PNG（意見書附圖、Excel 圖說、頁面下載共用）。page=1 → 範本頁式（A3 橫式整頁，書表預覽用）。"""
    from app import cases as C
    from app.maps.layers import case_layers
    from app.maps.render import MODE_TITLE, render_png
    from app.maps.zoning import get_zoning_store
    from app.spatial.roads_store import get_roads
    rec = C.get_case(cid)
    if not rec:
        raise HTTPException(404, "沒有這個案件")
    data = rec["data"]
    layers = case_layers(data, zoning=get_zoning_store(), roads=get_roads(data=data), pad_m=900.0)
    try:
        sid = data.get("subject_parcel", {}).get("section_id", "")
        sub = f"案號 {data['case'].get('case_no', '')}　估價基準日 {data['case'].get('valuation_date', '')}　區段 {sid}"
        png = render_png(layers, mode, subtitle=sub, highlight=highlight or None, basemap=basemap, district=data["case"].get("district", ""), subject_section=sid, appraiser=appraiser or data["case"].get("appraiser", ""), page=page)
    except ValueError as e:
        raise HTTPException(422, str(e)) from e
    fname = f"{data['case'].get('case_no', 'case')}_{MODE_TITLE.get(mode, mode)}.png"
    return StreamingResponse(io.BytesIO(png), media_type="image/png",
                             headers={"Content-Disposition": f"inline; filename*=UTF-8''{urllib.parse.quote(fname)}", "Cache-Control": "no-store"})


@app.post("/api/cases/{cid}/duplicate")
def cases_duplicate(cid: str, request: Request, name: str | None = None):
    from app import cases as C
    rec = C.duplicate_case(cid, name)
    if not rec:
        raise HTTPException(404, "沒有這個案件")
    AUD.copy_log(cid, rec["id"], AUD.actor_from_headers(request.headers))
    return rec


@app.get("/api/cases/{cid}")
def cases_get(cid: str, touch: bool = True):
    from app import cases as C
    rec = C.get_case(cid)
    if not rec:
        raise HTTPException(404, "沒有這個案件")
    if touch:
        C.touch_opened(cid)
    return rec


@app.delete("/api/cases")
def cases_delete_all():
    """重置所有案件：清掉全部案件資料、操作紀錄與匯入的地籍圖檔（不可復原）。demo 前回到乾淨狀態用。"""
    from app import cases as C
    return {"ok": True, **C.delete_all_cases()}


@app.delete("/api/cases/{cid}")
def cases_delete(cid: str, request: Request):
    from app import cases as C
    if not C.delete_case(cid):
        raise HTTPException(404, "沒有這個案件")
    removed = []
    for f in CADASTRE_DIR.glob(f"{cid}*.geojson"):            # 本案匯入的地籍圖／區段圖檔一併移除
        try:
            f.unlink()
            removed.append(f.name)
        except OSError:
            pass
    AUD.log(cid, AUD.actor_from_headers(request.headers), "delete", "案件已刪除（紀錄保留）" + (f"；移除匯入檔 {len(removed)} 個" if removed else ""))
    return {"ok": True, "removed_files": removed}


class CadastreResolvePayload(CasePayload):
    manual_points: dict[str, list[float]] | None = None    # {parcel_id: [lon, lat]} 人工在圖上點的質心
    overwrite: bool = False


DEFAULT_CADASTRE = Path(__file__).resolve().parents[2] / "data" / "cadastre" / "default.geojson"


def _cadastre_provider_for(data: dict[str, Any]):
    """地籍界線來源順序：本案匯入的地籍圖 → 環境變數 CADASTRE_GEOJSON → data/cadastre/default.geojson（賽前放好的該區地籍圖）。"""
    import os

    from app.maps.layers import load_cadastre_features
    from app.spatial.cadastre import FileCadastreProvider
    case_feats = load_cadastre_features(data)
    if case_feats:
        return FileCadastreProvider(case_feats, section_field="section", lot_field="lot_key", source="需用土地人地籍圖")
    path = os.environ.get("CADASTRE_GEOJSON") or (str(DEFAULT_CADASTRE) if DEFAULT_CADASTRE.exists() else "")
    if path and Path(path).exists():
        return FileCadastreProvider.from_geojson(path, section_field=os.environ.get("CADASTRE_SECTION_FIELD", "段名"),
                                                 lot_field=os.environ.get("CADASTRE_LOT_FIELD", "地號"), source="地籍圖檔（預載）")
    return None


def _resolve_case_geometries(data: dict[str, Any], manual_points: dict[str, list[float]] | None, overwrite: bool) -> dict[str, Any]:
    """為缺 geometry 的比準地／比較標的補幾何：本案匯入地籍圖 → 預載地籍圖檔 → NLSC API（需申請）→ 人工質心合成。回傳每筆來源。"""
    import os

    from app.maps.layers import load_cadastre_features
    from app.spatial.cadastre import NLSCCadastreProvider, TwlandCadastreProvider, resolve_parcel_geometry
    fp = _cadastre_provider_for(data)
    nlsc = NLSCCadastreProvider() if os.environ.get("NLSC_API_KEY") else None
    twland = None if os.environ.get("TWLAND_OFFLINE") == "1" else TwlandCadastreProvider()      # 免金鑰的開放地籍查詢，排在地籍圖檔之後
    case_feats = load_cadastre_features(data)       # 點圖時若點在本案地籍圖某筆宗地內，直接用那筆的真實界線
    report: dict[str, Any] = {}
    pts = manual_points or {}
    for p in [data["subject_parcel"], *data["comparables"]]:
        pid = p.get("parcel_id")
        pt = pts.get(pid)
        if pt and case_feats:
            from shapely.geometry import Point as _Pt
            from shapely.geometry import shape as _shape
            hit = next((f for f in case_feats if _shape(f["geometry"]).contains(_Pt(pt[0], pt[1]))), None)
            if hit:
                pr = hit.get("properties") or {}
                p["geometry"] = hit["geometry"]
                p["geometry_source"] = "cadastre_file"
                p["geometry_note"] = f"地籍圖：點選位置落在 {pr.get('section') or ''}{pr.get('lot') or ''}地號，採該筆真實界線"
                report[pid] = {"ok": True, "source": "cadastre_file", "note": p["geometry_note"]}
                continue
        report[pid] = resolve_parcel_geometry(p, file_provider=fp, nlsc=nlsc, manual_point=pt, district=data["case"].get("district"),
                                              overwrite=overwrite, twland=twland)
    report["_providers"] = {"cadastre_file": bool(fp), "nlsc_api": bool(nlsc), "twland": bool(twland)}
    return report


@app.post("/api/cadastre/resolve")
def cadastre_resolve(payload: CadastreResolvePayload):
    """為缺 geometry 的比準地／比較標的補幾何。回傳案件 + 每筆來源。"""
    data = payload.model_dump(exclude={"manual_points", "overwrite"})
    report = _resolve_case_geometries(data, payload.manual_points, payload.overwrite)
    providers = report.pop("_providers")
    return {"data": data, "report": report, "providers": providers}


class ComparableLocatePayload(BaseModel):
    comp_no: int
    lon: float
    lat: float


@app.post("/api/cases/{cid}/comparables/locate")
def comparables_locate(cid: str, payload: ComparableLocatePayload, request: Request):
    """比較標的在圖上點位置：以該點＋實例土地面積合成示意範圍（synthetic），再重推宗地屬性、區段草稿、勘查表與設施距離。"""
    from shapely.geometry import mapping

    from app import cases as C
    from app.spatial.cadastre import synthesize_parcel_geometry
    rec = C.get_case(cid)
    if not rec:
        raise HTTPException(404, "沒有這個案件")
    comps = copy.deepcopy(rec["data"].get("comparables") or [])
    c = next((x for x in comps if int(x.get("comp_no") or 0) == payload.comp_no), None)
    if c is None:
        raise HTTPException(404, f"沒有比較標的{payload.comp_no}")
    poly = synthesize_parcel_geometry([payload.lon, payload.lat], c.get("area_m2"), c.get("width_m"), c.get("depth_m"))
    c["geometry"], c["geometry_source"] = mapping(poly), "synthetic"
    c["geometry_note"] = f"以人工指定位置＋實例土地面積 {c.get('area_m2') or '—'} m² 合成示意範圍，非地籍圖，距離結果需確認"
    for k in list((c.get("derived") or {}).keys()):          # 位置變了，推定值重算
        c.pop(k, None) if k not in ("parcel_id",) else None
    c["derived"] = {}
    rec, notes = _apply_comparables(rec, comps)
    AUD.log(cid, AUD.actor_from_headers(request.headers), "locate", f"比較標的{payload.comp_no} {c.get('parcel_id')} 在圖上設定位置")
    return {"case": rec, "notes": notes}


class ComparableSearchPayload(BaseModel):
    max_n: int = 3
    relax: bool = True          # §17 第3項：無適當實例放寬至基準日前一年
    neighbors: bool = True      # §19 第2項：鄰近鄉鎮


class ComparableApplyPayload(ComparableSearchPayload):
    ids: list[str]                                   # 實價登錄編號（候選清單的 id）
    building_costs: dict[str, float] | None = None   # 含建物者的建物成本價格（元）
    reasons: dict[str, str] | None = None            # 蒐集期間外參考案例的採用理由（手冊 p.77 問答四），寫進備註
    replace: bool = True


def _cadastre_lookup_for(data: dict[str, Any]):
    fp = _cadastre_provider_for(data)
    return fp.lookup if fp else None


def _locate_by_address(c: dict, data: dict) -> bool:
    """實價登錄門牌 → 路網路段中點 → 依土地面積合成示意範圍（geometry_source=address_estimate）。"""
    from shapely.geometry import mapping

    from app.spatial.cadastre import synthesize_parcel_geometry
    from app.spatial.locate import point_from_address
    from app.spatial.roads_store import get_roads
    pos = c.get("address") or (c.get("source") or {}).get("position") or ""
    try:
        hit = point_from_address(pos, data["case"].get("district"), get_roads)
    except Exception:  # noqa: BLE001 - 對不到就留人工
        hit = None
    if not hit:
        return False
    poly = synthesize_parcel_geometry([hit["lon"], hit["lat"]], c.get("area_m2"), c.get("width_m"), c.get("depth_m"))
    c["geometry"], c["geometry_source"], c["geometry_note"] = mapping(poly), "address_estimate", hit["note"]
    return True


def _apply_comparables(rec: dict, comps: list[dict]) -> tuple[dict, list[str]]:
    """比較標的寫入案件：區段外者建區段草稿（路網街廓＋勘查表推算），有界線者推定宗地屬性，最後重量設施距離。回傳 (rec, notes)。"""
    from app import cases as C
    from app.maps.zoning import get_zoning_store
    from app.spatial import service
    from app.spatial.bootstrap import block_from_roads, describe_range
    from app.spatial.geo import centroid
    from app.spatial.lot import derive_parcel_attributes
    from app.spatial.roads_store import get_roads
    from app.spatial.survey_draft import draft_section_survey
    from app.spatial.walking import get_walk_graph
    data = copy.deepcopy(rec["data"])
    reg, ind = _rulesets(data["case"])
    roads, zoning = get_roads(data=data), get_zoning_store()
    notes: list[str] = []
    subj_sec = data["subject_parcel"].get("section_id")
    data["sections"] = {k: v for k, v in data["sections"].items() if not k.startswith(f"{subj_sec}-C")}   # 舊的比較標的區段草稿清掉
    for c in comps:
        if c.get("geometry") is None:                        # 不在地籍圖：用實價登錄門牌對路網取推定位置，合成示意範圍
            _locate_by_address(c, data)
        croads = roads
        if c.get("geometry") is not None:
            from shapely.geometry import shape as _shape

            from app.spatial.area import pad_bbox
            cb = pad_bbox(_shape(c["geometry"]).bounds, 2500.0)
            croads = get_roads(cb)                            # 鄰近鄉鎮的實例（§19 第2項）用它自己周邊的路網，不受案件範圍限制
            d = derive_parcel_attributes(c, roads=croads, zoning=zoning, overwrite=False, district=data["case"].get("district") or "", store=service.get_poi_store())
            if d["notes"]:
                notes.append(f"比較標的{c['comp_no']}：{'；'.join(d['notes'])}")
            if c.get("geometry_source") == "address_estimate":
                notes.append(f"比較標的{c['comp_no']} {c['parcel_id']} 不在地籍圖內，位置依門牌路段推定（{c.get('geometry_note', '')}）；寬深形狀請人工填載，設施距離為推定")
        else:
            notes.append(f"比較標的{c['comp_no']} {c['parcel_id']} 不在已匯入的地籍圖內，門牌也對不到路網：請在圖上「設定比較標的位置」或匯入地籍圖；面積寬深等條件與設施距離請人工填載")
        sid = c["section_id"]
        if sid != subj_sec and sid not in data["sections"]:
            sec: dict[str, Any] = {"section_id": sid, "range_desc": "", "survey": C.blank_survey(), "status": "draft",
                                   "geometry_note": f"比較標的{c['comp_no']}所在區段（其他地區，查估辦法 §19 第2項）"}
            if c.get("geometry") is not None and croads.items:
                cc = centroid(c["geometry"])
                res = block_from_roads([(cc.x, cc.y)], croads)
                if res:
                    hit = zoning.at((cc.x, cc.y)) if len(zoning) else None
                    zone_label = c.get("zoning") or (hit["zone"] if hit and "道路" not in (hit["zone"] or "") else None)   # 門牌推定點落在路上會打到道路用地，改用實例的分區
                    sec["geometry"], sec["geometry_source"] = res["geometry"], "estimate_osm_block"
                    sec["range_desc"] = describe_range(res["geometry"], croads, zoning=zone_label, section_id=sid)["range_desc"]
                    try:
                        from app.spatial.area import bbox_key as _bk
                        wg = get_walk_graph(pad_bbox(_shape(sec["geometry"]).bounds, 3000.0)) if _bk(cb) != _bk(pad_bbox(_shape(data["subject_parcel"]["geometry"]).bounds, 2500.0)) else get_walk_graph(data=data)
                        draft_section_survey(sec, reg, store=service.get_poi_store(), zoning=zoning, roads=croads, osrm=service.get_osrm(), walk_graph=wg, subject=data.get('subject_parcel'))
                    except Exception as e:  # noqa: BLE001
                        notes.append(f"區段 {sid} 勘查表推算失敗：{e}")
            data["sections"][sid] = sec
    data["comparables"] = comps
    from app.market.index import fill_date_adjustments
    notes += fill_date_adjustments(data)
    try:
        f = service.fill_case(data, regional=reg, individual=ind, overwrite=False)
        notes += f["warnings"]
    except Exception as e:  # noqa: BLE001
        notes.append(f"設施距離量測失敗：{e}")
    rec = C.save_case(data, cid=rec["id"], submitted_table5=rec.get("submitted_table5"), submitted_table4=rec.get("submitted_table4"))
    return rec, notes


@app.post("/api/cases/{cid}/comparables/search")
def cases_comparables_search(cid: str, payload: ComparableSearchPayload):
    """實價登錄搜尋比較標的候選（不寫入案件）。"""
    from app import cases as C
    from app.market.lvr import search_comparables
    rec = C.get_case(cid)
    if not rec:
        raise HTTPException(404, "沒有這個案件")
    data = rec["data"]
    return search_comparables(data, max_n=payload.max_n, relax=payload.relax, neighbors=payload.neighbors, cadastre_lookup=_cadastre_lookup_for(data))


@app.post("/api/cases/{cid}/comparables/apply")
def cases_comparables_apply(cid: str, payload: ComparableApplyPayload, request: Request):
    """把勾選的實價登錄實例寫成本案比較標的（含建物者須附建物成本價格）。"""
    from app import cases as C
    from app.market.lvr import build_comparables, search_comparables
    rec = C.get_case(cid)
    if not rec:
        raise HTTPException(404, "沒有這個案件")
    if not payload.ids:
        raise HTTPException(422, "請至少勾選一件實例")
    if len(payload.ids) > 3:
        raise HTTPException(422, "比較標的最多三件（查估辦法 §19 第1項第1款）")
    data = rec["data"]
    search = search_comparables(data, max_n=payload.max_n, relax=True, neighbors=True, cadastre_lookup=_cadastre_lookup_for(data))
    try:
        comps = build_comparables(data, search, payload.ids, payload.building_costs, payload.reasons)
    except (KeyError, ValueError) as e:
        raise HTTPException(422, str(e)) from e
    rec2, notes = _apply_comparables(rec, comps)
    AUD.log(cid, AUD.actor_from_headers(request.headers), "comparables",
            f"採用實價登錄實例 {len(comps)} 件：{'、'.join(c['parcel_id'] for c in comps)}" + ("；" + "；".join(notes) if notes else ""))
    return {"rec": rec2, "notes": notes, "comparables": comps}


class FromLotPayload(BaseModel):
    parcel_id: str = ""
    manual_point: list[float] | None = None     # [lon, lat]：找不到地籍界線時，人工指定比準地位置
    overwrite: bool = False                     # 同一地號重跑時是否覆寫既有推定值（換地號一律重算）
    parcel_changed: bool = False                # 前端已先存新地號時，用這個告知換了地號
    with_comparables: bool = True               # 換地號或尚無比較標的時，自動到實價登錄找比較標的


@app.post("/api/cases/{cid}/from_lot")
def cases_from_lot(cid: str, payload: FromLotPayload, request: Request):
    """依比準地地號產生：地籍界線 → 宗地屬性 → 區段範圍草稿 → 勘查表推算 → 設施距離，存回案件並回報每一步。"""
    from app import cases as C
    from app.maps.zoning import get_zoning_store
    from app.spatial import service
    from app.spatial.lot import generate_from_lot
    from app.spatial.roads_store import get_roads
    from app.spatial.walking import get_walk_graph
    rec = C.get_case(cid)
    if not rec:
        raise HTTPException(404, "沒有這個案件")
    data = copy.deepcopy(rec["data"])
    reg, ind = _rulesets(data["case"])
    pre_steps: list[dict[str, Any]] = []
    parcel_id = payload.parcel_id
    if not (parcel_id or data["subject_parcel"].get("parcel_id")):     # 無地號：用區段範圍文字圍區段、依 §18 自動選比準地
        from app.spatial.lot import prepare_from_range
        pre_steps = prepare_from_range(data, roads=get_roads(data=data), zoning=get_zoning_store(), provider=_cadastre_provider_for(data))
        parcel_id = data["subject_parcel"].get("parcel_id") or ""
        payload.parcel_changed = bool(parcel_id)
    out = generate_from_lot(data, parcel_id=parcel_id, manual_point=payload.manual_point, overwrite=payload.overwrite, parcel_changed=payload.parcel_changed,
                            regional=reg, individual=ind, resolve_geometries=_resolve_case_geometries,
                            roads=get_roads(data=data), zoning=get_zoning_store(), store=service.get_poi_store(), osrm=service.get_osrm(), walk_graph=get_walk_graph(data=data))
    out["steps"] = pre_steps + out["steps"]
    rec = C.save_case(out["data"], cid=cid, submitted_table5=rec.get("submitted_table5"), submitted_table4=rec.get("submitted_table4"))
    if payload.with_comparables and out["data"]["subject_parcel"].get("geometry") is not None and (out["parcel_changed"] or not out["data"].get("comparables")):
        from app.market.lvr import build_comparables, search_comparables
        try:
            sr = search_comparables(out["data"], max_n=3, relax=True, neighbors=True, cadastre_lookup=_cadastre_lookup_for(out["data"]))
            if sr["chosen"]:
                comps = build_comparables(out["data"], sr, [c["id"] for c in sr["chosen"]])
                rec, notes = _apply_comparables(rec, comps)
                out["steps"].append({"step": "比較標的", "ok": True, "note": f"實價登錄自動選取 {len(comps)} 件：" + "、".join(f"{c['parcel_id']}（{c['transaction_date']}，{c['selection']['basis']}）" for c in comps)
                                     + ("；" + "；".join(notes) if notes else "")})
            else:
                st = sr["stats"]
                nb = sum(1 for c in sr["candidates"] if c.get("needs_building_cost"))
                out["steps"].append({"step": "比較標的", "ok": False,
                                     "note": f"實價登錄無可自動採用之純土地實例（蒐集期間 {sr['window']['text']}，放寬至 {sr['window']['relaxed_from']}；同用地別 {st['n_zone']} 筆，特殊情況排除 {st['n_excluded']} 筆"
                                             + (f"，其中 {nb} 筆含建物，填建物成本價格後可採用" if nb else "") + "）。請到「宗地條件與買賣實例」勾選候選或人工填寫"})
        except Exception as e:  # noqa: BLE001
            out["steps"].append({"step": "比較標的", "ok": False, "note": f"搜尋失敗：{e}"})
    ok = [s["step"] for s in out["steps"] if s.get("ok")]
    bad = [s["step"] for s in out["steps"] if not s.get("ok")]
    rec = C.patch_case(cid, last_fill={"at": C._now(), "parcel_id": out["data"]["subject_parcel"].get("parcel_id"), "steps": out["steps"]}) or rec
    AUD.log(cid, AUD.actor_from_headers(request.headers), "from_lot",
            f"比準地 {out['data']['subject_parcel'].get('parcel_id')}：完成 {'、'.join(ok) or '無'}" + (f"；未完成 {'、'.join(bad)}" if bad else ""))
    return {"rec": rec, "steps": out["steps"], "parcel_changed": out["parcel_changed"]}


@app.get("/api/cadastre/lots")
def cadastre_lots():
    """預載地籍圖（或 CADASTRE_GEOJSON）裡的地號清單，給一鍵建案的比準地下拉選單。"""
    fp = _cadastre_provider_for({"case": {}})
    if fp is None:
        return {"source": None, "lots": []}
    from shapely.geometry import shape

    from app.spatial.area import _districts
    lots = []
    sec_district: dict[str, str | None] = {}                 # 段 → 鄉鎮市區（取該段第一筆界線中心點落在哪個區界；一段不跨區）
    for (sec, lot), v in fp.index.items():
        try:
            main, sub = lot.split("-")
            key = (sec, int(main), int(sub))
        except ValueError:
            key = (sec, 0, 0)
        if sec not in sec_district:
            d = None
            try:
                g = v.get("geometry") if isinstance(v, dict) else None
                c = shape(g).centroid if g else None
                if c is not None:
                    d = next((n for n, dg in _districts() if dg.contains(c)), None)
            except Exception:  # noqa: BLE001 - 沒有區界檔或幾何壞掉就不標
                d = None
            sec_district[sec] = d
        lots.append({"section": sec, "lot": lot, "parcel_id": f"{sec}{lot.removesuffix('-0')}地號", "_k": key,
                     "district": sec_district[sec], "area_m2": (v.get("properties") or {}).get("area_m2")})
    lots.sort(key=lambda x: x["_k"])
    for x in lots:
        x.pop("_k")
    return {"source": fp.source, "n": len(lots), "lots": lots, "districts": sorted({d for d in sec_district.values() if d})}


@app.get("/api/cadastre/sections")
def cadastre_sections(county: str = "新北市", town: str = "金山區"):
    """NLSC 開放 API 段代碼（有快取；離線時只回快取）。"""
    from app.spatial.cadastre import CadastreError, SectionCodes
    try:
        sc = SectionCodes()
        cc = sc.counties().get(county.replace("台", "臺"))
        tc = sc.towns(cc).get(town) if cc else None
        if not (cc and tc):
            raise HTTPException(404, f"找不到 {county}{town}")
        return {"county_code": cc, "town_code": tc, "sections": sc.sections(cc, tc)}
    except CadastreError as e:
        raise HTTPException(503, str(e)) from e
    except Exception as e:
        raise HTTPException(503, f"NLSC API 無法連線：{e}") from e


class BlockPayload(BaseModel):
    hints: list[list[float]]                 # [[lon, lat], ...]；第一個是主點，其餘是「區段內」的參考點
    exclude_names: list[str] = []
    extra_names: list[str] = []
    zoning: str | None = None
    section_id: str | None = None


@app.post("/api/bootstrap/block")
def bootstrap_block(payload: BlockPayload):
    """點幾個位置 → 路網圍出的街廓 + 四至草稿（區段範圍描述），供估價師修改。"""
    from app.maps.zoning import get_zoning_store
    from app.spatial.area import pad_bbox
    from app.spatial.bootstrap import block_from_roads, describe_range
    from app.spatial.roads_store import get_roads
    hb = (min(h[0] for h in payload.hints), min(h[1] for h in payload.hints), max(h[0] for h in payload.hints), max(h[1] for h in payload.hints))
    roads = get_roads(pad_bbox(hb, 2500.0))
    if not roads.items:
        raise HTTPException(503, "沒有路網資料（ROADS_GEOJSON／data/osm/roads.sqlite）")
    res = block_from_roads([tuple(h) for h in payload.hints], roads, exclude_names=payload.exclude_names, extra_names=payload.extra_names)
    if not res:
        raise HTTPException(404, "路網圍不出包含這些點的街廓，請改點位置或指定道路")
    z = payload.zoning
    if z is None:
        hit = get_zoning_store().at(tuple(payload.hints[0]))
        z = hit["zone"] if hit else None
        zoning_note = f"分區取自使用分區圖（{hit['source']}），細分區（第X種）需人工確認" if hit else "無使用分區圖資"
    else:
        zoning_note = "使用者指定"
    rng = describe_range(res["geometry"], roads, zoning=z, section_id=payload.section_id)
    return {**res, **rng, "zoning": z, "zoning_note": zoning_note,
            "basis": "查估辦法 §10、§11；街廓與四至由 OSM 路網推得，為草稿"}


class ProposePayload(BaseModel):
    scope: dict[str, Any] | None = None          # 徵收範圍 GeoJSON
    parcels: list[dict[str, Any]] | None = None  # 或宗地幾何清單（取聯集）
    use_zoning: bool = True
    prefix: str = "P"


@app.post("/api/bootstrap/sections")
def bootstrap_sections(payload: ProposePayload):
    """深模式：徵收範圍 × 使用分區 → 區段草稿清單（status=draft）。"""
    from shapely.geometry import shape
    from shapely.ops import unary_union

    from app.maps.zoning import get_zoning_store
    from app.spatial.bootstrap import propose_sections
    from app.spatial.roads_store import get_roads
    if payload.scope is None and not payload.parcels:
        raise HTTPException(400, "需要 scope 或 parcels")
    scope = payload.scope or [p for p in payload.parcels]
    zoning_feats = None
    if payload.use_zoning:
        area = shape(scope) if isinstance(scope, dict) else unary_union([shape(g) for g in scope])
        zoning_feats = get_zoning_store().within_bbox(area.bounds) or None
    from app.spatial.area import pad_bbox
    area_b = (shape(scope) if isinstance(scope, dict) else unary_union([shape(g) for g in scope])).bounds
    drafts = propose_sections(scope, get_roads(pad_bbox(area_b, 2000.0)), zoning_features=zoning_feats, zoning_name_field="zone", prefix=payload.prefix)
    return {"sections": drafts, "count": len(drafts)}


@app.post("/api/maps/layers")
def maps_layers(payload: CasePayload):
    from app.maps.layers import case_layers
    from app.maps.zoning import get_zoning_store
    from app.spatial.roads_store import get_roads
    data = payload.model_dump()
    return case_layers(data, zoning=get_zoning_store(), roads=get_roads(data=data))


@app.get("/maps", response_class=HTMLResponse)
def maps_viewer():
    return (Path(__file__).resolve().parent / "maps" / "static" / "viewer.html").read_text(encoding="utf-8")


class ManualPOIPayload(BaseModel):
    type: str
    name: str
    lon: float
    lat: float
    note: str = ""
    by: str = ""


@app.post("/api/spatial/poi")
def spatial_add_poi(payload: ManualPOIPayload):
    """人工標定設施（OSM／政府資料沒有的），寫入 data/sources/manual_poi.geojson，來源標「人工標定」。"""
    from app.spatial import service
    from app.spatial.reference import facility_measurement
    if payload.type not in facility_measurement()["facility_types"]:
        raise HTTPException(400, f"type 須為 facility_measurement.json 的 key，例如 {list(facility_measurement()['facility_types'])[:6]}…")
    return service.add_manual_poi(payload.type, payload.name, payload.lon, payload.lat, note=payload.note, by=payload.by)


class LandValueSectionsPayload(BaseModel):
    parcels: list[dict[str, Any]]      # [{"parcel_id", "geometry", "land_value"}]
    prefix: str = "P"


@app.post("/api/bootstrap/sections_by_land_value")
def bootstrap_sections_by_land_value(payload: LandValueSectionsPayload):
    """地價區段圖拿不到時：同公告現值且相鄰的宗地 → 區段草稿。需要宗地幾何。"""
    from shapely.geometry import shape as _shape
    from shapely.ops import unary_union as _uu

    from app.spatial.area import pad_bbox
    from app.spatial.bootstrap import sections_from_land_values
    from app.spatial.roads_store import get_roads
    pb = _uu([_shape(p["geometry"]) for p in payload.parcels if p.get("geometry")]).bounds if any(p.get("geometry") for p in payload.parcels) else None
    drafts = sections_from_land_values(payload.parcels, get_roads(pad_bbox(pb, 2000.0)) if pb else get_roads(), prefix=payload.prefix)
    return {"sections": drafts, "count": len(drafts)}


# ------------------------------------------------------------------ 審查意見書（第 5 步）


class ReportPayload(VerifyPayload):
    polish: bool = False          # 用 LLM 潤稿（需 LLM_PROVIDER 設好）；守門失敗自動退回模板句
    reviewer: str = ""
    review_date: str | None = None
    decisions: dict[str, Any] | None = None   # 承辦裁決 {finding_key: {decision, note}}
    reviewer_role: str = ""                   # officer | reviewer | appraiser（落款用）
    case_id: str | None = None                 # 有給就附三張圖說連結
    record: bool = False                       # True 才記操作紀錄（頁面預覽不記；下載 Word／PDF 由各該端點記）


@app.post("/api/report")
def report(payload: ReportPayload, request: Request):
    """審查：跑 verify → 逐條意見（每句帶 [F-xxx] 與法源）→ Markdown；可選 LLM 潤稿。"""
    from app.report.opinion import build_report, polish_items, to_markdown
    reg, ind = _rulesets(payload.case)
    data = payload.model_dump(exclude={"polish", "reviewer", "review_date", "decisions", "reviewer_role", "case_id"})
    result = run_case(reg, ind, data)
    findings = collect_findings(reg, ind, data, result, payload.submitted_table5, payload.submitted_table4)
    computed = {"table5": {k: v.to_dict() for k, v in result["table5"].items()}, "table4": result["table4"].to_dict()}
    reviewer = payload.reviewer.strip()
    if reviewer and payload.reviewer_role in AUD.ROLES:
        reviewer = f"{reviewer}（{AUD.ROLES[payload.reviewer_role]}）"
    figures = None
    if payload.case_id:
        from app.maps.render import MODE_TITLE
        has_geom = bool(data["subject_parcel"].get("geometry") or any((s or {}).get("geometry") for s in (data.get("sections") or {}).values()))
        if has_geom:
            figures = [{"mode": m, "title": t, "url": f"/api/cases/{payload.case_id}/map.png?mode={m}"} for m, t in MODE_TITLE.items()]
    rep = build_report(data["case"], findings, computed, subject_parcel_id=data["subject_parcel"].get("parcel_id", ""),
                       reviewer=reviewer, review_date=payload.review_date, decisions=payload.decisions, figures=figures,
                       has_submitted=bool(payload.submitted_table4 or payload.submitted_table5))
    if payload.case_id and payload.record:
        AUD.log(payload.case_id, AUD.actor_from_headers(request.headers), "report", f"結論 {rep.conclusion_code}；審查人 {reviewer or '—'}")
    if payload.polish:
        from app.llm import LLMNotConfigured, get_provider
        try:
            rep.polish = polish_items([it for items in rep.sections.values() for it in items], get_provider())
        except LLMNotConfigured as e:
            rep.polish = {"status": "unavailable", "note": str(e)}
    return {"report": rep.to_dict(), "markdown": to_markdown(rep), "findings": findings}


# ------------------------------------------------------------------ 前端顯示用的名稱對照（設施類型中文、量測方式、基準表名稱）


@app.get("/api/meta")
def meta():
    from app.report.basis import LEGAL_BASIS
    from app.spatial.reference import facility_measurement
    fm = facility_measurement()
    rules_summary = {p.stem: _ruleset_summary(p) for p in sorted(RULES_DIR.glob("*.json")) if "_regional" in p.name or "_individual" in p.name}
    return {
        "facility_types": {k: v.get("label", k) for k, v in fm["facility_types"].items()},
        "facility_modes": {k: v.get("mode") for k, v in fm["facility_types"].items()},
        "measure_labels": {"walking": "步行距離", "straight": "直線距離", "straight_estimated": "直線估算"},
        "origin_labels": {"parcel_centroid": "宗地中心點", "parcel_frontage": "臨路邊界", "section_boundary": "地價區段邊界", "subject_parcel": "比準地"},
        "geometry_sources": {"cadastre_file": "地籍圖", "section_map": "地價區段圖", "nlsc_api": "國土測繪中心地籍查詢", "synthetic": "依清冊面積合成（示意，非地籍圖）",
                             "estimate_osm_block": "依路網推估之街廓（草稿）"},
        "tables": {"1": "地價區段勘查表", "5": "影響地價區域因素分析明細表", "4": "比較法調查估價表", "7": "宗地個別因素清冊", "2": "買賣實例調查估價表"},
        "checklist": {"iii": "勘查表各細項優劣等級依評價基準明細表填列", "iv": "地價區段劃設與圖說", "v": "買賣實例蒐集期間與正常單價",
                      "vi": "區域因素分析明細表等級與修正百分比", "vii": "比較法調查估價表差異率、跨表抄填與權重",
                      "ix": "比準地地價估計表", "x": "宗地市價估計表", "p.53(十二)": "比準地比較價格尾數"},
        "rulesets": rules_summary,
        "legal_basis": LEGAL_BASIS,
        "zoning_bcr_far": __import__("app.spatial.admin", fromlist=["load_table"]).load_table(),
        "roles": AUD.ROLES,
    }


# ------------------------------------------------------------------ 勘查表推算（表1 為產出）


class SurveyDraftPayload(CasePayload):
    section_id: str | None = None
    overwrite: bool = False


@app.post("/api/spatial/survey_draft")
def spatial_survey_draft(payload: SurveyDraftPayload):
    """只有區段基本資訊時，依分區圖／路網／設施資料庫推算勘查表可推的欄位，其餘列為需人工填載。回傳整個案件（sections 已更新）。"""
    from app.maps.zoning import get_zoning_store
    from app.spatial import service
    from app.spatial.roads_store import get_roads
    from app.spatial.survey_draft import draft_section_survey
    from app.spatial.walking import get_walk_graph
    reg, _ind = _rulesets(payload.case)
    data = payload.model_dump(exclude={"section_id", "overwrite"})
    sid = payload.section_id or data["subject_parcel"].get("section_id") or next(iter(data["sections"]), None)
    if not sid or sid not in data["sections"]:
        raise HTTPException(404, f"找不到區段 {sid}")
    r = draft_section_survey(data["sections"][sid], reg, store=service.get_poi_store(), zoning=get_zoning_store(), roads=get_roads(data=data),
                             osrm=service.get_osrm(), walk_graph=get_walk_graph(data=data), overwrite=payload.overwrite, subject=data.get("subject_parcel"))
    from app.spatial.admin import fill_section_admin
    adm = fill_section_admin(data["sections"][sid], district=data["case"].get("district") or "", overwrite=payload.overwrite)
    filled = sorted(set(r["filled"]) | set(adm))
    manual = {k: v for k, v in r["manual"].items() if k not in adm}
    return {"data": data, "section_id": sid, "filled": filled, "suggestions": r["suggestions"], "manual": manual, "warnings": r["warnings"]}
