"""
地政局正式範本 xlsx 寫入器：把引擎結果直接填進主辦提供的三份 Excel 範本（表3 地價區段勘查表、表5-1 影響地價區域因素
分析明細表（住宅用地）、表4 比較法調查估價表），格線、合併格、字型全部沿用範本，只寫值。

範本放在 app/templates/official/；每份範本有二十多張工作表（表7 清冊、表1 買賣實例…），輸出時只留目標工作表。
格位對照：docs/12_official_templates.md。寫值一律寫在合併範圍的左上格（openpyxl 規則）。
"""
from __future__ import annotations

import io
import re
import zipfile
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from openpyxl.styles import Alignment
from openpyxl.workbook.workbook import Workbook
from openpyxl.worksheet.worksheet import Worksheet

from app.engine.rules import GradeError, RuleSet, grade, level_number
from app.engine.tables import Table4, Table5, _get, _parcel_obs

from .xlsx import level_label

TEMPLATE_DIR = Path(__file__).resolve().parents[1] / "templates" / "official"
TEMPLATES = {
    "t3": (TEMPLATE_DIR / "表3地價區段勘查表.xlsx", "表3區段勘查表", "表3_地價區段勘查表"),
    "t5": (TEMPLATE_DIR / "表5影響地價區域因素分析明細表.xlsx", "表5-1區域因素明細表(住)", "表5-1_影響地價區域因素分析明細表"),
    "t4": (TEMPLATE_DIR / "表4比較法調查估價表.xlsx", "表4比較法調查估價表", "表4_比較法調查估價表"),
    "t2": (TEMPLATE_DIR / "表4比較法調查估價表.xlsx", "94表2收益法 ", "表2_收益法調查估價表"),
    "t14": (TEMPLATE_DIR / "表4比較法調查估價表.xlsx", "114表14比準地地價估計表", "表14_比準地地價估計表"),
}
T2_COST_SHEET = "95表2收益附-成本法"
PCT = '0.00"%"'
MONEY = "#,##0"
WRAP = Alignment(wrap_text=True, vertical="center")


# ------------------------------------------------------------------ 小工具


def _open(key: str) -> tuple[Workbook, Worksheet]:
    """載入範本，只留目標工作表。"""
    path, sheet, _ = TEMPLATES[key]
    wb = load_workbook(path)
    ws = wb[sheet]
    for other in list(wb.worksheets):
        if other.title != sheet:
            wb.remove(other)
    return wb, ws


def _anchor(ws: Worksheet, coord: str) -> str:
    """coord 落在合併範圍內 → 回傳左上格；否則原格。"""
    for rng in ws.merged_cells.ranges:
        if coord in rng:
            return rng.start_cell.coordinate
    return coord


def _set(ws: Worksheet, coord: str, value: Any, fmt: str | None = None, wrap: bool = False) -> None:
    c = ws[_anchor(ws, coord)]
    c.value = value
    if fmt and isinstance(value, (int, float)):
        c.number_format = fmt
    if wrap:
        c.alignment = Alignment(horizontal=c.alignment.horizontal, vertical=c.alignment.vertical or "center", wrap_text=True)


def _mark(ws: Worksheet, coord: str, in_section: bool | None) -> None:
    """「○本區段內　○本區段外(距」→ 依設施在區段內／外把對應的 ○ 換成 ●。in_section None 不動。"""
    if in_section is None:
        return
    a = _anchor(ws, coord)
    t = ws[a].value
    if not isinstance(t, str):
        return
    if in_section:
        t = t.replace("○本區段內", "●本區段內", 1)
    else:
        t = t.replace("○本區段外", "●本區段外", 1)
    ws[a].value = t


def _radio(ws: Worksheet, coord: str, on: bool) -> None:
    """單獨的 ○ 格（大型車站、殯葬、污染的選項圓圈）。"""
    a = _anchor(ws, coord)
    ws[a].value = "●" if on else "○"


def _num(v: Any) -> float | int | None:
    if isinstance(v, bool) or v is None:
        return None
    if isinstance(v, (int, float)):
        return int(v) if float(v).is_integer() else v
    if isinstance(v, dict):
        for k in ("width_m", "distance_m", "value"):
            if isinstance(v.get(k), (int, float)):
                return _num(v[k])
    return None


def _txt(v: Any) -> str:
    if v is None:
        return ""
    if isinstance(v, bool):
        return "有" if v else "無"
    if isinstance(v, dict):
        return str(v.get("name") or v.get("value") or "")
    if isinstance(v, list):
        return "、".join(str(x.get("name") if isinstance(x, dict) else x) for x in v if x)
    return str(v)


def _first(lst: Any, types: tuple[str, ...] | None = None) -> dict | None:
    for f in lst or []:
        if isinstance(f, dict) and (types is None or f.get("type") in types):
            return f
    return None


def _fac(ws: Worksheet, f: dict | None, name_cell: str | None, mark_cell: str | None, dist_cell: str | None, prefix: str = "名稱：") -> None:
    """一列設施：名稱格、○本區段內／外 標記、距離格。沒有設施 → 名稱填「無」。"""
    if not f or not f.get("name"):
        if name_cell:
            _set(ws, name_cell, f"{prefix}無" if prefix else "無")
        return
    if name_cell:
        _set(ws, name_cell, f"{prefix}{f['name']}")
    ins = f.get("in_section")
    if mark_cell:
        _mark(ws, mark_cell, ins)
    if dist_cell and ins is False and _num(f.get("distance_m")) is not None:
        _set(ws, dist_cell, _num(f.get("distance_m")))


def _checks(text: str, selected: Any) -> str:
    """「□整平或填挖基地　□開挖水溝…」→ 選到的換 ■。selected 可為 list[str] 或 dict[str,bool]。"""
    if not isinstance(text, str) or not selected:
        return text
    chosen = [k for k, v in selected.items() if v] if isinstance(selected, dict) else list(selected)
    for k in chosen:
        text = re.sub(r"□(\s*" + re.escape(str(k)) + ")", r"■\1", text, count=1)
    return text


# ------------------------------------------------------------------ 表5-1


def _template_group_rows(ws: Worksheet, col_a: str = "A", col_b: str = "B") -> tuple[dict[int, list[int]], dict[int, int], int]:
    """掃描範本：A 欄「主要項目(n)」→ 該組細項列；B 欄「百分比小計」→ 小計列；「影響地價區域因素總修正數」→ 總計列。"""
    groups: dict[int, list[int]] = {}
    subtotal: dict[int, int] = {}
    cur: int | None = None
    total_row = 0
    for r in range(1, ws.max_row + 1):
        a = ws[f"{col_a}{r}"].value
        b = ws[f"{col_b}{r}"].value
        if isinstance(a, str):
            m = re.search(r"\((\d)\)", a)
            if m and "總修正數" not in a:
                cur = int(m.group(1))
                groups.setdefault(cur, [])
            elif "總修正數" in a:
                total_row = r
                cur = None
        if cur is not None:
            if isinstance(b, str) and "百分比小計" in b:
                subtotal[cur] = r
                cur = None
            elif r not in (subtotal.get(cur),):
                groups[cur].append(r)
    return groups, subtotal, total_row


def fill_table5(data: dict, tables: dict[int, Table5], rs: RuleSet, meta: dict) -> Workbook:
    wb, ws = _open("t5")
    case = data["case"]
    comps = sorted(tables.items())[:3]
    labels = meta.get("comp_labels") or {}
    _set(ws, "A2", f"案號：{case.get('case_no', '')}")
    subject_sec = comps[0][1].subject_section if comps else data["subject_parcel"].get("section_id", "")
    _set(ws, "C3", subject_sec)
    blocks = [("E", "G", "G2", "E3"), ("H", "J", "J2", "H3"), ("K", "M", "M2", "K3")]      # (等級欄, 修正欄, 實例編號格, 區段格)
    for k, (lv_col, pct_col, label_cell, sec_cell) in enumerate(blocks):
        if k < len(comps):
            comp_no, t5 = comps[k]
            _set(ws, label_cell, labels.get(comp_no, comp_no))
            _set(ws, sec_cell, t5.comparable_section)
    groups, subtotal_rows, total_row = _template_group_rows(ws)
    rows_by_rule = {comp_no: {row.rule_id: row for row in t5.rows} for comp_no, t5 in comps}
    warnings: list[str] = []
    for g in sorted(rs.groups):
        rules = [ru for ru in rs.rules if ru.group == g]
        trows = groups.get(g, [])
        if len(rules) > len(trows):
            warnings.append(f"範本第 {g} 組只有 {len(trows)} 列，基準表有 {len(rules)} 項，多出的未寫入：{'、'.join(ru.name for ru in rules[len(trows):])}")
        for ru, r in zip(rules, trows, strict=False):
            if not ws[f"B{r}"].value:
                _set(ws, f"B{r}", ru.name)
            srow = rows_by_rule[comps[0][0]].get(ru.id) if comps else None
            if srow and not ru.is_manual:
                _set(ws, f"C{r}", level_label(ru, srow.subject_level))
            for k, (lv_col, pct_col, _l, _s) in enumerate(blocks):
                if k >= len(comps):
                    continue
                row = rows_by_rule[comps[k][0]].get(ru.id)
                if row and not ru.is_manual:
                    _set(ws, f"{lv_col}{r}", level_label(ru, row.comparable_level))
                    _set(ws, f"{pct_col}{r}", row.pct if row.pct is not None else "-", fmt=PCT)
        sr = subtotal_rows.get(g)
        if sr:
            for k, (_lv, pct_col, _l, _s) in enumerate(blocks):
                if k < len(comps):
                    _set(ws, f"{pct_col}{sr}", comps[k][1].group_subtotals.get(g, 0.0), fmt=PCT)
    if total_row:
        for k, (_lv, pct_col, _l, _s) in enumerate(blocks):
            if k < len(comps):
                _set(ws, f"{pct_col}{total_row}", comps[k][1].total_pct, fmt=PCT)
    notes = meta.get("notes") or {}
    note_row = total_row + 1
    _set(ws, f"C{note_row}", notes.get("subject") or "", wrap=True)
    for k, (lv_col, _p, _l, _s) in enumerate(blocks):
        if k < len(comps):
            _set(ws, f"{lv_col}{note_row}", (notes.get("comparables") or {}).get(str(comps[k][0]), ""), wrap=True)
    _set(ws, f"C{note_row + 1}", notes.get("table5_case") or notes.get("case") or "", wrap=True)
    wb.properties.subject = "；".join(warnings)
    return wb


# ------------------------------------------------------------------ 表4


T4_ITEM_ROW = {7: 9, 8: 10, 9: 11, 10: 12, 11: 13, 12: 14, 13: 15, 14: 16, 15: 17, 16: 18, 17: 19, 18: 20, 19: 21, 20: 22, 21: 23,
               22: 24, 23: 25, 24: 26, 25: 27, 6: 28}
T4_NAME_NUM_ROWS = set(range(15, 24))            # 道路種類～停車方便性：名稱在第一格、數字在第二格、第三格是單位 M
T4_BLOCKS = [("G", "H", "J"), ("K", "L", "N"), ("O", "P", "R")]   # (值格, 數字格, 差異率格)


def _split_obs(v: Any) -> tuple[str, float | int | None]:
    """觀測值 → (顯示文字, 數字)。設施 dict → (名稱, 距離)；寬度 dict → ('', 寬度)；純數字 → ('', n)。"""
    if isinstance(v, list):
        v = v[0] if v else None
    if isinstance(v, dict):
        return (str(v.get("name") or ""), _num(v))
    if isinstance(v, bool) or v is None:
        return (_txt(v), None)
    if isinstance(v, (int, float)):
        return ("", _num(v))
    return (str(v), None)


def fill_table4(data: dict, t4: Table4, rs: RuleSet, meta: dict) -> Workbook:
    wb, ws = _open("t4")
    case, subject = data["case"], data["subject_parcel"]
    comps = t4.comparables[:3]
    comp_data = {c["comp_no"]: c for c in data.get("comparables") or []}
    labels = meta.get("comp_labels") or {}
    _set(ws, "L1", case.get("valuation_date", ""))
    _set(ws, "P1", case.get("case_no", ""))
    _set(ws, "F2", subject.get("serial_no") or "")
    for k, (vc, nc, pc) in enumerate(T4_BLOCKS):
        if k >= len(comps):
            continue
        c = comps[k]
        cd = comp_data.get(c.comp_no, {})
        _set(ws, {"G": "J2", "K": "N2", "O": "R2"}[vc], labels.get(c.comp_no, c.comp_no))
        _set(ws, f"{vc}4", cd.get("address") or c.parcel_id)
        _set(ws, f"{vc}5", c.normal_unit_price, fmt=MONEY)
        _set(ws, f"{vc}6", cd.get("transaction_date") or "")
        _set(ws, f"{pc}6", c.date_adjustment_pct, fmt=PCT)
        _set(ws, f"{vc}7", c.price_at_valuation_date, fmt=MONEY)
        _set(ws, f"{vc}8", c.section_id)
        _set(ws, f"{pc}8", c.regional_adjustment_pct, fmt=PCT)
    _set(ws, "D4", subject.get("address") or subject.get("parcel_id") or "")
    _set(ws, "D8", subject.get("section_id") or "")
    rows_by_item = {c.comp_no: {row.item_no: row for row in c.rows} for c in comps}
    for ru in rs.rules:
        r = T4_ITEM_ROW.get(ru.item_no or 0)
        if not r:
            continue
        sv = _parcel_obs(ru, subject)
        name, n = _split_obs(sv)
        if ru.item_no == 14 and isinstance(subject.get("front_road"), dict):
            name = subject["front_road"].get("name") or name
        if r in T4_NAME_NUM_ROWS:
            if name:
                _set(ws, f"D{r}", name)
            if n is not None:
                _set(ws, f"E{r}", n)
        else:
            _set(ws, f"D{r}", name if name else (n if n is not None else ""))
        for k, (vc, nc, pc) in enumerate(T4_BLOCKS):
            if k >= len(comps):
                continue
            row = rows_by_item[comps[k].comp_no].get(ru.item_no)
            cd = comp_data.get(comps[k].comp_no, {})
            cv = _parcel_obs(ru, cd) if cd else (row.comparable_value if row else None)
            cname, cn = _split_obs(cv)
            if ru.item_no == 14 and isinstance(cd.get("front_road"), dict):
                cname = cd["front_road"].get("name") or cname
            if r in T4_NAME_NUM_ROWS:
                if cname:
                    _set(ws, f"{vc}{r}", cname)
                if cn is not None:
                    _set(ws, f"{nc}{r}", cn)
            else:
                _set(ws, f"{vc}{r}", cname if cname else (cn if cn is not None else ""))
            _set(ws, f"{pc}{r}", row.pct if (row and row.pct is not None) else "-", fmt=PCT)
    for k, (vc, nc, pc) in enumerate(T4_BLOCKS):
        if k >= len(comps):
            continue
        c = comps[k]
        _set(ws, f"{vc}29", c.individual_total_pct, fmt=PCT)
        _set(ws, f"{vc}30", c.abs_sum_pct, fmt=PCT)
        _set(ws, {"G": "I30", "K": "M30", "O": "Q30"}[vc], c.similarity or "")
        _set(ws, f"{vc}31", round(c.trial_price), fmt=MONEY)
        _set(ws, {"G": "I31", "K": "M31", "O": "Q31"}[vc], f"{c.weight_pct:g}%" if c.weight_pct is not None else "")
    _set(ws, "G32", t4.subject_comparison_price, fmt=MONEY)
    notes = meta.get("notes") or {}
    _set(ws, "D33", notes.get("subject") or "", wrap=True)
    for k, (vc, _n, _p) in enumerate(T4_BLOCKS):
        if k < len(comps):
            c = comps[k]
            txt = (notes.get("comparables") or {}).get(str(c.comp_no)) or (comp_data.get(c.comp_no, {}).get("date_adjustment") or {}).get("note") or ""
            _set(ws, f"{vc}33", txt, wrap=True)
    _set(ws, "D34", notes.get("case") or "", wrap=True)
    a35 = ws[_anchor(ws, "A35")].value or ""
    if meta.get("fill_date"):
        ws[_anchor(ws, "A35")].value = a35.replace("填寫日期：", f"填寫日期：{meta['fill_date']}", 1)
    _set(ws, "K36", f"不動產估價師：{meta.get('appraiser', '')}")
    return wb


# ------------------------------------------------------------------ 表3


def _levels(sv: dict, regional: RuleSet, override: dict) -> dict[str, tuple[int | None, int]]:
    out: dict[str, tuple[int | None, int]] = {}
    for rule in regional.rules:
        f = rule.survey_field
        if not f or rule.is_manual:
            continue
        if f in override and isinstance(override[f], dict):
            out[f] = (override[f].get("num"), override[f].get("of") or len(rule.levels))
            continue
        try:
            out[f] = (level_number(rule, grade(rule, _get(sv, f))), len(rule.levels))
        except GradeError:
            out[f] = (None, len(rule.levels))
    return out


# 左半：(survey_field, 等級數字格, 等級總數格)
T3_LEVEL_CELLS = [
    ("land_control.urban_plan", "B4", "C4"), ("land_control.zoning", "B5", "C5"), ("land_control.bcr", "B6", "C6"), ("land_control.far", "B7", "C7"),
    ("land_control.building_prohibited", "B8", "C8"), ("land_control.building_restricted", "B9", "C9"),
    ("transport.main_road_width_m", "B11", "C11"), ("transport.avg_road_width_m", "B12", "C12"), ("transport.major_station", "B13", "C13"),
    ("transport.bus_stop", "B17", "C17"), ("transport.interchange", "B19", "C19"), ("transport.road_development", "B23", "C23"),
    ("natural.sunlight", "B24", "C24"), ("natural.view", "B25", "C25"), ("natural.slope", "B26", "C26"), ("natural.drainage", "B27", "C27"), ("natural.terrain", "B28", "C28"),
    ("improvement.building_site", "B31", "C31"), ("public.school", "B35", "C35"), ("public.market", "B39", "C39"), ("public.park", "B42", "C42"),
    ("public.tourism", "M4", "N4"), ("public.parking", "M6", "N6"), ("public.service", "M8", "N8"),
    ("special.utility", "M14", "N14"), ("special.funeral", "M18", "N18"), ("special.waste", "M22", "N22"), ("pollution.source", "M25", "N25"),
    ("commerce.department_store", "M30", "N30"), ("commerce.bank", "M32", "N32"), ("commerce.entertainment", "M34", "N34"), ("commerce.hotel", "M36", "N36"),
    ("commerce.foot_traffic", "M38", "N38"), ("commerce.shop_ratio_pct", "M39", "N39"), ("other", "M40", "N40"),
]


def _fill_survey_sheet(ws: Worksheet, section: dict, case: dict, regional: RuleSet, meta: dict) -> None:
    sv = section.get("survey") or {}
    lc, tr, na, pu, sp, po, co = (sv.get(k) or {} for k in ("land_control", "transport", "natural", "public", "special", "pollution", "commerce"))
    ex, im = sv.get("extra") or {}, sv.get("improvement") or {}
    _set(ws, "A2", f"{case.get('district', '')}（{case.get('land_use', '')}）")
    _set(ws, "B3", case.get("valuation_date", ""))
    _set(ws, "G3", section.get("section_id", ""))
    _set(ws, "L3", section.get("range_desc", ""), wrap=True)
    # 等級數字
    for f, num_cell, of_cell in T3_LEVEL_CELLS:
        lv = _levels(sv, regional, meta.get("level_numbers") or {}).get(f)
        if lv:
            n, of = lv
            if n is not None:
                _set(ws, num_cell, n)
            _set(ws, of_cell, of)
    # 土地使用管制
    _set(ws, "H4", _txt(lc.get("urban_plan")))
    _set(ws, "H5", _txt(lc.get("zoning")))
    if _num(lc.get("bcr")) is not None:
        _set(ws, "H6", f"{_num(lc.get('bcr'))}%")
    if _num(lc.get("far")) is not None:
        _set(ws, "H7", f"{_num(lc.get('far'))}%")
    _set(ws, "H8", {True: "有", False: "無"}.get(lc.get("building_prohibited"), ""))
    _set(ws, "H9", {True: "有", False: "無"}.get(lc.get("building_restricted"), ""))
    # 交通運輸
    mr = tr.get("main_road_width_m")
    if isinstance(mr, dict):
        _set(ws, "G11", mr.get("name") or "")
        if _num(mr) is not None:
            _set(ws, "J11", _num(mr))
    elif _num(mr) is not None:
        _set(ws, "J11", _num(mr))
    if _num(tr.get("avg_road_width_m")) is not None:
        _set(ws, "G12", _num(tr.get("avg_road_width_m")))
    stations = tr.get("major_station") or []
    for row, types in ((13, ("hsr_station",)), (14, ("rail_station",)), (15, ("intercity_bus_station",)), (16, ("mrt_station",))):
        f = _first(stations, types)
        _radio(ws, f"E{row}", bool(f))
        if f:
            _set(ws, f"F{row}", f.get("name") or "")
            _mark(ws, f"H{row}", f.get("in_section"))
            if f.get("in_section") is False and _num(f.get("distance_m")) is not None:
                _set(ws, f"J{row}", _num(f.get("distance_m")))
    bs = _first(tr.get("bus_stop"))
    _fac(ws, bs, "F17", "H17", "J17", prefix="")
    dens = (bs or {}).get("density")
    if dens:
        g18 = ws[_anchor(ws, "G18")].value or ""
        ws[_anchor(ws, "G18")].value = g18.replace(f"○{dens}", f"●{dens}", 1)
    _fac(ws, _first(tr.get("interchange")), "F19", "I19", "J19", prefix="")
    _set(ws, "I20", _txt(ex.get("settlement")))
    _set(ws, "I21", _txt(ex.get("distribution_center")))
    _set(ws, "I22", _txt(ex.get("consumer_market")))
    _set(ws, "I23", _txt(tr.get("road_development")))
    # 自然條件
    for row, key, src in ((24, "sunlight", na), (25, "view", na), (26, "slope", na), (27, "drainage", na), (28, "terrain", na), (29, "wind", ex), (30, "soil", ex)):
        _set(ws, f"I{row}", _txt(src.get(key)))
    # 土地改良
    for cell in ("E31", "E32"):
        ws[_anchor(ws, cell)].value = _checks(ws[_anchor(ws, cell)].value, im.get("building_site"))
    for cell in ("E33", "E34"):
        ws[_anchor(ws, cell)].value = _checks(ws[_anchor(ws, cell)].value, im.get("farmland"))
    # 公共建設（左）
    def options(rows: list[int], facs: list[dict], opts: list[tuple[str, tuple[str, ...] | None]]) -> None:
        used: set[int] = set()
        for row, (label, types) in zip(rows, opts, strict=False):
            hit = None
            for i, f in enumerate(facs or []):
                if i in used:
                    continue
                if types is None or f.get("type") in types or (f.get("subtype") == label):
                    hit, _ = f, used.add(i)
                    break
            _radio(ws, f"E{row}", bool(hit))
            _set(ws, f"F{row}", f"{label}：{hit['name']}" if hit and hit.get("name") else label)
            if hit:
                _mark(ws, f"I{row}", hit.get("in_section"))
                if hit.get("in_section") is False and _num(hit.get("distance_m")) is not None:
                    _set(ws, f"J{row}", _num(hit.get("distance_m")))

    options([35, 36, 37, 38], pu.get("school") or [], [("國小", ("school",)), ("國中", ("school",)), ("高中", ("school",)), ("大專院校", ("school",))])
    options([39, 40, 41], pu.get("market") or [], [("傳統市場", ("traditional_market",)), ("超級市場", ("supermarket",)), ("超大型購物中心", ("hypermarket",))])
    options([42, 43, 44], pu.get("park") or [], [("里鄰公園", ("park",)), ("一般公園", ("park",)), ("廣場.徒步區", ("plaza", "pedestrian_zone"))])
    # 右半：公共建設
    _fac(ws, _first(pu.get("tourism")), "R4", "T4", "U4")
    _fac(ws, _first(pu.get("parking")), "R6", "T6", "U6")
    _fac(ws, _first(pu.get("service")), "R8", "T8", "U8")
    _set(ws, "R10", _txt(ex.get("power")))
    _set(ws, "R11", _txt(ex.get("industrial_water")))
    _set(ws, "R12", _txt(ex.get("industrial_waste")))
    # 特殊設施
    util = sp.get("utility") or []
    _fac(ws, _first(util, ("substation", "hv_tower")), "R14", "T14", "U14")
    _fac(ws, _first(util, ("gas_tank", "oil_tank", "gas_station")), "R16", "T16", "U16")
    fun = sp.get("funeral") or []
    for row, t in ((18, "cemetery"), (19, "funeral_home"), (20, "crematorium"), (21, "columbarium")):
        f = _first(fun, (t,))
        _radio(ws, f"P{row}", bool(f))
        _fac(ws, f, f"R{row}", f"T{row}", f"U{row}")
    waste = sp.get("waste") or []
    for row, t in ((22, "sewage_plant"), (23, "landfill"), (24, "incinerator")):
        f = _first(waste, (t,))
        _radio(ws, f"P{row}", bool(f))
        _fac(ws, f, f"R{row}", f"T{row}", f"U{row}")
    srcs = po.get("source") or []
    for row, lab in ((25, "水污染"), (26, "噪音污染"), (27, "廢氣污染"), (28, "廢棄物污染"), (29, "其他污染")):
        f = next((x for x in srcs if x.get("subtype") == lab), None)
        _radio(ws, f"O{row}", bool(f))
        _fac(ws, f, f"R{row}", f"T{row}", f"U{row}")
    # 工商活動
    for row, key in ((30, "department_store"), (32, "bank"), (34, "entertainment"), (36, "hotel")):
        f = _first(co.get(key))
        _set(ws, f"S{row}", f["name"] if f and f.get("name") else "無")
        if f:
            _set(ws, f"S{row + 1}", f.get("count", 1))
            _mark(ws, f"T{row + 1}", f.get("in_section"))
            if f.get("in_section") is False and _num(f.get("distance_m")) is not None:
                _set(ws, f"U{row + 1}", _num(f.get("distance_m")))
    _set(ws, "R38", _txt(co.get("foot_traffic")))
    if co.get("shop_ratio_pct") is not None:
        _set(ws, "R39", f"{_num(co.get('shop_ratio_pct'))}%以上作為店舖")
    _set(ws, "Q40", _txt(sv.get("other")), wrap=True)
    _set(ws, "R42", _txt(ex.get("building_density")))
    _set(ws, "R43", _txt(ex.get("building_type")))
    lus = ex.get("land_use_status")
    if lus:
        chosen = [k for k, v in lus.items() if v] if isinstance(lus, dict) else list(lus)
        q44 = ws[_anchor(ws, "Q44")].value or ""
        for k in chosen:
            q44 = q44.replace(f"○{k}", f"●{k}", 1)
        ws[_anchor(ws, "Q44")].value = q44
    _set(ws, "A45", f"勘查日期：{section.get('survey_date') or ''}")
    _set(ws, "R46", f"不動產估價師：{meta.get('appraiser', '')}")


def fill_table3(data: dict, regional: RuleSet, meta: dict, section_ids: list[str] | None = None) -> Workbook:
    """每個區段一張工作表（複製範本工作表），比準地區段排第一。"""
    wb, ws = _open("t3")
    case = data["case"]
    sections = data.get("sections") or {}
    ids = section_ids or [data["subject_parcel"].get("section_id")] + [s for s in sections if s != data["subject_parcel"].get("section_id")]
    ids = [s for s in ids if s in sections]
    if not ids:
        _fill_survey_sheet(ws, {}, case, regional, meta)
        return wb
    made = []
    for sid in ids:
        target = wb.copy_worksheet(ws)
        target.title = f"表3 {sid}"[:31]
        target.sheet_view.zoomScale = ws.sheet_view.zoomScale
        _fill_survey_sheet(target, sections[sid], case, regional, meta)
        made.append(target)
    wb.remove(ws)
    wb.active = 0
    return wb


# ------------------------------------------------------------------ 打包


# ------------------------------------------------------------------ 表2 收益法調查估價表（含附表─成本法）與表14 比準地地價估計表


def _roc_text(v: Any) -> str:
    t = str(v or "").strip()
    return f"{int(t[:-4])}年{t[-4:-2]}月{t[-2:]}日" if t.isdigit() and len(t) >= 6 else t


def _frac(v: Any) -> float | None:
    n = _num(v)
    return None if n is None else round(n / 100, 6)


def fill_table2(data: dict, inc: dict, meta: dict) -> Workbook:
    """範本「94表2收益法」＋「95表2收益附-成本法」：值一律由 engine/income.compute_income 算好寫入（覆蓋範本公式），比率格寫小數（範本格式為百分比）。
    素地（inc.mode == "land"）建物相關欄位免填（手冊 p.38 (3)），附表成本法移除。格位對照見 docs/12。"""
    path = TEMPLATES["t2"][0]
    wb = load_workbook(path)
    for other in list(wb.worksheets):
        if other.title not in (TEMPLATES["t2"][1], T2_COST_SHEET):
            wb.remove(other)
    ws = wb[TEMPLATES["t2"][1]]
    for sh in wb.worksheets:                                   # 範本這兩張是隱藏工作表：輸出時設為顯示，表2 為使用中工作表
        sh.sheet_state = "visible"
    wb.active = wb.worksheets.index(ws)
    case = data.get("case") or {}
    sp = data.get("subject_parcel") or {}
    subj = (data.get("income") or {}).get("subject") or {}
    building = inc.get("mode") == "building"
    dash = "-"
    _set(ws, "J1", case.get("valuation_date"))
    _set(ws, "P1", sp.get("section_id"))
    _set(ws, "U1", sp.get("serial_no") or "")
    _set(ws, "C2", _txt(sp.get("address") or sp.get("parcel_id")))
    _set(ws, "J2", subj.get("address") or (dash if not building else ""))
    _set(ws, "T2", subj.get("building_no") or (dash if not building else ""))
    _set(ws, "C3", inc.get("income_area_m2"))
    _set(ws, "E3", inc.get("land_share_m2"))
    _set(ws, "G3", subj.get("total_floors") if building else dash)
    _set(ws, "J3", subj.get("level") if building else dash)
    _set(ws, "P3", round(inc["rebuild_total"]) if building and inc.get("rebuild_total") else dash, MONEY)
    _set(ws, "T3", round(inc["cost_total"]) if building and inc.get("cost_total") else dash, MONEY)
    _set(ws, "F4", inc.get("est_monthly_rent"), MONEY)
    _set(ws, "R4", _roc_text(case.get("valuation_date")))
    for c, k in (("F5", "annual_rent"), ("R5", "annual_rent"), ("F6", "deposit"), ("F7", "deposit_income"), ("R7", "other_income"), ("F8", "gross_income"), ("K9", "egi")):
        _set(ws, c, inc.get(k), MONEY)
    _set(ws, "R6", _frac(inc.get("deposit_rate_pct")))
    _set(ws, "R8", inc.get("idle_months"))
    exp = inc.get("expenses") or {}
    _set(ws, "F10", exp.get("land_value_tax"), MONEY)
    _set(ws, "F12", round(exp.get("management") or 0, 2), MONEY)
    _set(ws, "R12", exp.get("other"), MONEY)
    if building:
        _set(ws, "F11", exp.get("house_tax"), MONEY)
        _set(ws, "F13", round(exp.get("insurance") or 0, 2), MONEY)
        _set(ws, "R10", exp.get("maintenance"), MONEY)
        _set(ws, "R11", exp.get("replacement"), MONEY)
        _set(ws, "R13", inc.get("depreciation_rate"))
        _set(ws, "F16", _frac(inc.get("building_cap_rate_pct")))
        _set(ws, "F17", inc.get("building_noi"), MONEY)
        _set(ws, "O15", (inc.get("sources") or {}).get("cap_rate_building_pct", {}).get("source", ""), wrap=True)
    else:
        for c in ("F11", "F13", "R10", "R11", "R13", "F16", "F17", "O15"):
            _set(ws, c, dash)
    _set(ws, "K14", inc.get("total_expense"), MONEY)
    _set(ws, "F15", inc.get("noi"), MONEY)
    _set(ws, "F18", inc.get("other_deduction"), MONEY)
    _set(ws, "R18", inc.get("land_noi"), MONEY)
    _set(ws, "F19", _frac(inc.get("land_cap_rate_pct")))
    _set(ws, "O19", (inc.get("sources") or {}).get("cap_rate_land_pct", {}).get("source", ""), wrap=True)
    _set(ws, "F20", inc.get("land_income_total"), MONEY)
    _set(ws, "F21", inc.get("land_income_unit"), MONEY)
    _set(ws, "F22", inc.get("income_price") if building else dash, MONEY)
    rows = inc.get("examples") or []
    for i in range(3):
        r = 24 + i
        ex = rows[i] if i < len(rows) else None
        vals = [("B", ex and ex.get("example_no")), ("C", ex and ex.get("section_id")), ("D", ex and ex.get("area_m2")), ("E", ex and ex.get("total_rent")),
                ("F", ex and ex.get("unit_rent")), ("G", ex and ex.get("rent_type")), ("H", ex and _frac(ex.get("situation_pct"))), ("I", ex and ex.get("rent_date")),
                ("J", ex and _frac(ex.get("date_pct"))), ("K", ex and _frac(ex.get("regional_pct"))), ("M", ex and _frac(ex.get("individual_pct"))),
                ("O", ex and _frac(ex.get("abs_sum_pct"))), ("Q", ex and ex.get("n_adjusted")), ("R", ex and ex.get("trial_rent")), ("T", ex and _frac(ex.get("weight_pct")))]
        for col, v in vals:
            _set(ws, f"{col}{r}", v if ex else None)
    _set(ws, "U24", inc.get("est_monthly_rent"), MONEY)
    expl = [f"編號{ex.get('example_no')}：{'；'.join(ex.get('flags') or [])}" + (f"，情況調整 {ex.get('situation_pct'):+.2f}%" if ex.get("situation_pct") else "")
            for ex in rows if ex.get("flags") or ex.get("situation_pct")]
    _set(ws, "C27", "；".join(expl) or "無特殊情況", wrap=True)
    checks = [f"{k}：{v.get('value')}（{v.get('source')}）" for k, v in (inc.get("sources") or {}).items() if v.get("check")]
    _set(ws, "A30", "備註欄：" + "；".join(filter(None, ["系統推定參數需估價師確認：" + "；".join(checks) if checks else "", "；".join(inc.get("issues") or [])])), wrap=True)
    _set(ws, "A32", f"填寫日期：{meta.get('fill_date') or '      年      月      日'}         承辦員：")
    _set(ws, "N33", f"不動產估價師：{meta.get('appraiser') or ''}")
    wa = wb[T2_COST_SHEET]
    if not building or not inc.get("cost"):
        wb.remove(wa)
        wb.active = 0
        return wb
    c = inc["cost"]
    b = subj.get("building") or {}
    _set(wa, "N1", sp.get("serial_no") or "")
    _set(wa, "D3", subj.get("address") or "")
    _set(wa, "D4", _txt(sp.get("parcel_id")))
    _set(wa, "D5", subj.get("building_no") or "")
    _set(wa, "H5", c.get("structure"))
    _set(wa, "D6", c.get("floors_above"))
    _set(wa, "D7", c.get("floors_below"))
    _set(wa, "H6", c.get("reg_area_m2") or inc.get("income_area_m2"))
    _set(wa, "H7", c.get("calc_area_m2"))
    _set(wa, "D8", subj.get("level"))
    _set(wa, "H8", c.get("years"))
    for cell, key in (("M3", "unit_cost_m2"), ("L4", "adj_rate"), ("M5", "unit_cost_adj"), ("L6", "design_rate"), ("M6", "design"), ("M7", "cum1"),
                      ("L8", "adv_rate"), ("M8", "advertising"), ("L9", "mgmt_rate"), ("M9", "management"), ("L10", "tax_rate"), ("M10", "tax"), ("M11", "cum2"),
                      ("D10", "own_rate"), ("E10", "own_ratio"), ("G10", "w_own"), ("D11", "loan_rate"), ("E11", "loan_ratio"), ("G11", "w_loan"), ("E12", "presale_ratio"),
                      ("G12", "w_presale"), ("I10", "avg_rate"), ("G13", "installment"), ("G14", "years"), ("G15", "interest_rate"), ("L12", "interest_rate"), ("M12", "interest"),
                      ("M13", "cum3"), ("L14", "profit_rate"), ("M14", "profit"), ("M16", "rebuild_unit"), ("D18", "age"), ("H18", "life"), ("H19", "residual"),
                      ("M20", "cost_total_calc")):
        _set(wa, cell, c.get(key))
    _set(wa, "D12", 0)
    _set(wa, "N4", _roc_text(case.get("valuation_date")))
    _set(wa, "J11", round((c.get("adv_rate") or 0) + (c.get("mgmt_rate") or 0) + (c.get("tax_rate") or 0), 6))
    _set(wa, "J13", round(1 + (c.get("interest_rate") or 0), 6))
    _set(wa, "J15", round(1 + (c.get("profit_rate") or 0), 6))
    _set(wa, "F18", c.get("remaining"))
    _set(wa, "H20", round(c.get("accum_dep_unit") or 0, 2))
    _set(wa, "M18", round(c.get("cost_unit") or 0, 2))
    comp = str(b.get("completed") or "")
    _set(wa, "A18", int(comp[:-4]) if comp.isdigit() and len(comp) >= 6 else "")
    _set(wa, "C18", int(comp[-4:-2]) if comp.isdigit() and len(comp) >= 6 else "")
    _set(wa, "H22", round((c.get("unit_cost_adj") or 0) * 3.305785))
    _set(wa, "A24", f"填寫日期：{meta.get('fill_date') or ''}         承辦員：")
    return wb


def fill_table14(data: dict, t4: Table4, decision: dict, meta: dict) -> Workbook:
    """範本「114表14比準地地價估計表」：比較價格×權重＋收益價格×權重 → 比準地地價（值由 engine/income.land_price_decision 算好寫入）。"""
    from app.spatial.cadastre import split_parcel_id
    wb, ws = _open("t14")
    ws.sheet_state = "visible"
    wb.active = 0
    case = data.get("case") or {}
    sp = data.get("subject_parcel") or {}
    sec, lot = split_parcel_id(sp.get("parcel_id") or "")
    _set(ws, "B2", case.get("case_no"))
    _set(ws, "I2", case.get("valuation_date"))
    for r in range(5, 10):
        for col in "ABCDEFGHIJK":
            _set(ws, f"{col}{r}", None)
    _set(ws, "A5", sp.get("section_id"))
    _set(ws, "B5", sp.get("serial_no") or "")
    _set(ws, "C5", (case.get("district") or "").replace("新北市", ""))
    _set(ws, "D5", sec)
    _set(ws, "E5", lot)
    _set(ws, "F5", decision.get("comparison_price"), MONEY)
    _set(ws, "G5", decision.get("comparison_weight"))
    _set(ws, "H5", decision.get("income_price") if decision.get("income_price") is not None else "-", MONEY)
    _set(ws, "I5", decision.get("income_weight"))
    _set(ws, "J5", decision.get("land_price"), MONEY)
    _set(ws, "K5", decision.get("reason") or "；".join(decision.get("issues") or []), wrap=True)
    _set(ws, "A10", f"填寫日期：{meta.get('fill_date') or '   年   月   日'}    承辦員：")
    _set(ws, "J11", f"不動產估價師：{meta.get('appraiser') or ''}")
    return wb


def official_workbooks(data: dict, result: dict, regional: RuleSet, individual: RuleSet, meta: dict) -> dict[str, Workbook]:
    out = {"t3": fill_table3(data, regional, meta), "t5": fill_table5(data, result["table5"], regional, meta), "t4": fill_table4(data, result["table4"], individual, meta)}
    if result.get("income"):                                     # 收益法啟用：加表2（含附表成本法）與表14 比準地地價估計表
        out["t2"] = fill_table2(data, result["income"], meta)
        out["t14"] = fill_table14(data, result["table4"], result.get("land_price_decision") or {}, meta)
    return out


def official_zip(data: dict, result: dict, regional: RuleSet, individual: RuleSet, meta: dict) -> bytes:
    case_no = data["case"].get("case_no") or "case"
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as z:
        for key, wb in official_workbooks(data, result, regional, individual, meta).items():
            b = io.BytesIO()
            wb.save(b)
            z.writestr(f"{case_no}_{TEMPLATES[key][2]}.xlsx", b.getvalue())
    return buf.getvalue()


def workbook_bytes(wb: Workbook) -> bytes:
    b = io.BytesIO()
    wb.save(b)
    return b.getvalue()


__all__ = ["fill_table2", "fill_table3", "fill_table4", "fill_table5", "fill_table14", "official_workbooks", "official_zip", "workbook_bytes"]
