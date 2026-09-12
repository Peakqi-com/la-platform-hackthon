"""
表1 地價區段勘查表 → Excel（版面照範本第 1 頁：左右兩大欄、每列「等級數字 | 等級總數 | 細項 | 內容」）。

等級數字欄：估價師填的值若有（meta.level_numbers，來自 pdf_forms 的 submitted.table1）就用它；
否則用區域因素基準表對勘查事實判等級（rules.grade → level_number），判不出留白。
●○ 勾選、「名稱：X ○本區段內 ●本區段外(距 N M)」的寫法照範本。
"""
from __future__ import annotations

from typing import Any

from openpyxl.styles import Alignment, Border, Font
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from app.engine.rules import GradeError, RuleSet, grade, level_number
from app.engine.tables import _get

from .xlsx import FONT, HEAD_FILL, LEFT, _cell, _merge

TOP_LEFT = Alignment(horizontal="left", vertical="top", wrap_text=True)


def _r(on: bool | None) -> str:
    return "●" if on else "○"


def _d(f: dict) -> str:
    v = f.get("distance_m")
    return f"{int(v) if isinstance(v, (int, float)) and float(v).is_integer() else (v if v is not None else '')}"


def fac_line(f: dict | None, empty_label: str = "無") -> str:
    if not f or not f.get("name"):
        return f"名稱：{empty_label} ○本區段內 ○本區段外(距   M)"
    ins = f.get("in_section")
    return f"名稱：{f['name']} {_r(ins is True)}本區段內 {_r(ins is False)}本區段外(距 {_d(f) if ins is False else ''} M)"


def count_line(f: dict | None) -> str:
    if not f or not f.get("name"):
        return "名稱：無  數量:\n○本區段內 ○本區段外(距   M)"
    ins = f.get("in_section")
    return f"名稱：{f['name']}  數量:{f.get('count', 1)}\n{_r(ins is True)}本區段內 {_r(ins is False)}本區段外(距 {_d(f) if ins is False else ''} M)"


def option_lines(facs: list[dict], options: list[tuple[str, tuple[str, ...]]], name_first: bool = False) -> str:
    """大型車站／學校／市場／公園：每個選項一行；有對應設施 → ●名稱 + 區段內外，否則 ○選項名。"""
    used: set[int] = set()
    lines = []
    for label, types in options:
        hit = None
        for i, f in enumerate(facs):
            if i in used:
                continue
            if f.get("type") in types or (not types and True):
                hit = f
                used.add(i)
                break
        if hit:
            ins = hit.get("in_section")
            lines.append(f"● {hit['name']} {_r(ins is True)}本區段內 {_r(ins is False)}本區段外(距 {_d(hit) if ins is False else ''} M)")
        else:
            lines.append(f"○ {label} ○本區段內 ○本區段外(距   M)")
    return "\n".join(lines)


def checks_line(selected: Any, options: list[str], box: str = "□", on: str = "■") -> str:
    sel = set(selected or []) if isinstance(selected, list) else ({selected} if selected else set())
    return " ".join(f"{on if o in sel else box}{o}" for o in options)


def _num(v: Any, unit: str = "") -> str:
    if v is None:
        return ""
    if isinstance(v, dict):
        v = v.get("value")
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    return f"{v}{unit}" if v is not None else ""


def _txt(v: Any) -> str:
    return "" if v is None else str(v)


def _build_rows(sv: dict, regional: RuleSet, level_override: dict) -> tuple[list, list]:
    """回傳 (left_rows, right_rows)，每列 = (group, label, value_text, survey_field or None)。"""
    L: list[tuple] = []
    R: list[tuple] = []
    lc, tr, na, pu, sp, po, co = (sv.get(k) or {} for k in ("land_control", "transport", "natural", "public", "special", "pollution", "commerce"))
    ex, im = sv.get("extra") or {}, sv.get("improvement") or {}

    def first(lst, types=None):
        for f in lst or []:
            if types is None or f.get("type") in types:
                return f
        return None

    g = "土地使用管制"
    L += [(g, "都市計畫(內外)", _txt(lc.get("urban_plan")), "land_control.urban_plan"),
          (g, "使用分區(使用地類別)", _txt(lc.get("zoning")), "land_control.zoning"),
          (g, "建蔽率", _num(lc.get("bcr"), "%"), "land_control.bcr"),
          (g, "容積率", _num(lc.get("far"), "%"), "land_control.far"),
          (g, "有無禁止建築", {True: "有", False: "無"}.get(lc.get("building_prohibited"), ""), "land_control.building_prohibited"),
          (g, "有無限制建築(整體開發、面積限制、高度限制)", {True: "有", False: "無"}.get(lc.get("building_restricted"), ""), "land_control.building_restricted")]
    g = "交通運輸"
    mr = tr.get("main_road_width_m") or {}
    L += [(g, "主要道路", f"名稱：{_txt(mr.get('name') if isinstance(mr, dict) else '')}    寬度： {_num(mr)} M", "transport.main_road_width_m"),
          (g, f"區段內道路平均寬度  {_num(tr.get('avg_road_width_m'))} M", "", "transport.avg_road_width_m"),
          (g, "大型車站", option_lines(tr.get("major_station") or [], [("無高鐵站", ("hsr_station",)), ("無火車站", ("rail_station",)),
                                                                   ("無客運站", ("intercity_bus_station",)), ("無捷運站", ("mrt_station",))]), "transport.major_station")]
    bs = first(tr.get("bus_stop"))
    dens = (bs or {}).get("density")
    L += [(g, "站牌", fac_line(bs) + "\n密集程度： " + " ".join(f"{_r(dens == o)}{o}" for o in ("非常密集", "密集", "不密集")), "transport.bus_stop"),
          (g, "交流道", fac_line(first(tr.get("interchange")), "無交流道"), "transport.interchange"),
          (g, "接近聚落程度", _txt(ex.get("settlement")), None),
          (g, "接近運銷中心程度", _txt(ex.get("distribution_center")), None),
          (g, "接近消費市場程度", _txt(ex.get("consumer_market")), None),
          (g, "區段內道路規劃及闢建程度", _txt(tr.get("road_development")), "transport.road_development")]
    g = "自然條件"
    L += [(g, "日  照", _txt(na.get("sunlight")), "natural.sunlight"), (g, "景  觀", _txt(na.get("view")), "natural.view"),
          (g, "傾 斜 度", _txt(na.get("slope")), "natural.slope"), (g, "保（排）水之良否", _txt(na.get("drainage")), "natural.drainage"),
          (g, "地  勢", _txt(na.get("terrain")), "natural.terrain"), (g, "風  勢", _txt(ex.get("wind")), None), (g, "土  質", _txt(ex.get("soil")), None)]
    g = "土地改良"
    L += [(g, "建築基地改良", checks_line(im.get("building_site"), ["整平或填挖基地", "開挖水溝", "水土保持", "鋪築道路", "埋設管道", "修築駁嵌", "其他＿＿"]), "improvement.building_site"),
          (g, "農地改良", checks_line(im.get("farmland"), ["耕地整理", "水土保持", "土壤改良", "修築農路", "灌溉", "排水", "防風", "防砂", "堤防", "其他＿＿"]), "improvement.farmland")]
    g = "公共建設"
    schools = pu.get("school") or []
    L += [(g, "學  校", option_lines(schools, [("國小", ("school",)), ("國中", ("school",)), ("高中", ("school",)), ("大專院校", ("school",))]), "public.school"),
          (g, "市  場", option_lines(pu.get("market") or [], [("傳統市場", ("traditional_market",)), ("超級市場", ("supermarket",)), ("超大型購物中心", ("hypermarket",))]), "public.market"),
          (g, "公 園\n廣 場\n徒步區", option_lines(pu.get("park") or [], [("里鄰公園", ("park",)), ("一般公園", ("park",)), ("廣場.徒步區", ("plaza", "pedestrian_zone"))]), "public.park")]

    g = "公共建設"
    R += [(g, "觀光遊憩設施", fac_line(first(pu.get("tourism"))), "public.tourism"),
          (g, "停車場地", fac_line(first(pu.get("parking"))), "public.parking"),
          (g, "接近服務性設施的程度", fac_line(first(pu.get("service")), ""), "public.service"),
          (g, "電力資源", _txt(ex.get("power")), None), (g, "產業用水及設施", _txt(ex.get("industrial_water")), None),
          (g, "污廢水及廢棄物處理設施", fac_line(first(ex.get("industrial_waste")) if isinstance(ex.get("industrial_waste"), list) else None), None)]
    g = "特殊設施"
    util = sp.get("utility") or []
    R += [(g, "電業\n氣體\n燃料｜變電所或高壓鐵塔", fac_line(first(util, ("substation", "hv_tower"))), "special.utility"),
          (g, "電業\n氣體\n燃料｜瓦斯槽或儲油槽", fac_line(first(util, ("gas_tank", "oil_tank", "gas_station"))), "special.utility")]
    fun = sp.get("funeral") or []
    for lab, t in (("墓  地", "cemetery"), ("殯儀館", "funeral_home"), ("火葬場", "crematorium"), ("納骨塔", "columbarium")):
        f = first(fun, (t,))
        R.append((g, f"殯葬｜{_r(bool(f))}{lab}", fac_line(f), "special.funeral"))
    waste = sp.get("waste") or []
    for lab, t in (("污水處理場", "sewage_plant"), ("垃圾場或掩埋場", "landfill"), ("焚化爐", "incinerator")):
        f = first(waste, (t,))
        R.append((g, f"廢棄物處理｜{_r(bool(f))}{lab}", fac_line(f), "special.waste"))
    g = "環境污染"
    srcs = po.get("source") or []
    for lab in ("水污染", "噪音污染", "廢氣污染", "廢棄物污染", "其他污染"):
        f = next((x for x in srcs if x.get("subtype") == lab), None)
        R.append((g, f"{_r(bool(f))}{lab}", fac_line(f), "pollution.source"))
    g = "工商活動"
    R += [(g, "百貨公司", count_line(first(co.get("department_store"))), "commerce.department_store"),
          (g, "金融機構", count_line(first(co.get("bank"))), "commerce.bank"),
          (g, "娛樂設施", count_line(first(co.get("entertainment"))), "commerce.entertainment"),
          (g, "大型展示中心或觀光飯店", count_line(first(co.get("hotel"))), "commerce.hotel"),
          (g, "顧客之通行量", _txt(co.get("foot_traffic")), "commerce.foot_traffic"),
          (g, "店鋪之毗連狀態", f"{_num(co.get('shop_ratio_pct'))}%以上作為店舖" if co.get("shop_ratio_pct") is not None else "", "commerce.shop_ratio_pct")]
    R += [("其他影響因素", "", _txt(sv.get("other")), "other")]
    R += [("房屋建築現況", "建築密度", _txt(ex.get("building_density")), None), ("房屋建築現況", "建築型態", _txt(ex.get("building_type")), None)]
    R += [("土地利用現況", "", checks_line(ex.get("land_use_status"), ["商業用", "住宅用", "工業用", "住商混合", "住工混合", "農作用", "漁牧用", "空地", "公共設施", "其他＿＿"], "○", "●"), None)]
    return L, R


def _levels(sv: dict, regional: RuleSet, override: dict) -> dict[str, tuple[int | None, int | None]]:
    out: dict[str, tuple[int | None, int | None]] = {}
    for rule in regional.rules:
        f = rule.survey_field
        if not f or rule.is_manual:
            continue
        if f in override and isinstance(override[f], dict):
            out[f] = (override[f].get("num"), override[f].get("of") or len(rule.levels))
            continue
        try:
            lv = grade(rule, _get(sv, f))
            out[f] = (level_number(rule, lv), len(rule.levels))
        except GradeError:
            out[f] = (None, len(rule.levels))
    return out


def write_table1(ws: Worksheet, section: dict, case: dict, regional: RuleSet, meta: dict) -> None:
    ws.title = "表1"
    widths = [4, 3, 3, 16, 6, 6, 6, 6, 8, 4, 3, 3, 15, 8, 8, 12]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    sv = section.get("survey") or {}
    lv = _levels(sv, regional, meta.get("level_numbers") or {})
    _cell(ws, 1, 1, "表1  地價區段勘查表", align=LEFT, bold=True, size=14)
    ws.cell(row=1, column=1).border = Border()
    ws.cell(row=2, column=1, value=case.get("district", "")).font = Font(name=FONT, size=10)
    _merge(ws, 3, 1, 3, 2, "年\n期", fill=HEAD_FILL)
    _merge(ws, 3, 3, 3, 4, case.get("valuation_date", ""))
    _cell(ws, 3, 5, "區段編號", fill=HEAD_FILL)
    _merge(ws, 3, 6, 3, 7, section.get("section_id", ""))
    _merge(ws, 3, 8, 3, 9, "區段範圍\n(公共設施保留地請填明毗鄰非保留地區段號)", fill=HEAD_FILL, size=7)
    _merge(ws, 3, 10, 3, 16, section.get("range_desc") or "", align=TOP_LEFT)
    ws.row_dimensions[3].height = 36
    L, R = _build_rows(sv, regional, meta.get("level_numbers") or {})
    r0 = 4
    n = max(len(L), len(R))

    def put_half(rows: list, base_col: int, val_cols: int) -> None:
        gstart = None
        for i in range(n):
            r = r0 + i
            if i >= len(rows):
                for c in range(base_col, base_col + 4 + val_cols):
                    _cell(ws, r, c)
                continue
            group, label, value, field = rows[i]
            num, of = lv.get(field, (None, None)) if field else (None, None)
            _cell(ws, r, base_col + 1, num)
            _cell(ws, r, base_col + 2, of)
            lab_txt = label.split("｜")[-1] if "｜" in label else label
            _cell(ws, r, base_col + 3, lab_txt, align=LEFT, size=8)
            _merge(ws, r, base_col + 4, r, base_col + 3 + val_cols, value, align=TOP_LEFT, size=8)
            lines = value.count("\n") + 1
            ws.row_dimensions[r].height = max(ws.row_dimensions[r].height or 15, 13 * lines + 4)
            if gstart is None or rows[gstart - r0][0] != group:
                if gstart is not None:
                    _merge(ws, gstart, base_col, r - 1, base_col, "\n".join(rows[gstart - r0][0]), fill=HEAD_FILL, size=8)
                gstart = r
        if gstart is not None and rows:
            _merge(ws, gstart, base_col, r0 + len(rows) - 1, base_col, "\n".join(rows[gstart - r0][0]), fill=HEAD_FILL, size=8)

    put_half(L, 1, 5)
    put_half(R, 10, 3)
    r = r0 + n + 1
    for c, txt in ((1, f"勘查日期：{section.get('survey_date') or meta.get('survey_date', '')}"), (5, "承辦員："), (9, "課（股）長："), (13, "主任（局、處長）：")):
        ws.cell(row=r, column=c, value=txt).font = Font(name=FONT, size=9)
    ws.cell(row=r + 3, column=13, value=f"不動產估價師：{meta.get('appraiser', '')}").font = Font(name=FONT, size=9)
    ws.page_setup.orientation = "portrait"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 1
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_area = f"A1:P{r + 3}"
